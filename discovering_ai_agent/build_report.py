"""Build a black-text DOCX with six theory pages and executed notebook screenshots."""

import argparse
import json
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

BASE = Path(__file__).resolve().parent
BLACK = RGBColor(0, 0, 0)


def black_run(run, *, bold=False, size=None):
    run.font.name = "Times New Roman"
    run.font.color.rgb = BLACK
    run.bold = bold
    if size:
        run.font.size = Pt(size)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Times New Roman")
    return run


def paragraph(doc, text, *, style=None, centered=False, size=None):
    p = doc.add_paragraph(style=style)
    black_run(p.add_run(text), size=size)
    if centered:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Cm(0)
    return p


def table(doc, headers, rows):
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = False
    table.style = "Table Grid"
    widths = [4.2, 12.3] if len(headers) == 2 else [3.6, 5.7, 7.2]
    for row_index, values in enumerate([headers, *rows]):
        row = table.rows[0] if row_index == 0 else table.add_row()
        properties = row._tr.get_or_add_trPr()
        properties.append(OxmlElement("w:cantSplit"))
        if row_index == 0:
            properties.append(OxmlElement("w:tblHeader"))
        for index, (cell, text) in enumerate(zip(row.cells, values)):
            cell.width = Cm(widths[index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.08
            black_run(p.add_run(text), bold=row_index == 0, size=11.5)
            tcpr = cell._tc.get_or_add_tcPr()
            margins = OxmlElement("w:tcMar")
            for edge in ("top", "bottom", "left", "right"):
                value = OxmlElement("w:" + edge)
                value.set(qn("w:w"), "70")
                value.set(qn("w:type"), "dxa")
                margins.append(value)
            tcpr.append(margins)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = OxmlElement("w:" + edge)
        for key, value in {"val": "single", "sz": "4", "color": "000000"}.items():
            border.set(qn("w:" + key), value)
        borders.append(border)
    table._tbl.tblPr.append(borders)
    return table


def screenshot(doc, number, caption):
    p = paragraph(doc, caption, style="Caption", centered=True, size=10.5)
    p.paragraph_format.keep_with_next = True
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_after = Pt(5)
    run = p.add_run()
    image = run.add_picture(str(BASE / f"screenshots/cell-{number:02d}.png"), width=Cm(16.3))
    image._inline.docPr.set("descr", caption)


def heading(doc, text, level=1):
    return doc.add_heading(text, level=level)


def new_page(doc, title):
    heading(doc, title).paragraph_format.page_break_before = True


def configure(doc):
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin, section.bottom_margin = Cm(2.0), Cm(2.0)
    section.left_margin, section.right_margin = Cm(2.5), Cm(2.0)
    section.footer_distance = Cm(1.0)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size, normal.font.color.rgb = "Times New Roman", Pt(13), BLACK
    normal.paragraph_format.line_spacing = 1.2
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.first_line_indent = Cm(0.6)
    normal.paragraph_format.widow_control = True
    for style in doc.styles:
        if style.type in (1, 2):
            style.font.color.rgb = BLACK
            style.font.name = "Times New Roman"
            fonts = style._element.get_or_add_rPr().get_or_add_rFonts()
            for key in list(fonts.attrib):
                if "theme" in key.lower():
                    del fonts.attrib[key]
            for key in ("ascii", "hAnsi", "eastAsia", "cs"):
                fonts.set(qn("w:" + key), "Times New Roman")
        for border in list(style._element.iter(qn("w:pBdr"))):
            border.getparent().remove(border)
    for name, size in [("Heading 1", 16), ("Heading 2", 13.5), ("Title", 20), ("Subtitle", 14)]:
        style = doc.styles[name]
        style.font.name, style.font.size, style.font.color.rgb = "Times New Roman", Pt(size), BLACK
        style.font.bold = name != "Subtitle"
        style.paragraph_format.first_line_indent = Cm(0)
        style.paragraph_format.space_before = Pt(7 if name == "Heading 2" else 0)
        style.paragraph_format.space_after = Pt(5 if name == "Heading 2" else 8)
        style.paragraph_format.keep_with_next = True
    doc.styles["Caption"].font.color.rgb = BLACK
    doc.styles["Caption"].paragraph_format.first_line_indent = Cm(0)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.paragraph_format.first_line_indent = Cm(0)
    black_run(footer.add_run(), size=10)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    text = OxmlElement("w:r")
    value = OxmlElement("w:t")
    value.text = "1"
    text.append(value)
    field.append(text)
    footer._p.append(field)


def force_black(doc):
    for root in [doc._element, doc.styles.element, *[section.footer._element for section in doc.sections]]:
        for color in root.iter(qn("w:color")):
            color.attrib.clear()
            color.set(qn("w:val"), "000000")
    for part in doc.part.package.parts:
        if str(part.partname) != "/word/theme/theme1.xml":
            continue
        theme = parse_xml(part.blob)
        a = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
        for scheme in theme.iter(a + "clrScheme"):
            for entry in scheme:
                label = entry.tag.rsplit("}", 1)[-1]
                if label.startswith("accent") or label in {"hlink", "folHlink", "dk2"}:
                    for child in list(entry):
                        entry.remove(child)
                    value = OxmlElement("a:srgbClr")
                    value.set("val", "000000")
                    entry.append(value)
        from lxml import etree
        part._blob = etree.tostring(theme, xml_declaration=True, encoding="UTF-8", standalone=True)


def build(revision):
    data = json.loads((BASE / "report_content.json").read_text())
    result = json.loads((BASE / "evidence/results.json").read_text())
    doc = Document()
    configure(doc)
    paragraph(doc, data["title"], style="Title", centered=True)
    paragraph(doc, data["subtitle"], style="Subtitle", centered=True)
    paragraph(doc, data["author"], centered=True, size=11)
    for index, page in enumerate(data["pages"]):
        h = heading(doc, page["heading"])
        h.paragraph_format.page_break_before = bool(index)
        for block in page["blocks"]:
            if "p" in block:
                paragraph(doc, block["p"])
            elif "h2" in block:
                heading(doc, block["h2"], level=2)
            elif "flow" in block:
                paragraph(doc, block["flow"], centered=True, size=12)
            elif "table" in block:
                table(doc, block["table"]["headers"], block["table"]["rows"])

    new_page(doc, "PHẦN THỰC HÀNH")
    heading(doc, "7 Thiết lập notebook và môi trường", level=2)
    paragraph(doc, "Thực hành Q-learning trên mê cung 10 × 10 theo ví dụ GeeksforGeeks [4]. Mã được viết lại thành các hàm nhỏ; giữ ma trận, bốn hướng di chuyển, phần thưởng và lịch giảm epsilon của ví dụ. Notebook chạy bằng Python 3.11 trong môi trường WSL hiện có, dùng NumPy trên CPU.")
    screenshot(doc, 1, "Hình 1  Cell khởi tạo và output kiểm tra môi trường")
    paragraph(doc, "Mê cung có 43 ô đi được. Bảng Q có kích thước 10 × 10 × 4, tương ứng 400 giá trị; các phần tử khởi tạo bằng không. Điểm bắt đầu là (0, 0), đích là (9, 9). Mỗi tọa độ viết theo thứ tự hàng rồi cột.", size=12)
    paragraph(doc, "Thiết lập: 5.000 lượt, seed 42, alpha 0,1, gamma 0,9; epsilon từ 0,5 giảm theo hệ số 0,995 và không dưới 0,01. Tới đích nhận +50; va tường hoặc ra ngoài nhận −10; bước thường nhận −1.", size=12)

    new_page(doc, "8 Môi trường và cập nhật kinh nghiệm")
    paragraph(doc, "Hàm valid kiểm tra ô đi được. Hàm transition trả trạng thái tiếp theo, phần thưởng và cờ kết thúc. Va tường hoặc ra ngoài kết thúc lượt; trạng thái được giữ tại ô trước va chạm, như quy ước của ví dụ [4].", size=12)
    screenshot(doc, 2, "Hình 2  Hàm môi trường và ba output kiểm tra")
    paragraph(doc, "choose_action cân bằng thử ngẫu nhiên và chọn giá trị Q lớn nhất. update_q điều chỉnh giá trị của lựa chọn vừa thực hiện. Nếu lượt đã kết thúc, giá trị tương lai bằng không; không cộng thêm giá trị ở trạng thái đích.", size=12)
    screenshot(doc, 3, "Hình 3  Chọn hành động và kiểm tra cập nhật Q ở đích")

    new_page(doc, "9 Huấn luyện tác tử")
    paragraph(doc, "Mỗi lượt bắt đầu lại từ START. Tác tử chọn hành động, nhận phản hồi rồi cập nhật Q. Seed cố định giúp chạy lại có cùng kết quả. Giới hạn 400 bước mỗi lượt là chốt an toàn được bổ sung; lượt đạt giới hạn được ghi là truncated, không bị coi là đã tới đích.", size=12)
    screenshot(doc, 4, "Hình 4  Toàn bộ hàm train và output sau 5.000 lượt")
    paragraph(doc, "Các mốc trong output là thống kê từ lượt thực thi notebook. Không có lượt nào bị cắt bởi giới hạn bước. Epsilon giảm dần tới 0,01, nên kể cả cuối huấn luyện tác tử vẫn thỉnh thoảng chọn hành động khám phá.", size=12)

    new_page(doc, "10 Kiểm tra tác tử sau khi học")
    paragraph(doc, "Đánh giá dưới đây tắt khám phá và luôn chọn argmax của Q. Hàm greedy_path không lọc trước hướng va tường và không tự tìm đường thay tác tử. Nếu chính sách va chạm, lặp hoặc quá số bước, hàm báo đúng trạng thái đó thay vì ép đường đi tới đích.")
    screenshot(doc, 5, "Hình 5  Chính sách greedy và đường đi thực tế")
    table(doc, ["Kết quả đánh giá", "Giá trị thực thi"], [["Trạng thái cuối", result["greedy_status"]], ["Số bước di chuyển", str(result["learned_moves"])], ["Tổng thưởng không chiết khấu", str(result["greedy_return"])]])
    paragraph(doc, "Đường đi có 25 tọa độ, tương ứng 24 lần di chuyển. Có 23 bước thường nhận −1 và bước cuối tới đích nhận +50, nên tổng thưởng là 27. Tổng thưởng này không phải giá trị Q: Q được học với gamma 0,9 nên có xét chiết khấu phần thưởng tương lai.")
    paragraph(doc, "Chỉ nhìn đường vẽ chưa đủ kết luận ngắn nhất. Cần kiểm từng bước không xuyên tường, đúng bốn hướng được phép và kết thúc tại GOAL. Trang sau dùng một thuật toán độc lập để kiểm tra thêm độ dài tối thiểu.")

    new_page(doc, "11 Đối chiếu độc lập bằng BFS")
    paragraph(doc, "BFS duyệt lưới không trọng số theo từng lớp khoảng cách, chỉ được gọi sau khi Q-learning đã học xong. Hàm này không sửa Q, không cung cấp hành động cho train và không được dùng để thay đường đi greedy.", size=12)
    screenshot(doc, 6, "Hình 6  Mã BFS và output đối chiếu đường đi")
    paragraph(doc, "valid_path kiểm tra các ô đều đi được và hai tọa độ liên tiếp kề nhau theo hàng hoặc cột. Output xác nhận đường đi hợp lệ và độ dài Q-learning bằng BFS: đều 24 bước. Vì vậy có thể kết luận đường đi greedy là ngắn nhất trên đúng mê cung cố định này, không khái quát sang mọi mê cung hoặc mọi seed.", size=12)

    new_page(doc, "12 Quan sát đường đi học được")
    paragraph(doc, "Hình dưới là output do notebook tự vẽ từ biến path, không phải hình lấy từ bài hướng dẫn. Ô trắng là đường đi, ô xám là tường; đường đen nối các hành động đã chọn. S và G lần lượt đánh dấu đầu và đích.", size=12)
    screenshot(doc, 7, "Hình 7  Code vẽ và output đường đi sau huấn luyện")
    paragraph(doc, "Trục Row tăng từ trên xuống đúng quy ước của ma trận. Đường đi được vẽ qua tâm các ô hợp lệ và đi qua những đoạn bắt buộc của mê cung. Đây là minh họa trực quan cho kết quả đã được kiểm bằng valid_path và BFS, không phải bằng chứng duy nhất để khẳng định thành công.", size=12)

    new_page(doc, "13 Phần thưởng và tổng hợp output")
    paragraph(doc, "Đường xám ghi tổng thưởng từng lượt; đường đen là trung bình trượt 100 lượt. Đường cong có dao động vì huấn luyện vẫn khám phá. Không diễn giải một đường trung bình tăng thành bảo đảm mọi lượt đều thành công.", size=12)
    screenshot(doc, 8, "Hình 8  Code và output biểu đồ phần thưởng")
    screenshot(doc, 9, "Hình 9  Cell tổng hợp các kết quả thực thi")
    paragraph(doc, "Trong 1.000 lượt cuối, tỷ lệ tới đích là 87,6% và thưởng trung bình 20,951. Hai số này khác kết quả greedy 24 bước, thưởng 27 vì điều kiện chọn hành động khác nhau. Thí nghiệm chỉ chạy một seed trên một mê cung; chưa đánh giá khả năng tổng quát hóa.", size=11.5)

    new_page(doc, "TÀI LIỆU THAM KHẢO")
    sources = json.loads((BASE / "sources/manifest.json").read_text())["sources"]
    for source in sources:
        text = f'[{source["id"]}] {source["publisher"]}. {source["title"]}.'
        if source.get("author"):
            text += " Tác giả " + source["author"] + "."
        if source.get("published"):
            text += " Ngày công bố 25/03/2025."
        paragraph(doc, text, size=12)
        p = paragraph(doc, source["url"], size=11)
        p.paragraph_format.first_line_indent = Cm(0)
        paragraph(doc, "Truy cập ngày 23/10/2026.", size=11)
    heading(doc, "Ghi chú về phần thực hành", level=2)
    paragraph(doc, "Mê cung, hành động, phần thưởng và lịch học được tham khảo từ [4]. Seed, giới hạn bước, kiểm tra terminal, đánh giá greedy nghiêm ngặt và đối chứng BFS được bổ sung cho việc kiểm chứng. Các hình code và output là ảnh chụp trình duyệt từ bản HTML xuất của notebook đã thực thi, với màu chữ được đặt đen; không dùng output mẫu của website.", size=12)
    force_black(doc)
    doc.core_properties.title = data["title"]
    doc.core_properties.subject = "Independent assignment on AI agents and tabular Q learning"
    doc.core_properties.author = "Triệu Tuấn Anh"
    destination = BASE / "qa/candidates" / revision / "DISCOVERING_AI_Agent.docx"
    destination.parent.mkdir(parents=True, exist_ok=True)
    if destination.exists():
        raise FileExistsError("Use a new candidate revision")
    doc.save(destination)
    print(json.dumps({"candidate": str(destination), "theory_planned_pages": 6, "screenshots": 9}, ensure_ascii=False))


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--revision", required=True)
    build(parser.parse_args().revision)
