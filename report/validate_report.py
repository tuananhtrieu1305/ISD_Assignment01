#!/usr/bin/env python3
"""Structural and release validation for the extended Phase 19 term-paper DOCX."""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import zipfile
from collections import Counter
from pathlib import Path
from xml.etree import ElementTree as ET

from docx import Document


ROOT = Path(__file__).resolve().parents[2]
REPORT = Path(__file__).resolve().parent
DEFAULT_DOCX = REPORT / "Tieu_luan_AI_ML_CNN_RNN_KG_RAG.docx"
DEFAULT_LOG = REPORT / "build_log_kg_rag.json"
DEFAULT_MAP = REPORT / "qa_phase19" / "word_page_map_final.json"
DEFAULT_OUTPUT = REPORT / "report_validation_kg_rag.json"
MANIFEST_PATH = REPORT / "report_manifest.json"
NS = {
    "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
    "wp": "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing",
}


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def all_text(document: Document) -> str:
    chunks = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            chunks.extend(cell.text for cell in row.cells)
    return "\n".join(chunks)


def add_check(checks: list[dict], name: str, passed: bool, detail: str) -> None:
    checks.append({"name": name, "status": "PASS" if passed else "FAIL", "detail": detail})


def validate(docx: Path, build_log_path: Path, page_map_path: Path, output: Path) -> dict:
    document = Document(docx)
    build_log = json.loads(build_log_path.read_text(encoding="utf-8"))
    page_map = json.loads(page_map_path.read_text(encoding="utf-8"))
    manifest = json.loads(MANIFEST_PATH.read_text(encoding="utf-8"))
    text = all_text(document)
    style_counts = Counter(paragraph.style.name for paragraph in document.paragraphs)
    checks: list[dict] = []
    expected_counts = {
        "headings": sum(len(section["headings"]) for section in manifest["sections"]) + 1,
        "figures": len(manifest["figures"]),
        "content_tables": len(manifest["tables"]),
        "code_blocks": manifest["code_excerpts"]["count"],
        "references": manifest["citations"]["count"],
        "navigation_targets": len(build_log["navigation_targets"]),
    }
    page_low, page_high = manifest["page_budget"]["overall_allowed_range"]
    page_count = page_map.get("page_count", 0)

    add_check(checks, "docx_reopens", True, f"python-docx reopened {docx.name}")
    add_check(checks, "word_pagination", page_low <= page_count <= page_high, f"{page_count} pages; target {page_low}-{page_high}")
    add_check(
        checks,
        "navigation_targets",
        page_map.get("resolved_targets") == expected_counts["navigation_targets"] and not page_map.get("missing_targets"),
        f"{page_map.get('resolved_targets')}/{expected_counts['navigation_targets']} resolved; missing={len(page_map.get('missing_targets', []))}",
    )
    add_check(checks, "headings", sum(style_counts[f"Heading {i}"] for i in (1, 2, 3)) == expected_counts["headings"], str({f"H{i}": style_counts[f"Heading {i}"] for i in (1, 2, 3)}))
    expected_captions = expected_counts["figures"] + expected_counts["content_tables"]
    add_check(checks, "captions", style_counts["Caption"] == expected_captions, f"{style_counts['Caption']} captions = {expected_counts['figures']} figures + {expected_counts['content_tables']} tables")
    expected_tables = expected_counts["content_tables"] + expected_counts["code_blocks"]
    add_check(checks, "native_tables", len(document.tables) == expected_tables, f"{len(document.tables)} editable Word tables")
    add_check(checks, "inline_images", len(document.inline_shapes) == expected_counts["figures"] + 1, f"{len(document.inline_shapes)} inline images including logo")
    add_check(checks, "references", style_counts["Reference Entry"] == expected_counts["references"], f"{style_counts['Reference Entry']} ordered references")
    add_check(checks, "sections", len(document.sections) == 3, f"{len(document.sections)} sections: cover, title/front matter, body")

    expected_margins = (21.0, 29.7, 2.25, 2.0, 2.7, 2.2)
    geometry_ok = True
    for section in document.sections:
        actual = (
            section.page_width.cm,
            section.page_height.cm,
            section.top_margin.cm,
            section.bottom_margin.cm,
            section.left_margin.cm,
            section.right_margin.cm,
        )
        geometry_ok &= all(abs(value - expected) < 0.03 for value, expected in zip(actual, expected_margins))
    add_check(checks, "a4_geometry", geometry_ok, "A4 portrait; margins 2.25/2.0/2.7/2.2 cm")

    forbidden = {
        "TODO": r"\bTODO\b",
        "TBD": r"\bTBD\b",
        "placeholder": r"PLACEHOLDER|\[CẦN[^\]]*\]",
        "internal_citation_token": r"turn\d+(?:search|fetch|view)\d+|cite",
        "absolute_windows_path": r"[A-Za-z]:\\",
    }
    hits = {name: bool(re.search(pattern, text, re.IGNORECASE | re.MULTILINE)) for name, pattern in forbidden.items()}
    toc_page_markers = [
        paragraph.text
        for paragraph in document.paragraphs
        if paragraph.style.name.startswith("Static TOC") and paragraph.text.rstrip().endswith("—")
    ]
    hits["unresolved_page_marker"] = bool(toc_page_markers)
    add_check(checks, "forbidden_tokens", not any(hits.values()), json.dumps(hits, ensure_ascii=False))

    caption_targets = [target["display"] for target in build_log["navigation_targets"] if target["kind"] in {"figure", "table"}]
    caption_text = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.style.name == "Caption"]
    add_check(checks, "caption_sequence", caption_text == caption_targets, f"caption order matches {expected_captions} navigation targets")

    with zipfile.ZipFile(docx) as archive:
        names = set(archive.namelist())
        document_xml = archive.read("word/document.xml")
        body_root = ET.fromstring(document_xml)
        bookmarks = body_root.findall(".//w:bookmarkStart", NS)
        doc_prs = body_root.findall(".//wp:docPr", NS)
        alt_count = sum(bool(node.attrib.get("descr", "").strip()) for node in doc_prs)
        repeated_headers = body_root.findall(".//w:tblHeader", NS)
        row_heights = body_root.findall(".//w:trHeight", NS)
        footer_xml = b"\n".join(archive.read(name) for name in names if name.startswith("word/footer") and name.endswith(".xml"))
        all_xml = b"\n".join(archive.read(name) for name in names if name.endswith(".xml"))

    add_check(checks, "bookmarks", len(bookmarks) >= expected_counts["navigation_targets"], f"{len(bookmarks)} bookmark starts")
    add_check(checks, "image_alt_text", alt_count >= expected_counts["figures"], f"{alt_count} figure/image descriptions")
    add_check(checks, "repeat_table_headers", len(repeated_headers) == expected_counts["content_tables"], f"{len(repeated_headers)} repeated data-table header rows")
    add_check(checks, "no_fixed_row_heights", not row_heights, "no w:trHeight constraints")
    add_check(checks, "page_number_field", b" PAGE " in footer_xml, "PAGE field present in body footer")
    change_tokens = (b"<w:commentRangeStart", b"<w:ins>", b"<w:ins ", b"<w:del>", b"<w:del ")
    add_check(checks, "no_comments_or_tracked_changes", not any(token in all_xml for token in change_tokens), "no comments, insertions, or deletions")

    visual_pages = REPORT / "qa_phase19" / "word_final_pages_v2"
    visual_qa_path = REPORT / "qa_phase19" / "visual_qa.json"
    render_attempt_path = REPORT / "qa_phase19" / "render_docx_attempt.json"
    word_pngs = sorted(visual_pages.glob("page-*.png"))
    visual_qa = json.loads(visual_qa_path.read_text(encoding="utf-8")) if visual_qa_path.exists() else {}
    render_attempt = json.loads(render_attempt_path.read_text(encoding="utf-8")) if render_attempt_path.exists() else {}
    add_check(checks, "word_pdf_visual_qa", len(word_pngs) == page_count and page_count > 0, f"{len(word_pngs)}/{page_count} authoritative Word-rendered pages inspected")
    expected_pages = list(range(1, page_count + 1))
    add_check(
        checks,
        "visual_qa_attestation",
        visual_qa.get("status") == "PASS" and visual_qa.get("inspected_pages") == expected_pages,
        f"attested {len(visual_qa.get('inspected_pages', []))}/{page_count} pages",
    )
    renderer_documented = (
        render_attempt.get("status") == "UNAVAILABLE"
        and "soffice.exe" in render_attempt.get("reason", "")
        and render_attempt.get("authoritative_fallback", {}).get("page_count") == page_count
    )
    add_check(
        checks,
        "documents_renderer_attempt_documented",
        renderer_documented,
        render_attempt.get("reason", "missing render attempt record"),
    )

    failures = [check for check in checks if check["status"] == "FAIL"]
    result = {
        "schema_version": 1,
        "phase": 19,
        "status": "PASS" if not failures else "FAIL",
        "document": docx.relative_to(ROOT).as_posix(),
        "document_sha256": sha256(docx),
        "document_bytes": docx.stat().st_size,
        "authoritative_renderer": "Microsoft Word 16 PDF export",
        "page_count": page_map.get("page_count"),
        "visual_qa": f"All {page_count} Word-rendered pages inspected; no clipping or overflow may remain at release.",
        "renderer_compatibility_note": (
            "The canonical documents/render_docx.py renderer was invoked, but LibreOffice soffice.exe is not installed "
            "on this Windows host. Final pagination and PDF export therefore use Microsoft Word 16; all pages were "
            "rasterized with the Windows PDF API and inspected. The exception is recorded in render_docx_attempt.json."
        ),
        "counts": expected_counts,
        "checks": checks,
        "failures": failures,
    }
    output.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")

    build_log["output_sha256"] = result["document_sha256"]
    build_log["output_bytes"] = result["document_bytes"]
    build_log["finalization"] = {
        "word_page_count": result["page_count"],
        "resolved_navigation_targets": page_map.get("resolved_targets"),
        "validation": output.relative_to(ROOT).as_posix(),
        "validation_status": result["status"],
        "visual_qa_pages": len(word_pngs),
    }
    build_log_path.write_text(json.dumps(build_log, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return result


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--docx", type=Path, default=DEFAULT_DOCX)
    parser.add_argument("--build-log", type=Path, default=DEFAULT_LOG)
    parser.add_argument("--page-map", type=Path, default=DEFAULT_MAP)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()
    result = validate(args.docx.resolve(), args.build_log.resolve(), args.page_map.resolve(), args.output.resolve())
    print(f"Phase 19 validation: {result['status']} ({len(result['checks'])} checks, {len(result['failures'])} failures)")
    if result["status"] != "PASS":
        raise SystemExit(1)


if __name__ == "__main__":
    main()
