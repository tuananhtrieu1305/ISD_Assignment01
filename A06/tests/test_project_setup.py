"""Contract tests for the Assignment 06 initialization milestone."""

from __future__ import annotations

from pathlib import Path

import openpyxl
import pandas as pd


PROJECT_ROOT = Path(__file__).resolve().parents[1]
CUSTOMER_PATH = PROJECT_ROOT / "datasets" / "customer" / "online_retail_II.xlsx"
STOCK_PATH = PROJECT_ROOT / "datasets" / "stock" / "AAPL_2015_2025.csv"
RAW_RETAIL_COLUMNS = [
    "Invoice",
    "StockCode",
    "Description",
    "Quantity",
    "InvoiceDate",
    "Price",
    "Customer ID",
    "Country",
]
STOCK_COLUMNS = ["Date", "Open", "High", "Low", "Close", "Adj Close", "Volume"]


def test_required_project_structure_exists() -> None:
    required = [
        "datasets/customer",
        "datasets/stock",
        "datasets/processed",
        "notebooks",
        "src",
        "tests",
        "models/pytorch",
        "models/keras",
        "results/figures",
        "results/metrics",
        "results/predictions",
        "report",
    ]
    assert all((PROJECT_ROOT / relative).is_dir() for relative in required)
    assert (PROJECT_ROOT / "PROJECT_SPEC.md").is_file()
    assert (PROJECT_ROOT / "CURRENT_STATE.md").is_file()
    assert (PROJECT_ROOT / "README.md").is_file()


def test_customer_workbook_has_locked_raw_schema() -> None:
    workbook = openpyxl.load_workbook(CUSTOMER_PATH, read_only=True, data_only=True)
    try:
        assert workbook.sheetnames == ["Year 2009-2010", "Year 2010-2011"]
        for worksheet in workbook.worksheets:
            header = [cell.value for cell in next(worksheet.iter_rows(max_row=1))]
            assert header == RAW_RETAIL_COLUMNS
    finally:
        workbook.close()


def test_stock_csv_has_locked_schema_and_interval() -> None:
    stock = pd.read_csv(STOCK_PATH)
    dates = pd.to_datetime(stock["Date"], errors="raise")
    assert stock.columns.tolist() == STOCK_COLUMNS
    assert len(stock) == 2766
    assert dates.min() >= pd.Timestamp("2015-01-01")
    assert dates.max() <= pd.Timestamp("2025-12-31")
    assert dates.is_monotonic_increasing
    assert not dates.duplicated().any()
    assert not stock.isna().any().any()


def test_protocol_locks_core_experiment_decisions() -> None:
    spec = (PROJECT_ROOT / "PROJECT_SPEC.md").read_text(encoding="utf-8")
    required_terms = [
        "Sequence length: 8 weeks",
        "Sequence length: 30 trading days",
        "approximately 70% train, 15% validation, and 15% test",
        "Vanilla RNN",
        "SEED = 42",
        "CustomerID",
        "auto_adjust=False",
    ]
    assert all(term in spec for term in required_terms)
