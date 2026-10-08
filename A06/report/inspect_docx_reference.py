"""Read-only style inventory for the A05 formatting reference."""

from __future__ import annotations

import json
import sys
from collections import Counter
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


def emu_cm(value):
    return None if value is None else round(value.cm, 3)


def pt(value):
    return None if value is None else round(value.pt, 3)


def style_info(style):
    pf = style.paragraph_format
    rf = style.font
    rfonts = style._element.rPr.rFonts if style._element.rPr is not None else None
    return {
        "name": style.name,
        "font_name": rf.name,
        "font_ascii": None if rfonts is None else rfonts.get(qn("w:ascii")),
        "font_hAnsi": None if rfonts is None else rfonts.get(qn("w:hAnsi")),
        "font_size_pt": pt(rf.size),
        "bold": rf.bold,
        "italic": rf.italic,
        "color": None if rf.color is None or rf.color.rgb is None else str(rf.color.rgb),
        "alignment": None if pf.alignment is None else str(pf.alignment),
        "space_before_pt": pt(pf.space_before),
        "space_after_pt": pt(pf.space_after),
        "line_spacing": pf.line_spacing,
        "first_line_indent_cm": emu_cm(pf.first_line_indent),
        "left_indent_cm": emu_cm(pf.left_indent),
    }


def paragraph_info(p, index):
    pf = p.paragraph_format
    runs = []
    for r in p.runs:
        if not r.text:
            continue
        runs.append({
            "text": r.text[:120],
            "font": r.font.name,
            "size_pt": pt(r.font.size),
            "bold": r.bold,
            "italic": r.italic,
            "color": None if r.font.color is None or r.font.color.rgb is None else str(r.font.color.rgb),
        })
    return {
        "index": index,
        "text": p.text[:300],
        "style": p.style.name,
        "alignment": None if pf.alignment is None else str(pf.alignment),
        "space_before_pt": pt(pf.space_before),
        "space_after_pt": pt(pf.space_after),
        "line_spacing": pf.line_spacing,
        "first_line_indent_cm": emu_cm(pf.first_line_indent),
        "left_indent_cm": emu_cm(pf.left_indent),
        "runs": runs,
    }


path = Path(sys.argv[1])
doc = Document(path)
section = doc.sections[0]
interesting = {"Normal", "normal", "Title", "Heading 1", "Heading 2", "Heading 3", "Caption"}
styles = [style_info(s) for s in doc.styles if s.name in interesting]

selected = []
for i, p in enumerate(doc.paragraphs):
    text = p.text.strip()
    if i < 50 or text.startswith(("I.", "1.1.", "Bảng 1.", "Hình 1.", "V. ", "X. ")):
        selected.append(paragraph_info(p, i))

table_styles = Counter(t.style.name if t.style else "(none)" for t in doc.tables)
out = {
    "document": str(path),
    "paragraphs": len(doc.paragraphs),
    "tables": len(doc.tables),
    "inline_shapes": len(doc.inline_shapes),
    "sections": len(doc.sections),
    "section": {
        "page_width_cm": emu_cm(section.page_width),
        "page_height_cm": emu_cm(section.page_height),
        "top_margin_cm": emu_cm(section.top_margin),
        "bottom_margin_cm": emu_cm(section.bottom_margin),
        "left_margin_cm": emu_cm(section.left_margin),
        "right_margin_cm": emu_cm(section.right_margin),
        "header_distance_cm": emu_cm(section.header_distance),
        "footer_distance_cm": emu_cm(section.footer_distance),
        "different_first_page": section.different_first_page_header_footer,
    },
    "styles": styles,
    "table_styles": table_styles,
    "selected_paragraphs": selected,
}
print(json.dumps(out, ensure_ascii=False, indent=2, default=str))
