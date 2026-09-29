"""Build the final Vietnamese academic DOCX report for Assignment 06.

The report is generated only from the audited A06 artifacts.  This script does
not train models, rebuild preprocessing, or modify datasets/model results.
"""

from __future__ import annotations

import json
import math
import re
import sys
import textwrap
import zipfile
from pathlib import Path
from typing import Iterable, Sequence

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import (
    WD_ALIGN_PARAGRAPH,
    WD_BREAK,
    WD_LINE_SPACING,
    WD_TAB_ALIGNMENT,
    WD_TAB_LEADER,
)
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Cm, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


SCRIPT_PATH = Path(__file__).resolve()
REPORT_DIR = SCRIPT_PATH.parent
ROOT = REPORT_DIR.parent
OUTPUT_PATH = REPORT_DIR / "A06_Assignment_Report.docx"
ASSETS_DIR = REPORT_DIR / "docx_assets"

ACCENT = "17365D"
ACCENT_2 = "2F75B5"
LIGHT_BLUE = "D9EAF7"
PALE_BLUE = "F3F8FC"
PALE_GOLD = "FFF4D6"
PALE_GREEN = "EAF4EA"
PALE_RED = "FCE8E6"
GRID = "9EADBA"
TEXT = "1F2933"
MUTED = "5B6770"

FIGURES: list[tuple[str, str]] = [
    ("2.1", "RNN được unroll qua ba bước thời gian với tham số dùng chung"),
    ("3.1", "Hoạt động mua hàng theo tháng trong Online Retail II"),
    ("3.2", "Phân phối chi tiêu, tần suất mua và số đơn của khách hàng"),
    ("3.3", "Minh họa tuần không hoạt động được điền bằng 0"),
    ("3.4", "Phân bố nhãn mua hàng tuần kế tiếp ở bước EDA nguyên mẫu"),
    ("3.5", "Diễn biến Close và Volume của AAPL trong snapshot đã khóa"),
    ("3.6", "Close của AAPL cùng MA(7) và MA(30) dùng cho EDA"),
    ("3.7", "Phân phối tỷ suất sinh lợi ngày của AAPL"),
    ("4.1", "Sơ đồ chia tập theo thời gian và nguyên tắc chống leakage"),
    ("5.1", "Loss huấn luyện và validation của PyTorch Customer RNN"),
    ("5.2", "Confusion matrix trên TEST của PyTorch Customer RNN"),
    ("5.3", "Đường ROC trên TEST của PyTorch Customer RNN"),
    ("5.4", "Loss huấn luyện và validation của PyTorch Stock RNN"),
    ("5.5", "Close thực tế và dự báo trên TEST của PyTorch Stock RNN"),
    ("5.6", "So sánh PyTorch Stock RNN với naive baseline"),
    ("6.1", "Loss huấn luyện và validation của Keras Customer SimpleRNN"),
    ("6.2", "Confusion matrix trên TEST của Keras Customer SimpleRNN"),
    ("6.3", "Đường ROC trên TEST của Keras Customer SimpleRNN"),
    ("6.4", "Loss huấn luyện và validation của Keras Stock SimpleRNN"),
    ("6.5", "Close thực tế và dự báo trên TEST của Keras Stock SimpleRNN"),
    ("6.6", "So sánh Keras Stock SimpleRNN với naive baseline"),
    ("7.1", "So sánh metric customer giữa PyTorch và Keras"),
    ("7.2", "So sánh confusion matrix customer giữa hai framework"),
    ("7.3", "So sánh ROC customer giữa hai framework"),
    ("7.4", "So sánh metric stock giữa hai RNN và naive baseline"),
    ("7.5", "So sánh dự báo stock theo thời gian trên cùng TEST"),
]

TABLES: list[tuple[str, str]] = [
    ("1.1", "Phạm vi dữ liệu và nhiệm vụ học máy của A06"),
    ("2.1", "Ký hiệu trong một Vanilla RNN cell"),
    ("2.2", "Các kiểu ánh xạ input/output của mô hình chuỗi"),
    ("2.3", "So sánh Vanilla RNN, LSTM và GRU"),
    ("3.1", "Cấu trúc workbook Online Retail II"),
    ("3.2", "Kết quả kiểm tra và làm sạch giao dịch customer"),
    ("3.3", "Định nghĩa năm đặc trưng customer-week"),
    ("3.4", "Ý nghĩa các trường giá AAPL"),
    ("3.5", "Kết quả kiểm tra dữ liệu AAPL"),
    ("4.1", "Shape và khoảng thời gian của các split sản xuất"),
    ("4.2", "Các kiểm soát chống data leakage đã được audit"),
    ("5.1", "Thiết lập PyTorch Customer RNN"),
    ("5.2", "Kết quả TEST của PyTorch Customer RNN"),
    ("5.3", "Thiết lập và kết quả PyTorch Stock RNN"),
    ("6.1", "Thiết lập Keras Customer SimpleRNN"),
    ("6.2", "Kết quả TEST của Keras Customer SimpleRNN"),
    ("6.3", "Thiết lập và kết quả Keras Stock SimpleRNN"),
    ("7.1", "So sánh kết quả customer trên cùng TEST"),
    ("7.2", "So sánh kết quả stock trên cùng TEST"),
    ("7.3", "Khác biệt thực hành giữa PyTorch và Keras"),
]

# Filled after the first render-and-review pass.  The builder validates that no
# dash/placeholder remains in the final front matter.
TOC_PAGES = {
    "I. TỔNG QUAN": "5",
    "II. CƠ SỞ LÝ THUYẾT VỀ RNN": "6",
    "III. PHÂN TÍCH DỮ LIỆU": "9",
    "3.1. UCI Online Retail II": "9",
    "3.2. AAPL lịch sử": "12",
    "IV. TIỀN XỬ LÝ VÀ THIẾT KẾ THÍ NGHIỆM": "15",
    "V. THỰC NGHIỆM PYTORCH": "17",
    "5.1. Customer classification": "17",
    "5.2. Stock regression": "20",
    "VI. THỰC NGHIỆM KERAS": "22",
    "6.1. Customer classification": "22",
    "6.2. Stock regression": "25",
    "VII. SO SÁNH HAI FRAMEWORK": "28",
    "VIII. HẠN CHẾ VÀ HƯỚNG PHÁT TRIỂN": "31",
    "IX. KẾT LUẬN": "32",
    "TÀI LIỆU THAM KHẢO": "33",
}

FIGURE_PAGES = {
    "2.1": "6", "3.1": "10", "3.2": "10", "3.3": "11", "3.4": "11",
    "3.5": "13", "3.6": "13", "3.7": "14", "4.1": "15", "5.1": "18",
    "5.2": "18", "5.3": "19", "5.4": "20", "5.5": "21", "5.6": "21",
    "6.1": "23", "6.2": "23", "6.3": "24", "6.4": "25", "6.5": "26",
    "6.6": "27", "7.1": "28", "7.2": "29", "7.3": "29", "7.4": "30",
    "7.5": "30",
}

TABLE_PAGES = {
    "1.1": "5", "2.1": "6", "2.2": "7", "2.3": "8", "3.1": "9",
    "3.2": "9", "3.3": "11", "3.4": "12", "3.5": "12", "4.1": "15",
    "4.2": "16", "5.1": "17", "5.2": "17", "5.3": "20", "6.1": "22",
    "6.2": "22", "6.3": "25", "7.1": "28", "7.2": "29", "7.3": "30",
}


def read_json(relative_path: str) -> dict:
    path = ROOT / relative_path
    if not path.is_file():
        raise FileNotFoundError(path)
    return json.loads(path.read_text(encoding="utf-8"))


def ensure_sources() -> None:
    required = [
        "PROJECT_SPEC.md",
        "CURRENT_STATE.md",
        "report/A06_AUDIT_REPORT.md",
        "results/metrics/preprocessing_metadata.json",
        "results/metrics/pytorch_customer_metrics.json",
        "results/metrics/keras_customer_metrics.json",
        "results/metrics/pytorch_stock_metrics.json",
        "results/metrics/keras_stock_metrics.json",
        "results/metrics/framework_comparison.json",
        "results/metrics/environment_verification.json",
    ]
    required.extend(f"notebooks/{i:02d}_{name}.ipynb" for i, name in [
        (1, "RNN_Fundamentals"), (2, "Customer_Behavior_Data"), (3, "Stock_Data"),
        (4, "PyTorch_Customer_RNN"), (5, "PyTorch_Stock_RNN"),
        (6, "Keras_Customer_RNN"), (7, "Keras_Stock_RNN"),
        (8, "Framework_Comparison"),
    ])
    for path in required:
        if not (ROOT / path).is_file():
            raise FileNotFoundError(f"Thiếu source bắt buộc: {path}")

    audit = (ROOT / "report/A06_AUDIT_REPORT.md").read_text(encoding="utf-8")
    if not re.search(r"\*\*PASSED\*\*\s*$", audit):
        raise RuntimeError("Audit gate chưa đạt PASSED; không tạo báo cáo cuối.")


def validate_metric_contract(pre: dict, ptc: dict, kc: dict, pts: dict, ks: dict, comp: dict) -> None:
    assert pre["seed"] == ptc["seed"] == kc["seed"] == pts["seed"] == ks["seed"] == 42
    assert ptc["data"]["split_shapes"] == kc["data"]["split_shapes"]
    assert math.isclose(pts["naive_MAE"], ks["naive_MAE"], abs_tol=1e-12)
    assert math.isclose(pts["naive_RMSE"], ks["naive_RMSE"], abs_tol=1e-12)
    assert math.isclose(pts["naive_R2"], ks["naive_R2"], abs_tol=1e-12)
    assert comp["customer"]["same_test_keys_and_labels_verified"] is True
    assert comp["stock"]["same_dates_actual_and_naive_verified"] is True
    assert comp["stock"]["rnn_beats_naive_baseline"] is False


def set_run_font(run, name: str = "Times New Roman", size: float | None = None,
                 bold: bool | None = None, italic: bool | None = None,
                 color: str | None = None) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top: int = 90, start: int = 100, bottom: int = 90, end: int = 100) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def remove_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "nil")


def set_fixed_cell_width(cell, width_cm: float) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width_cm * 567)))
    tc_w.set(qn("w:type"), "dxa")


def set_table_layout(table, widths_cm: Sequence[float]) -> None:
    """Set table, grid, and cell widths explicitly for deterministic renderers."""
    table.autofit = False
    total_twips = sum(int(width * 567) for width in widths_cm)
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total_twips))
    tbl_w.set(qn("w:type"), "dxa")
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_cm:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(int(width * 567)))
        grid.append(grid_col)

    for row in table.rows:
        for idx, width in enumerate(widths_cm):
            if idx < len(row.cells):
                set_fixed_cell_width(row.cells[idx], width)


def add_page_field(paragraph) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    run._r.extend([begin, instr, separate])
    result_run = paragraph.add_run("1")
    set_run_font(result_run, size=8.5, color=MUTED)
    end_run = paragraph.add_run()
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    end_run._r.append(end)


def add_hyperlink(paragraph, text: str, url: str) -> None:
    rel_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), ACCENT_2)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_fonts = OxmlElement("w:rFonts")
    r_fonts.set(qn("w:ascii"), "Times New Roman")
    r_fonts.set(qn("w:hAnsi"), "Times New Roman")
    r_pr.extend([r_fonts, color, underline])
    run.append(r_pr)
    t = OxmlElement("w:t")
    t.text = text
    run.append(t)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def set_alt_text(inline_shape, title: str, description: str) -> None:
    doc_pr = inline_shape._inline.docPr
    doc_pr.set("title", title)
    doc_pr.set("descr", description)


def set_paragraph_box(paragraph, fill: str, border_color: str = "CAD6E0") -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    shd = p_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        p_pr.append(shd)
    shd.set(qn("w:fill"), fill)
    borders = p_pr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        p_pr.append(borders)
    for edge_name in ("top", "left", "bottom", "right"):
        edge = borders.find(qn(f"w:{edge_name}"))
        if edge is None:
            edge = OxmlElement(f"w:{edge_name}")
            borders.append(edge)
        edge.set(qn("w:val"), "single")
        edge.set(qn("w:sz"), "5")
        edge.set(qn("w:space"), "4")
        edge.set(qn("w:color"), border_color)


def _wrap_pixels(text: str, font: ImageFont.FreeTypeFont, max_width: int) -> list[str]:
    lines: list[str] = []
    for raw_line in str(text).splitlines() or [""]:
        words = raw_line.split()
        if not words:
            lines.append("")
            continue
        current = words[0]
        for word in words[1:]:
            candidate = f"{current} {word}"
            if font.getlength(candidate) <= max_width:
                current = candidate
            else:
                lines.append(current)
                current = word
        lines.append(current)
    return lines or [""]


def render_table_image(number: str, headers: Sequence[str], rows: Sequence[Sequence[str]],
                       widths: Sequence[float], font_size: float) -> Path:
    """Rasterize a high-resolution academic table for deterministic DOCX rendering."""
    ASSETS_DIR.mkdir(parents=True, exist_ok=True)
    path = ASSETS_DIR / f"table_{number.replace('.', '_')}.png"
    canvas_width = 2200
    outer = 12
    usable = canvas_width - 2 * outer
    total_width = float(sum(widths))
    col_widths = [int(usable * float(w) / total_width) for w in widths]
    col_widths[-1] += usable - sum(col_widths)
    body_px = max(25, int(font_size * 3.45))
    header_px = body_px + 2
    font_path = "C:/Windows/Fonts/arial.ttf"
    bold_path = "C:/Windows/Fonts/arialbd.ttf"
    body_font = ImageFont.truetype(font_path, body_px)
    header_font = ImageFont.truetype(bold_path, header_px)
    line_gap = 7
    pad_x, pad_y = 14, 13

    wrapped_header = [_wrap_pixels(str(value), header_font, col_widths[i] - 2 * pad_x)
                      for i, value in enumerate(headers)]
    wrapped_rows = [
        [_wrap_pixels(str(value), body_font, col_widths[i] - 2 * pad_x)
         for i, value in enumerate(row)] for row in rows
    ]

    def row_height(cells: Sequence[Sequence[str]], font: ImageFont.FreeTypeFont) -> int:
        bbox = font.getbbox("Ag")
        line_h = bbox[3] - bbox[1] + line_gap
        return max(58, max(len(lines) for lines in cells) * line_h + 2 * pad_y)

    heights = [row_height(wrapped_header, header_font)]
    heights.extend(row_height(row, body_font) for row in wrapped_rows)
    canvas_height = sum(heights) + 2 * outer
    image = Image.new("RGB", (canvas_width, canvas_height), "white")
    draw = ImageDraw.Draw(image)
    y = outer
    all_rows = [wrapped_header, *wrapped_rows]
    for row_idx, (cells, height) in enumerate(zip(all_rows, heights)):
        x = outer
        fill = "#D9EAF7" if row_idx == 0 else ("#F3F8FC" if row_idx % 2 == 0 else "#FFFFFF")
        font = header_font if row_idx == 0 else body_font
        color = "#17365D" if row_idx == 0 else "#1F2933"
        for col_idx, (lines, width) in enumerate(zip(cells, col_widths)):
            draw.rectangle([x, y, x + width, y + height], fill=fill, outline="#7F8F9C", width=2)
            line_h = font.getbbox("Ag")[3] - font.getbbox("Ag")[1] + line_gap
            block_h = len(lines) * line_h - line_gap
            text_y = y + max(pad_y, (height - block_h) // 2)
            for line in lines:
                if row_idx == 0 or col_idx > 0:
                    text_x = x + (width - font.getlength(line)) / 2
                else:
                    text_x = x + pad_x
                draw.text((text_x, text_y), line, fill=color, font=font)
                text_y += line_h
            x += width
        y += height
    image.save(path, dpi=(300, 300))
    return path


def render_diagram_image(number: str, lines: Sequence[str]) -> Path:
    ASSETS_DIR.mkdir(parents=True, exist_ok=True)
    path = ASSETS_DIR / f"figure_{number.replace('.', '_')}.png"
    width = 2100
    row_h = 112
    height = len(lines) * row_h + 24
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    regular = ImageFont.truetype("C:/Windows/Fonts/consola.ttf", 34)
    bold = ImageFont.truetype("C:/Windows/Fonts/consolab.ttf", 36)
    y = 12
    for idx, line in enumerate(lines):
        fill = "#D9EAF7" if idx == 0 else ("#F3F8FC" if idx % 2 else "#FFFFFF")
        draw.rounded_rectangle([12, y, width - 12, y + row_h - 4], radius=12,
                               fill=fill, outline="#7F8F9C", width=2)
        font = bold if idx == 0 else regular
        x = (width - font.getlength(line)) / 2
        bbox = font.getbbox(line)
        text_y = y + (row_h - (bbox[3] - bbox[1])) / 2 - bbox[1] - 2
        draw.text((x, text_y), line, font=font, fill="#17365D")
        y += row_h
    image.save(path, dpi=(300, 300))
    return path


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.25)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.7)
    section.right_margin = Cm(2.2)
    section.header_distance = Cm(0.9)
    section.footer_distance = Cm(0.8)
    section.different_first_page_header_footer = True

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(11.5)
    normal.font.color.rgb = RGBColor.from_string(TEXT)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.widow_control = True

    for style_name, size, color, before, after in [
        ("Title", 21, ACCENT, 0, 12),
        ("Heading 1", 15, ACCENT, 10, 7),
        ("Heading 2", 13, ACCENT_2, 8, 5),
        ("Heading 3", 11.5, ACCENT, 6, 3),
    ]:
        style = doc.styles[style_name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    if "FrontMatterTitle" not in [s.name for s in doc.styles]:
        style = doc.styles.add_style("FrontMatterTitle", WD_STYLE_TYPE.PARAGRAPH)
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        style.font.size = Pt(16)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(ACCENT)
        style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        style.paragraph_format.space_after = Pt(12)

    if "Caption" in [s.name for s in doc.styles]:
        cap = doc.styles["Caption"]
    else:
        cap = doc.styles.add_style("Caption", WD_STYLE_TYPE.PARAGRAPH)
    cap.font.name = "Times New Roman"
    cap._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    cap._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    cap._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    cap.font.size = Pt(10)
    cap.font.italic = True
    cap.font.color.rgb = RGBColor.from_string(MUTED)
    cap.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_before = Pt(3)
    cap.paragraph_format.space_after = Pt(8)
    cap.paragraph_format.keep_together = True

    code = doc.styles.add_style("CodeBlock", WD_STYLE_TYPE.PARAGRAPH)
    code.font.name = "Consolas"
    code._element.rPr.rFonts.set(qn("w:ascii"), "Consolas")
    code._element.rPr.rFonts.set(qn("w:hAnsi"), "Consolas")
    code.font.size = Pt(8.5)
    code.paragraph_format.left_indent = Cm(0.25)
    code.paragraph_format.right_indent = Cm(0.25)
    code.paragraph_format.space_after = Pt(0)
    code.paragraph_format.line_spacing = 1.0

    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = hp.add_run("ASSIGNMENT 06  •  RECURRENT NEURAL NETWORK")
    set_run_font(run, size=8.5, bold=True, color=MUTED)
    p_pr = hp._p.get_or_add_pPr()
    bottom = OxmlElement("w:pBdr")
    edge = OxmlElement("w:bottom")
    edge.set(qn("w:val"), "single")
    edge.set(qn("w:sz"), "4")
    edge.set(qn("w:space"), "1")
    edge.set(qn("w:color"), "CAD6E0")
    bottom.append(edge)
    p_pr.append(bottom)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = fp.add_run("A06 • Báo cáo RNN với PyTorch và Keras")
    set_run_font(fr, size=8.5, color=MUTED)

    # Tell Word to refresh fields when opened; the visible report is static and
    # does not depend on this setting, but the footer PAGE field benefits from it.
    settings = doc.settings._element
    update = OxmlElement("w:updateFields")
    update.set(qn("w:val"), "true")
    settings.append(update)


def add_page_break(doc: Document) -> None:
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def add_body(doc: Document, text: str, *, bold_prefix: str | None = None) -> None:
    p = doc.add_paragraph()
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        set_run_font(r1, bold=True)
        r2 = p.add_run(text[len(bold_prefix):])
        set_run_font(r2)
    else:
        r = p.add_run(text)
        set_run_font(r)


def add_bullets(doc: Document, items: Iterable[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Cm(0.65)
        p.paragraph_format.first_line_indent = Cm(-0.25)
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(item)
        set_run_font(r, size=11.2)


def add_numbered(doc: Document, items: Iterable[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.left_indent = Cm(0.65)
        p.paragraph_format.first_line_indent = Cm(-0.25)
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(item)
        set_run_font(r, size=11.2)


def add_equation(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(7)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run(text)
    set_run_font(r, name="Cambria Math", size=12.5, italic=True, color=ACCENT)


def add_callout(doc: Document, title: str, text: str, fill: str = PALE_BLUE) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.25)
    p.paragraph_format.right_indent = Cm(0.25)
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(7)
    set_paragraph_box(p, fill)
    r = p.add_run(f"{title}: ")
    set_run_font(r, bold=True, color=ACCENT)
    r2 = p.add_run(text)
    set_run_font(r2, size=10.8)


def add_code(doc: Document, title: str, code: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.keep_with_next = True
    r = p.add_run(title)
    set_run_font(r, size=10, bold=True, color=ACCENT_2)
    lines = code.strip("\n").splitlines()
    for idx, line in enumerate(lines):
        paragraph = doc.add_paragraph()
        paragraph.style = doc.styles["CodeBlock"]
        paragraph.paragraph_format.keep_with_next = idx < len(lines) - 1
        paragraph.paragraph_format.left_indent = Cm(0.3)
        paragraph.paragraph_format.right_indent = Cm(0.3)
        set_paragraph_box(paragraph, "F5F7F9", "D5DDE3")
        run = paragraph.add_run(line or " ")
        set_run_font(run, name="Consolas", size=8.5, color="263238")
    if lines:
        paragraph.paragraph_format.space_after = Pt(6)


def add_caption(doc: Document, prefix: str, number: str, text: str) -> None:
    p = doc.add_paragraph(style="Caption")
    p.paragraph_format.keep_with_next = False
    r1 = p.add_run(f"{prefix} {number}. ")
    set_run_font(r1, size=10, bold=True, italic=True, color=ACCENT)
    r2 = p.add_run(text)
    set_run_font(r2, size=10, italic=True, color=MUTED)


def add_table(doc: Document, number: str, title: str, headers: Sequence[str],
              rows: Sequence[Sequence[str]], widths: Sequence[float] | None = None,
              font_size: float = 9.2) -> None:
    cap = doc.add_paragraph(style="Caption")
    cap.paragraph_format.keep_with_next = True
    r1 = cap.add_run(f"Bảng {number}. ")
    set_run_font(r1, size=10, bold=True, italic=True, color=ACCENT)
    r2 = cap.add_run(title)
    set_run_font(r2, size=10, italic=True, color=MUTED)

    if widths is None:
        widths = [15.6 / len(headers)] * len(headers)
    path = render_table_image(number, headers, rows, widths, font_size)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(7)
    shape = p.add_run().add_picture(str(path), width=Cm(15.6))
    set_alt_text(shape, f"Bảng {number}", title)


def add_figure(doc: Document, number: str, path: str, title: str,
               width_cm: float = 15.6) -> None:
    full_path = ROOT / path
    if not full_path.is_file():
        raise FileNotFoundError(full_path)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(0)
    shape = p.add_run().add_picture(str(full_path), width=Cm(width_cm))
    set_alt_text(shape, f"Hình {number}", title)
    add_caption(doc, "Hình", number, title)


def add_diagram(doc: Document, number: str, title: str, lines: Sequence[str]) -> None:
    path = render_diagram_image(number, lines)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    shape = p.add_run().add_picture(str(path), width=Cm(14.8))
    set_alt_text(shape, f"Hình {number}", title)
    add_caption(doc, "Hình", number, title)


def add_front_list(doc: Document, title: str, items: Sequence[tuple[str, str]], pages: dict[str, str], prefix: str) -> None:
    p = doc.add_paragraph(title, style="FrontMatterTitle")
    p.paragraph_format.space_after = Pt(10)
    for number, item in items:
        if number not in pages or not pages[number].isdigit():
            raise RuntimeError(f"Thiếu số trang cuối cho {prefix} {number}")
        p0 = doc.add_paragraph()
        p0.paragraph_format.space_after = Pt(3)
        p0.paragraph_format.tab_stops.add_tab_stop(Cm(15.3), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        r = p0.add_run(f"{prefix} {number}. {item}\t{pages[number]}")
        set_run_font(r, size=10.5)


def add_toc(doc: Document) -> None:
    p = doc.add_paragraph("MỤC LỤC", style="FrontMatterTitle")
    entries = [
        (0, "I. TỔNG QUAN"),
        (0, "II. CƠ SỞ LÝ THUYẾT VỀ RNN"),
        (0, "III. PHÂN TÍCH DỮ LIỆU"),
        (1, "3.1. UCI Online Retail II"),
        (1, "3.2. AAPL lịch sử"),
        (0, "IV. TIỀN XỬ LÝ VÀ THIẾT KẾ THÍ NGHIỆM"),
        (0, "V. THỰC NGHIỆM PYTORCH"),
        (1, "5.1. Customer classification"),
        (1, "5.2. Stock regression"),
        (0, "VI. THỰC NGHIỆM KERAS"),
        (1, "6.1. Customer classification"),
        (1, "6.2. Stock regression"),
        (0, "VII. SO SÁNH HAI FRAMEWORK"),
        (0, "VIII. HẠN CHẾ VÀ HƯỚNG PHÁT TRIỂN"),
        (0, "IX. KẾT LUẬN"),
        (0, "TÀI LIỆU THAM KHẢO"),
    ]
    for level, title in entries:
        page = TOC_PAGES.get(title, "")
        if not page.isdigit():
            raise RuntimeError(f"Thiếu số trang TOC cuối: {title}")
        p0 = doc.add_paragraph()
        p0.paragraph_format.left_indent = Cm(0.6 * level)
        p0.paragraph_format.space_after = Pt(4)
        p0.paragraph_format.tab_stops.add_tab_stop(Cm(15.3 - 0.6 * level), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        r = p0.add_run(f"{title}\t{page}")
        set_run_font(r, size=10.8 if level else 11.2, bold=(level == 0), color=ACCENT if level == 0 else TEXT)


def heading(doc: Document, level: int, text: str, *, new_page: bool = False) -> None:
    if new_page:
        add_page_break(doc)
    p = doc.add_paragraph(text, style=f"Heading {level}")
    if level == 1:
        p.paragraph_format.page_break_before = False


def fmt4(value: float) -> str:
    return f"{float(value):.4f}"


def fmt2(value: float) -> str:
    return f"{float(value):.2f}"


def build_report() -> None:
    ensure_sources()
    pre = read_json("results/metrics/preprocessing_metadata.json")
    ptc = read_json("results/metrics/pytorch_customer_metrics.json")
    kc = read_json("results/metrics/keras_customer_metrics.json")
    pts = read_json("results/metrics/pytorch_stock_metrics.json")
    ks = read_json("results/metrics/keras_stock_metrics.json")
    comp = read_json("results/metrics/framework_comparison.json")
    env = read_json("results/metrics/environment_verification.json")
    validate_metric_contract(pre, ptc, kc, pts, ks, comp)

    doc = Document()
    configure_document(doc)

    # Cover
    for _ in range(4):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("ASSIGNMENT 06")
    set_run_font(r, size=16, bold=True, color=ACCENT_2)
    p.paragraph_format.space_after = Pt(12)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run("MẠNG NƠ-RON HỒI QUY CHO DỮ LIỆU CHUỖI")
    set_run_font(r, size=23, bold=True, color=ACCENT)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Phân tích hành vi khách hàng và dự báo giá đóng cửa AAPL\nbằng PyTorch và Keras")
    set_run_font(r, size=14, italic=True, color=MUTED)
    p.paragraph_format.space_after = Pt(28)
    line = doc.add_paragraph()
    line.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = line.add_run("────────────────────────")
    set_run_font(r, size=12, color=ACCENT_2)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("BÁO CÁO HỌC THUẬT TỔNG KẾT DỰ ÁN A06")
    set_run_font(r, size=12, bold=True, color=ACCENT)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Vanilla RNN • Binary Classification • Regression • Chronological Evaluation")
    set_run_font(r, size=10.5, color=MUTED)
    for _ in range(7):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Tháng 9 năm 2026")
    set_run_font(r, size=11.5, bold=True, color=ACCENT)

    # Front matter: fixed page breaks keep navigation deterministic.
    add_page_break(doc)
    add_toc(doc)
    add_page_break(doc)
    add_front_list(doc, "DANH MỤC HÌNH", FIGURES, FIGURE_PAGES, "Hình")
    add_page_break(doc)
    add_front_list(doc, "DANH MỤC BẢNG", TABLES, TABLE_PAGES, "Bảng")

    # I
    heading(doc, 1, "I. TỔNG QUAN", new_page=True)
    heading(doc, 2, "1.1. Bối cảnh và yêu cầu")
    add_body(doc, "A06 khảo sát Recurrent Neural Network (RNN) từ nền tảng toán học đến hai bài toán chuỗi có bản chất khác nhau: dự báo hành vi mua hàng tuần kế tiếp và dự báo giá Close của phiên giao dịch kế tiếp. Dự án dùng cùng một protocol theo thời gian, triển khai Vanilla RNN bằng PyTorch và Keras, rồi so sánh trên đúng các mẫu TEST đã khóa.")
    add_body(doc, "Mục tiêu của báo cáo không phải chứng minh RNN luôn vượt trội, mà là trình bày trọn vẹn chu trình khoa học: hiểu dữ liệu, tạo sequence đúng quan hệ nhân quả, chống data leakage, huấn luyện bằng validation, đánh giá trên TEST một lần và diễn giải kết quả theo đúng artifact thực tế.")
    add_callout(doc, "Trạng thái kiểm toán", "Toàn bộ 8 notebook đã chạy lại từ output trống; 91/91 code cell thành công, 50/50 pytest passed, bốn model tải lại được, prediction tái tạo khớp và không phát hiện data leakage. Trạng thái cuối: PASSED.", PALE_GREEN)

    heading(doc, 2, "1.2. Dữ liệu, framework và nhiệm vụ")
    add_table(doc, "1.1", "Phạm vi dữ liệu và nhiệm vụ học máy của A06",
              ["Nhánh", "Nguồn đã khóa", "Input", "Target", "Bài toán"], [
                  ["Customer", "UCI Online Retail II (1.067.371 dòng raw)", "8 tuần × 5 đặc trưng", "Có mua ở tuần kế tiếp", "Binary classification"],
                  ["Stock", "AAPL local CSV, 2015-01-02–2025-12-31", "30 phiên × 5 đặc trưng", "Close của phiên kế tiếp", "Regression"],
              ], widths=[2.0, 4.1, 3.1, 3.6, 2.8], font_size=8.8)
    add_body(doc, f"Môi trường khóa là Conda rnn312 với Python {env['environment']['python_version']}; PyTorch {env['environment']['package_versions']['torch']}, TensorFlow {env['environment']['package_versions']['tensorflow']} và Keras {env['environment']['package_versions']['keras']} đều chạy trên CPU. SEED = 42 được dùng xuyên suốt. Duration vì vậy chỉ mô tả lần chạy cụ thể, không phải benchmark phổ quát.")
    heading(doc, 2, "1.3. Nguyên tắc thực nghiệm")
    add_bullets(doc, [
        "Dữ liệu được chia theo target timestamp; không dùng random shuffle để tạo split.",
        "Feature scaler và target scaler chỉ fit trên TRAIN; validation/TEST chỉ transform.",
        "Class weight của customer chỉ tính từ nhãn TRAIN.",
        "Early stopping và chọn checkpoint chỉ dựa trên validation loss; TEST không tham gia tuning.",
        "Hai framework đọc cùng NPZ artifacts, giữ nguyên thứ tự mẫu và metadata target.",
        "Stock RNN phải được so với naive persistence baseline: Close kế tiếp ≈ Close cuối của cửa sổ.",
    ])

    # II
    heading(doc, 1, "II. CƠ SỞ LÝ THUYẾT VỀ RNN", new_page=True)
    heading(doc, 2, "2.1. Dữ liệu chuỗi và giới hạn của mạng feed-forward")
    add_body(doc, "Dữ liệu chuỗi là tập quan sát có thứ tự. Một mức giá chỉ có ý nghĩa khi đặt sau các phiên trước đó; một tuần khách hàng không mua cũng khác nhau tùy lịch sử hoạt động. Với văn bản, đổi thứ tự token có thể đổi toàn bộ nghĩa. Vì vậy, hoán vị các time step thường phá vỡ thông tin.")
    add_body(doc, "Mạng feed-forward chuẩn xử lý xₜ → NN → yₜ như từng quan sát độc lập. Nó không có biến trạng thái mang thông tin từ bước t−1 sang t. RNN bổ sung hidden state hₜ₋₁, biến phép tính hiện tại thành hàm của cả input mới và tóm tắt quá khứ.")

    heading(doc, 2, "2.2. RNN cell, hidden state và weight sharing")
    add_equation(doc, "hₜ = tanh(Wₓₕxₜ + Wₕₕhₜ₋₁ + bₕ)")
    add_equation(doc, "ŷₜ = Wₕᵧhₜ + bᵧ")
    add_table(doc, "2.1", "Ký hiệu trong một Vanilla RNN cell",
              ["Ký hiệu", "Ý nghĩa", "Vai trò"], [
                  ["x_t", "Vector đặc trưng tại time step t", "Thông tin mới"],
                  ["h_(t-1), h_t", "Hidden state trước và sau cập nhật", "Bộ nhớ nén của lịch sử"],
                  ["W_xh", "Ma trận input → hidden", "Biến đổi input"],
                  ["W_hh", "Ma trận hidden → hidden", "Truyền trạng thái qua thời gian"],
                  ["b_h", "Bias của hidden state", "Dịch chuyển kích hoạt"],
                  ["W_hy, b_y", "Tham số hidden → output", "Sinh dự báo/logit"],
                  ["tanh", "Hàm kích hoạt trong (−1, 1)", "Tạo phi tuyến"],
              ], widths=[2.2, 6.4, 7.0], font_size=9.1)
    add_body(doc, "Cùng Wₓₕ, Wₕₕ và bₕ được tái sử dụng ở mọi time step. Đây là weight sharing through time: mô hình học một quy tắc chuyển trạng thái thay vì một mạng riêng cho từng ngày/tuần. Nhờ vậy số tham số không tăng theo chiều dài sequence.")

    add_diagram(doc, "2.1", "RNN được unroll qua ba bước thời gian với tham số dùng chung", [
        "CÙNG W_xh, W_hh, b_h TẠI MỌI TIME STEP",
        "x1 + h0  →  [RNN cell]  →  h1",
        "x2 + h1  →  [RNN cell]  →  h2",
        "x3 + h2  →  [RNN cell]  →  h3  →  output",
    ])
    add_body(doc, "Unrolling không tạo thêm ba bộ tham số; nó chỉ biểu diễn cùng cell được gọi tuần tự. Trong A06, h cuối cùng đại diện cho 8 tuần customer hoặc 30 ngày giao dịch trước target.")

    heading(doc, 2, "2.3. Forward pass số nhỏ")
    add_body(doc, "Ví dụ dưới đây dùng vector một chiều và hidden size bằng 2. Mục đích là nhìn thấy rõ x₁ → h₁ → h₂ → h₃, không phải mô hình assignment.")
    add_code(doc, "Mã minh họa NumPy rút gọn", '''import numpy as np

W_xh = np.array([[0.5], [-0.4]])
W_hh = np.array([[0.2, 0.1], [-0.3, 0.4]])
b_h  = np.array([0.0, 0.1])
x = [np.array([1.0]), np.array([0.5]), np.array([-0.2])]

h = np.zeros(2)
for t, x_t in enumerate(x, start=1):
    h = np.tanh(W_xh @ x_t + W_hh @ h + b_h)
    print(f"h{t} = {np.round(h, 4)}")''')
    add_body(doc, "Mỗi hₜ phụ thuộc gián tiếp vào mọi input trước đó thông qua chuỗi trạng thái. Nếu đổi x₁, cả h₁, h₂ và h₃ đều thay đổi dù x₂, x₃ giữ nguyên.")

    heading(doc, 2, "2.4. Shape và các kiểu input/output")
    add_body(doc, "Một sequence đơn có shape (sequence_length, features). Khi ghép batch, shape chuẩn với batch_first=True là (batch_size, sequence_length, features). Ví dụ customer: (32, 8, 5); stock: (32, 30, 5). PyTorch nn.RNN trả output cho mọi time step và h_n là hidden state cuối của từng layer; Keras SimpleRNN mặc định chỉ trả output cuối, còn return_sequences=True trả toàn bộ chuỗi.")
    add_table(doc, "2.2", "Các kiểu ánh xạ input/output của mô hình chuỗi",
              ["Kiểu", "Ví dụ", "Liên hệ A06"], [
                  ["one-to-one", "Một vector → một nhãn", "Không phải bài toán chuỗi điển hình"],
                  ["one-to-many", "Ảnh → chuỗi mô tả", "Không sử dụng"],
                  ["many-to-one", "Nhiều time step → một dự báo", "Cả hai bài toán A06"],
                  ["many-to-many", "Dịch máy/gán nhãn từng token", "Không sử dụng"],
              ], widths=[3.0, 6.1, 6.5])

    heading(doc, 2, "2.5. Loss, BPTT và gradient")
    add_body(doc, "Dự báo được so với target để tạo loss. Backpropagation Through Time (BPTT) lan truyền gradient ngược từ output qua các bản sao thời gian của cell rồi cộng đóng góp lên bộ tham số dùng chung. Optimizer dùng gradient đó để cập nhật W và b. Với classification, A06 tối ưu binary cross-entropy; với regression, A06 tối ưu mean squared error trong không gian target đã chuẩn hóa.")
    add_body(doc, "Qua nhiều time step, gradient chứa tích lặp của các đạo hàm và Wₕₕ. Nếu độ lớn hiệu dụng nhỏ hơn 1, tích có thể tiến nhanh về 0 (vanishing gradient); nếu lớn hơn 1, tích có thể tăng mất kiểm soát (exploding gradient). Ví dụ 0,5²⁰ ≈ 9,54×10⁻⁷ trong khi 1,5²⁰ ≈ 3.325. Gradient clipping chặn norm quá lớn để ổn định cập nhật, nhưng không giải quyết tận gốc vanishing gradient.")

    heading(doc, 2, "2.6. Động lực của LSTM và GRU")
    add_body(doc, "LSTM bổ sung cell state và ba cơ chế điều tiết: forget gate quyết định phần ký ức bị quên, input gate kiểm soát thông tin mới được ghi, output gate chọn phần trạng thái được phát ra. GRU gọn hơn với update gate để cân bằng trạng thái cũ–mới và reset gate để quyết định mức sử dụng lịch sử khi tạo candidate state.")
    add_table(doc, "2.3", "So sánh Vanilla RNN, LSTM và GRU",
              ["Mô hình", "Cơ chế nhớ", "Ưu điểm", "Hạn chế / vai trò A06"], [
                  ["Vanilla RNN", "Một hidden state", "Đơn giản, minh bạch, ít tham số", "Khó giữ phụ thuộc dài; baseline chính"],
                  ["LSTM", "Cell state + 3 gate", "Kiểm soát dòng thông tin dài hạn", "Nhiều tham số; chỉ nêu hướng mở rộng"],
                  ["GRU", "Hidden state + 2 gate", "Gọn hơn LSTM, có gating", "Vẫn phức tạp hơn baseline; chưa huấn luyện"],
              ], widths=[2.4, 3.5, 4.7, 5.0], font_size=8.8)
    add_callout(doc, "Vì sao vẫn dùng Vanilla RNN?", "A06 ưu tiên hiểu đúng recurrence, sequence, hidden state và protocol đánh giá. LSTM/GRU là phần mở rộng hợp lý sau khi baseline và leakage controls đã rõ, không phải sự thay thế âm thầm cho kiến trúc được khóa.", PALE_GOLD)

    # III
    heading(doc, 1, "III. PHÂN TÍCH DỮ LIỆU", new_page=True)
    heading(doc, 2, "3.1. UCI Online Retail II")
    heading(doc, 3, "3.1.1. Nguồn, cấu trúc và thời gian")
    add_body(doc, "Online Retail II là dữ liệu transaction-level của một nhà bán lẻ trực tuyến tại Anh, ghi từng dòng sản phẩm trong hóa đơn từ 01/12/2009 đến 09/12/2011. Tính tuần tự của hành vi mua, cùng định danh khách hàng và thời điểm hóa đơn, cho phép tổng hợp một regular timeline theo customer-week.")
    sheets = pre["customer"]["raw_source"]["workbook"]["sheets"]
    add_table(doc, "3.1", "Cấu trúc workbook Online Retail II",
              ["Sheet", "Số dòng", "Cột raw"], [
                  [s["name"], f"{s['rows']:,}", ", ".join(s["columns"])] for s in sheets
              ], widths=[3.2, 2.2, 10.2], font_size=8.7)
    add_body(doc, "Ba tên cột lịch sử được chuẩn hóa chỉ trong DataFrame: Invoice → InvoiceNo, Price → UnitPrice và Customer ID → CustomerID. Workbook raw không bị ghi đè; SHA-256 trước và sau xử lý giữ nguyên.")

    heading(doc, 3, "3.1.2. Missing values, cancellation và cleaning")
    cc = pre["customer"]["cleaning_counts"]
    add_table(doc, "3.2", "Kết quả kiểm tra và làm sạch giao dịch customer",
              ["Hạng mục", "Giá trị", "Diễn giải"], [
                  ["Raw rows", f"{cc['raw_rows']:,}", "Tổng hai sheet"],
                  ["Exact duplicates", f"{cc['exact_duplicates_removed']:,}", "Loại trước khi cộng quantity/revenue"],
                  ["CustomerID thiếu/không hợp lệ", f"{cc['invalid_or_missing_customer_id_raw']:,}", "Không thể gán lịch sử cho customer"],
                  ["Cancellation lines", f"{cc['cancellation_lines_raw']:,}", "InvoiceNo bắt đầu bằng C"],
                  ["Quantity ≤ 0", f"{cc['nonpositive_quantity_lines_raw']:,}", "Không phải purchase dương"],
                  ["UnitPrice ≤ 0", f"{cc['nonpositive_unit_price_lines_raw']:,}", "Không tạo revenue mua hợp lệ"],
                  ["Cleaned rows", f"{cc['cleaned_rows']:,}", f"{cc['cleaned_customers']:,} customers; {cc['cleaned_invoices']:,} invoices"],
              ], widths=[4.5, 2.6, 8.5], font_size=8.9)
    add_body(doc, "Dữ liệu raw có 243.007 CustomerID thiếu và 4.382 Description thiếu; InvoiceDate không có giá trị không hợp lệ. Cancellation được nhận diện theo quy ước InvoiceNo bắt đầu bằng C: 19.494 dòng thuộc 8.292 hóa đơn hủy, trong đó 19.493 dòng có Quantity âm. Với nhiệm vụ dự báo purchase behavior, các dòng này được đếm và giải thích trước khi loại, vì chúng biểu diễn hoàn/hủy chứ không phải một giao dịch mua dương mới.")
    add_body(doc, "Cleaning sản xuất lần lượt loại duplicate, yêu cầu CustomerID nguyên dương và InvoiceDate hợp lệ, loại cancellation, giữ Quantity > 0 và UnitPrice > 0, sau đó tạo Revenue = Quantity × UnitPrice. Kết quả còn 779.425 dòng, nhưng raw data luôn được bảo toàn.")

    heading(doc, 3, "3.1.3. Hành vi theo thời gian và phân phối")
    add_figure(doc, "3.1", "results/figures/customer_eda/monthly_activity.png", "Hoạt động mua hàng theo tháng trong Online Retail II", 15.5)
    add_body(doc, "Số đơn, revenue và số khách hàng active cùng tăng rõ vào giai đoạn cuối năm trong cả hai chu kỳ. Điểm cuối 12/2011 thấp bất thường vì dữ liệu chỉ kéo dài đến ngày 09/12; đây là lý do pipeline chỉ dùng các tuần lịch đầy đủ đến 2011-11-28.")
    add_figure(doc, "3.2", "results/figures/customer_eda/customer_distributions.png", "Phân phối chi tiêu, tần suất mua và số đơn của khách hàng", 15.5)
    add_body(doc, "Các phân phối lệch phải: phần lớn customer hoạt động ít tuần và có ít đơn, trong khi một nhóm nhỏ rất tích cực. Vì invoice lines trong cùng hóa đơn không phải các đơn khác nhau, order_count bắt buộc dùng số InvoiceNo duy nhất.")
    add_figure(doc, "3.3", "results/figures/customer_eda/inactive_week_example.png", "Minh họa tuần không hoạt động được điền bằng 0", 15.2)
    add_body(doc, "RNN cần time step cách đều. Sau khi tạo lịch tuần thứ Hai–Chủ nhật, tuần không mua được điền 0 cho toàn bộ behavior features. Zero ở đây mang ý nghĩa quan sát được: không có purchase hợp lệ trong tuần, không phải missing value tùy tiện.")
    add_table(doc, "3.3", "Định nghĩa năm đặc trưng customer-week",
              ["Feature", "Định nghĩa"], [
                  ["total_spent", "Tổng Revenue của purchase lines trong tuần"],
                  ["total_quantity", "Tổng Quantity dương trong tuần"],
                  ["order_count", "Số InvoiceNo duy nhất trong tuần"],
                  ["unique_products", "Số StockCode duy nhất trong tuần"],
                  ["active_flag", "1 nếu có ít nhất một purchase hợp lệ, ngược lại 0"],
              ], widths=[4.2, 11.4])
    add_figure(doc, "3.4", "results/figures/customer_eda/prototype_label_balance.png", "Phân bố nhãn mua hàng tuần kế tiếp ở bước EDA nguyên mẫu", 8.8)
    add_body(doc, "EDA nguyên mẫu tạo 350.864 target customer-week: 93,61% không mua và 6,39% có mua. Mất cân bằng lớp dự báo rằng accuracy có thể cao ngay cả với mô hình thiên về class 0; vì vậy evaluation phải xem Precision, Recall, F1 và ROC-AUC. Class weight chưa được tính ở EDA mà chỉ được tính từ TRAIN sau split.")

    heading(doc, 2, "3.2. AAPL lịch sử", new_page=True)
    heading(doc, 3, "3.2.1. Nguồn, OHLCV và integrity")
    add_body(doc, "Snapshot AAPL được tải một lần qua yfinance với auto_adjust=False và lưu cố định tại datasets/stock/AAPL_2015_2025.csv. Mọi notebook sau chỉ đọc local CSV. Khoảng calendar yêu cầu là 2015-01-01–2025-12-31; ngày giao dịch thực tế đầu tiên là 2015-01-02 vì 01/01 không có phiên, ngày cuối là 2025-12-31.")
    add_table(doc, "3.4", "Ý nghĩa các trường giá AAPL",
              ["Trường", "Ý nghĩa trực quan"], [
                  ["Date", "Ngày giao dịch"], ["Open", "Giá đầu phiên"],
                  ["High", "Giá cao nhất trong phiên"], ["Low", "Giá thấp nhất trong phiên"],
                  ["Close", "Giá cuối phiên; target của A06"],
                  ["Adj Close", "Close điều chỉnh theo corporate actions; chỉ audit, không làm feature"],
                  ["Volume", "Số cổ phiếu giao dịch trong phiên"],
              ], widths=[3.2, 12.4])
    integrity = pre["stock"]["integrity"]
    add_table(doc, "3.5", "Kết quả kiểm tra dữ liệu AAPL",
              ["Kiểm tra", "Kết quả"], [
                  ["Shape", f"{integrity['raw_rows']:,} dòng × 7 cột"],
                  ["Date range", f"{integrity['date_min']} – {integrity['date_max']}"],
                  ["Missing / invalid date", "0 / 0"],
                  ["Duplicate date", str(integrity["duplicate_dates"])],
                  ["Non-positive OHLC / Volume", f"{integrity['nonpositive_price_rows']} / {integrity['nonpositive_volume_rows']}"],
                  ["High/Low inconsistent", f"{integrity['invalid_high_rows']} / {integrity['invalid_low_rows']}"],
                  ["Sort order", "Đã tăng dần; assert monotonic đạt"],
              ], widths=[6.6, 9.0])

    heading(doc, 3, "3.2.2. Trend, volatility và moving average")
    add_figure(doc, "3.5", "results/figures/stock_eda/close_volume_history.png", "Diễn biến Close và Volume của AAPL trong snapshot đã khóa", 15.5)
    add_body(doc, "Close có xu hướng tăng mạnh theo thang giá tuyệt đối nhưng trải qua nhiều giai đoạn giảm và biến động. Volume thay đổi mạnh theo phiên và không đồng biến đơn giản với price level. Độ lớn giá thay đổi theo thời gian cho thấy chuỗi không stationarity theo scale.")
    add_figure(doc, "3.6", "results/figures/stock_eda/close_moving_averages.png", "Close của AAPL cùng MA(7) và MA(30) dùng cho EDA", 15.5)
    add_body(doc, "MA(7) phản ứng nhanh hơn MA(30), còn MA(30) làm trơn tốt hơn nhưng trễ hơn. Hai đường này chỉ giúp quan sát xu hướng, không được đưa vào năm model features để tránh thay đổi contract.")
    add_figure(doc, "3.7", "results/figures/stock_eda/daily_return_distribution.png", "Phân phối tỷ suất sinh lợi ngày của AAPL", 13.4)
    add_body(doc, "Daily return tập trung quanh 0 nhưng có đuôi dày và một số phiên cực trị. Dự báo chính xác price level vì vậy khó: mô hình phải đồng thời theo kịp trend, regime biến động và cú sốc chưa xuất hiện trong 30 ngày lịch sử. Báo cáo không đưa ra khuyến nghị đầu tư.")
    add_callout(doc, "Task stock đã khóa", "Input = 30 phiên trước × 5 features [Open, High, Low, Close, Volume]. Target = unadjusted Close của phiên ngay sau cửa sổ. Đây là regression. Baseline = Close cuối cùng trong input window.", PALE_GOLD)

    # IV
    heading(doc, 1, "IV. TIỀN XỬ LÝ VÀ THIẾT KẾ THÍ NGHIỆM", new_page=True)
    heading(doc, 2, "4.1. Tạo sequence đúng quan hệ thời gian")
    add_body(doc, "Customer sequence cho target tuần t chứa đúng các tuần t−8,…,t−1 của cùng CustomerID; y là active_flag ở t. Stock sequence cho target ngày giao dịch t chứa đúng 30 dòng ngay trước t; y là Close tại t. Audit đã inverse-transform và đối chiếu exhaustive 350.864 customer samples cùng 2.736 stock samples với nguồn raw.")
    add_equation(doc, "Customer: Xᵢ = [wₜ₋₈,…,wₜ₋₁]  →  yᵢ = active_flag(wₜ)")
    add_equation(doc, "Stock: Xᵢ = [dₜ₋₃₀,…,dₜ₋₁]  →  yᵢ = Close(dₜ)")
    add_code(doc, "Mẫu logic tạo stock window từ pipeline", '''values = data[["Open", "High", "Low", "Close", "Volume"]].to_numpy()
windows = sliding_window_view(values, window_shape=30, axis=0)[:-1]
X = windows.transpose(0, 2, 1).copy()       # (samples, 30, 5)
y = values[30:, close_position].copy()     # Close ngay sau mỗi cửa sổ
naive = values[29:-1, close_position].copy()''')

    heading(doc, 2, "4.2. Chronological split và shape thực tế")
    add_diagram(doc, "4.1", "Sơ đồ chia tập theo thời gian và nguyên tắc chống leakage", [
        "PAST ─────────────────────────────────────────────────────────→ FUTURE",
        "TRAIN (fit scaler / fit model)  │  VALIDATION (early stopping)  │  TEST (final only)",
        "target timestamp quyết định split; input chỉ dùng lịch sử trước target",
    ])
    cp = pre["customer"]["splits"]
    sp = pre["stock"]["splits"]
    add_table(doc, "4.1", "Shape và khoảng thời gian của các split sản xuất",
              ["Task / split", "X shape", "Target range", "Số mẫu / class 1"], [
                  ["Customer TRAIN", str(tuple(cp['train']['X_shape'])), f"{cp['train']['target_week_min']} – {cp['train']['target_week_max']}", f"{cp['train']['y_shape'][0]:,} / {cp['train']['class_distribution']['class_1']:,}"],
                  ["Customer VAL", str(tuple(cp['val']['X_shape'])), f"{cp['val']['target_week_min']} – {cp['val']['target_week_max']}", f"{cp['val']['y_shape'][0]:,} / {cp['val']['class_distribution']['class_1']:,}"],
                  ["Customer TEST", str(tuple(cp['test']['X_shape'])), f"{cp['test']['target_week_min']} – {cp['test']['target_week_max']}", f"{cp['test']['y_shape'][0]:,} / {cp['test']['class_distribution']['class_1']:,}"],
                  ["Stock TRAIN", str(tuple(sp['train']['X_shape'])), f"{sp['train']['target_date_min']} – {sp['train']['target_date_max']}", f"{sp['train']['y_shape'][0]:,}"],
                  ["Stock VAL", str(tuple(sp['val']['X_shape'])), f"{sp['val']['target_date_min']} – {sp['val']['target_date_max']}", f"{sp['val']['y_shape'][0]:,}"],
                  ["Stock TEST", str(tuple(sp['test']['X_shape'])), f"{sp['test']['target_date_min']} – {sp['test']['target_date_max']}", f"{sp['test']['y_shape'][0]:,}"],
              ], widths=[3.2, 3.4, 5.5, 3.5], font_size=8.5)
    add_body(doc, "Các tỷ lệ gần 70/15/15 nhưng boundary được chọn tại target timestamp gần nhất và không tách cùng một target week/date qua nhiều split. Validation/TEST được phép dùng context lịch sử trước boundary; điều này không phải leakage vì mọi feature vẫn xảy ra trước target.")

    heading(doc, 2, "4.3. Scaling, class weight và leakage controls")
    add_body(doc, "StandardScaler của customer được fit trên 247.458 × 8 = 1.979.664 time steps TRAIN. Stock feature scaler được fit trên 1.915 × 30 = 57.450 time steps TRAIN; stock target scaler được fit trên đúng 1.915 TRAIN targets. Sau dự báo, y và prediction được inverse-transform về USD trước khi tính MAE/RMSE/R².")
    add_body(doc, f"TRAIN customer có 231.206 class 0 và 16.252 class 1. Positive weight = 231.206 / 16.252 = {pre['customer']['class_weight']['positive_weight']:.12f}; không dùng nhãn validation hoặc TEST để tính trọng số.")
    add_table(doc, "4.2", "Các kiểm soát chống data leakage đã được audit",
              ["Rủi ro", "Kiểm soát", "Kết quả"], [
                  ["Future week/day lọt vào input", "Đối chiếu raw từng window với target position", "PASS"],
                  ["Split chồng target range", "TRAIN max < VAL min < TEST min", "PASS"],
                  ["Scaler thấy validation/TEST", "Tái tính mean/scale từ raw TRAIN", "PASS"],
                  ["Class weight dùng TEST", "So khớp 231.206 / 16.252", "PASS"],
                  ["Early stopping dùng TEST", "Đọc training loop/callback và artifact", "PASS"],
                  ["Hai framework lệch sample", "So key/label/date từng dòng prediction", "PASS"],
              ], widths=[4.4, 8.2, 3.0], font_size=8.7)

    # V
    heading(doc, 1, "V. THỰC NGHIỆM PYTORCH", new_page=True)
    heading(doc, 2, "5.1. Customer classification")
    heading(doc, 3, "5.1.1. Model và training")
    add_body(doc, "CustomerRNN dùng nn.RNN(input_size=5, hidden_size=64, num_layers=1, batch_first=True). Tensor hidden có shape (1, batch, 64); hidden[-1] là representation cuối của 8 tuần và được đưa qua Linear(64, 1) để sinh một logit. Không có Sigmoid trong model vì BCEWithLogitsLoss gộp sigmoid và binary cross-entropy theo cách số học ổn định hơn.")
    add_code(doc, "Kiến trúc PyTorch customer thực tế", '''class CustomerRNN(nn.Module):
    def __init__(self):
        super().__init__()
        self.rnn = nn.RNN(5, 64, batch_first=True, nonlinearity="tanh")
        self.classifier = nn.Linear(64, 1)

    def forward(self, inputs):
        _, hidden = self.rnn(inputs)
        return self.classifier(hidden[-1]).squeeze(-1)''')
    add_table(doc, "5.1", "Thiết lập PyTorch Customer RNN",
              ["Thuộc tính", "Giá trị"], [
                  ["Parameters", f"{ptc['architecture']['parameter_count']:,}"],
                  ["Batch / optimizer", f"{ptc['training']['batch_size']} / Adam, lr={ptc['training']['learning_rate']}"],
                  ["Loss", f"BCEWithLogitsLoss, pos_weight={ptc['class_weighting']['positive_weight']:.6f}"],
                  ["Gradient clipping", f"norm ≤ {ptc['training']['gradient_clip_norm']}"],
                  ["Early stopping", f"max 30, patience {ptc['training']['patience']}, validation loss"],
                  ["Epochs / best epoch", f"{ptc['training']['epochs_actually_run']} / {ptc['training']['best_epoch']}"],
                  ["Best validation loss", fmt4(ptc['training']['best_validation_loss'])],
                  ["Device / duration", f"{ptc['device']} / {ptc['training']['training_duration_seconds']:.2f} s"],
              ], widths=[5.5, 10.1])
    pm = ptc["evaluation"]["metrics"]
    pcm = ptc["evaluation"]["confusion_matrix_tn_fp_fn_tp"]
    add_table(doc, "5.2", "Kết quả TEST của PyTorch Customer RNN",
              ["Accuracy", "Precision", "Recall", "F1", "ROC-AUC", "TN / FP / FN / TP"], [[
                  fmt4(pm['accuracy']), fmt4(pm['precision']), fmt4(pm['recall']), fmt4(pm['f1']), fmt4(pm['roc_auc']),
                  f"{pcm['tn']:,} / {pcm['fp']:,} / {pcm['fn']:,} / {pcm['tp']:,}"
              ]], widths=[2.2, 2.2, 2.2, 2.0, 2.2, 4.8], font_size=8.4)
    add_body(doc, "Accuracy 0,7582 không đủ để kết luận tốt vì TEST chỉ có khoảng 6,99% class 1. Recall 0,5722 cho thấy mô hình bắt được hơn một nửa các tuần mua, nhưng precision 0,1587 phản ánh nhiều false positive. ROC-AUC 0,7009 thể hiện khả năng xếp hạng tốt hơn ngẫu nhiên ở mức vừa phải.")
    add_figure(doc, "5.1", "results/figures/pytorch_customer/training_validation_loss.png", "Loss huấn luyện và validation của PyTorch Customer RNN", 14.8)
    add_figure(doc, "5.2", "results/figures/pytorch_customer/test_confusion_matrix.png", "Confusion matrix trên TEST của PyTorch Customer RNN", 12.5)
    add_figure(doc, "5.3", "results/figures/pytorch_customer/test_roc_curve.png", "Đường ROC trên TEST của PyTorch Customer RNN", 12.5)

    heading(doc, 2, "5.2. Stock regression", new_page=True)
    heading(doc, 3, "5.2.1. Model, scale và baseline")
    add_body(doc, "StockRNN giữ cùng ý tưởng nn.RNN(5, 64, batch_first=True) → Linear(64, 1), nhưng output là số thực không qua sigmoid. MSELoss được tối ưu trên target đã standardize. Best checkpoint được chọn theo validation MSE; TEST chỉ được inverse-transform và đánh giá sau khi model selection kết thúc.")
    add_table(doc, "5.3", "Thiết lập và kết quả PyTorch Stock RNN",
              ["Nhóm", "Chỉ số", "Giá trị"], [
                  ["Training", "Parameters / batch", f"{pts['parameter_count']:,} / {pts['batch_size']}"],
                  ["Training", "Epochs / best", f"{pts['epochs_actually_run']} / {pts['best_epoch']}"],
                  ["Training", "Best val MSE (scaled)", f"{pts['best_validation_loss']:.6f}"],
                  ["Training", "Device / duration", f"{pts['device']} / {pts['training_duration_seconds']:.2f} s"],
                  ["RNN TEST", "MAE / RMSE / R²", f"{fmt4(pts['MAE'])} / {fmt4(pts['RMSE'])} / {fmt4(pts['R2'])}"],
                  ["Naive TEST", "MAE / RMSE / R²", f"{fmt4(pts['naive_MAE'])} / {fmt4(pts['naive_RMSE'])} / {fmt4(pts['naive_R2'])}"],
              ], widths=[3.0, 5.1, 7.5], font_size=8.9)
    add_body(doc, "MAE và RMSE được báo cáo bằng USD; R² không có đơn vị. PyTorch RNN có MAE 20,3393 USD, RMSE 23,5788 USD và R² −0,0269; naive baseline có MAE 2,6177 USD, RMSE 3,8789 USD và R² 0,9722. Vì vậy kết quả thực tế không ủng hộ tuyên bố RNN vượt baseline.")
    add_figure(doc, "5.4", "results/figures/pytorch_stock/training_validation_loss.png", "Loss huấn luyện và validation của PyTorch Stock RNN", 14.8)
    add_figure(doc, "5.5", "results/figures/pytorch_stock/test_actual_vs_predicted.png", "Close thực tế và dự báo trên TEST của PyTorch Stock RNN", 15.5)
    add_body(doc, "Đường dự báo RNN bám một phần dao động ngắn hạn nhưng bị nén mạnh khi price level tăng cuối TEST, tạo bias thấp. Naive dùng giá gần nhất nên theo sát local level hơn nhiều trong một chuỗi giá có tính liên tục cao.")
    add_figure(doc, "5.6", "results/figures/pytorch_stock/rnn_vs_naive_metrics.png", "So sánh PyTorch Stock RNN với naive baseline", 14.8)

    # VI
    heading(doc, 1, "VI. THỰC NGHIỆM KERAS", new_page=True)
    heading(doc, 2, "6.1. Customer classification")
    heading(doc, 3, "6.1.1. SimpleRNN, fit() và callbacks")
    add_body(doc, "Keras nhận cùng input shape (8, 5). SimpleRNN(64, tanh) trả representation cuối, Dense(1, sigmoid) chuyển representation thành probability. binary_crossentropy là loss tương ứng. model.fit() quản lý batch loop; EarlyStopping monitor val_loss với patience 5 và restore_best_weights=True. TRAIN-derived class_weight được dùng nhất quán với PyTorch.")
    add_code(doc, "Kiến trúc Keras customer thực tế", '''model = keras.Sequential([
    keras.Input(shape=(8, 5), name="customer_sequence"),
    keras.layers.SimpleRNN(64, activation="tanh"),
    keras.layers.Dense(1, activation="sigmoid"),
])
model.compile(optimizer=keras.optimizers.Adam(1e-3),
              loss="binary_crossentropy")''')
    add_table(doc, "6.1", "Thiết lập Keras Customer SimpleRNN",
              ["Thuộc tính", "Giá trị"], [
                  ["Parameters", f"{kc['architecture']['parameter_count']:,}"],
                  ["Batch / optimizer", f"{kc['training']['batch_size']} / Adam, lr={kc['training']['learning_rate']}"],
                  ["Loss / class weight", f"binary_crossentropy / {kc['class_weighting']['positive_weight']:.6f}"],
                  ["EarlyStopping", f"patience {kc['training']['patience']}, restore_best_weights=True"],
                  ["Epochs / best epoch", f"{kc['training']['epochs_actually_run']} / {kc['training']['best_epoch']}"],
                  ["Best validation loss", fmt4(kc['training']['best_validation_loss'])],
                  ["Device / duration", f"{kc['device']} / {kc['training']['training_duration_seconds']:.2f} s"],
              ], widths=[5.5, 10.1])
    km = kc["evaluation"]["metrics"]
    kcm = kc["evaluation"]["confusion_matrix_tn_fp_fn_tp"]
    add_table(doc, "6.2", "Kết quả TEST của Keras Customer SimpleRNN",
              ["Accuracy", "Precision", "Recall", "F1", "ROC-AUC", "TN / FP / FN / TP"], [[
                  fmt4(km['accuracy']), fmt4(km['precision']), fmt4(km['recall']), fmt4(km['f1']), fmt4(km['roc_auc']),
                  f"{kcm['tn']:,} / {kcm['fp']:,} / {kcm['fn']:,} / {kcm['tp']:,}"
              ]], widths=[2.2, 2.2, 2.2, 2.0, 2.2, 4.8], font_size=8.4)
    add_body(doc, "Keras đạt Accuracy 0,7600, F1 0,2491 và ROC-AUC 0,7005. Confusion matrix gần như trùng PyTorch: mô hình ưu tiên recall cho minority class nhờ class weighting, đổi lại false positive cao. Không có lý do để diễn giải chênh lệch phần nghìn như ưu thế framework tổng quát.")
    add_figure(doc, "6.1", "results/figures/keras_customer/training_validation_loss.png", "Loss huấn luyện và validation của Keras Customer SimpleRNN", 14.8)
    add_figure(doc, "6.2", "results/figures/keras_customer/test_confusion_matrix.png", "Confusion matrix trên TEST của Keras Customer SimpleRNN", 12.5)
    add_figure(doc, "6.3", "results/figures/keras_customer/test_roc_curve.png", "Đường ROC trên TEST của Keras Customer SimpleRNN", 12.5)

    heading(doc, 2, "6.2. Stock regression", new_page=True)
    heading(doc, 3, "6.2.1. Linear output và real-scale evaluation")
    add_body(doc, "Keras Stock dùng SimpleRNN(64, tanh) → Dense(1, linear). Không dùng sigmoid vì target là giá thực. MSE được tối ưu trong standardized target space; predict() tạo scaled output rồi stock_target_scaler inverse-transform về USD. Cả actual và naive_close là cùng 410 TEST dates như PyTorch.")
    add_table(doc, "6.3", "Thiết lập và kết quả Keras Stock SimpleRNN",
              ["Nhóm", "Chỉ số", "Giá trị"], [
                  ["Training", "Parameters / batch", f"{ks['parameter_count']:,} / {ks['batch_size']}"],
                  ["Training", "Epochs / best", f"{ks['epochs_actually_run']} / {ks['best_epoch']}"],
                  ["Training", "Best val MSE (scaled)", f"{ks['best_validation_loss']:.6f}"],
                  ["Training", "Device / duration", f"{ks['device']} / {ks['training_duration_seconds']:.2f} s"],
                  ["RNN TEST", "MAE / RMSE / R²", f"{fmt4(ks['MAE'])} / {fmt4(ks['RMSE'])} / {fmt4(ks['R2'])}"],
                  ["Naive TEST", "MAE / RMSE / R²", f"{fmt4(ks['naive_MAE'])} / {fmt4(ks['naive_RMSE'])} / {fmt4(ks['naive_R2'])}"],
              ], widths=[3.0, 5.1, 7.5], font_size=8.9)
    add_body(doc, "Keras RNN có MAE 23,1947 USD, RMSE 27,8703 USD và R² −0,4347; kém naive baseline rõ rệt. Best epoch là 30, nghĩa là loss validation vẫn giảm đến giới hạn epoch và EarlyStopping chưa kích hoạt. Điều này không cấp phép dùng TEST để tăng epoch hoặc chọn lại model.")
    add_figure(doc, "6.4", "results/figures/keras_stock/training_validation_loss.png", "Loss huấn luyện và validation của Keras Stock SimpleRNN", 14.8)
    add_figure(doc, "6.5", "results/figures/keras_stock/test_actual_vs_predicted.png", "Close thực tế và dự báo trên TEST của Keras Stock SimpleRNN", 15.5)
    add_body(doc, "Tương tự PyTorch, dự báo Keras bị nén và thấp hơn actual ở nửa sau TEST. Một validation loss nhỏ trong scaled space không bảo đảm khả năng theo kịp price regime mới ở real scale.")
    add_figure(doc, "6.6", "results/figures/keras_stock/rnn_vs_naive_metrics.png", "So sánh Keras Stock SimpleRNN với naive baseline", 14.8)

    # VII
    heading(doc, 1, "VII. SO SÁNH HAI FRAMEWORK", new_page=True)
    heading(doc, 2, "7.1. Customer: kết quả gần tương đương")
    add_table(doc, "7.1", "So sánh kết quả customer trên cùng TEST",
              ["Model", "Accuracy", "Precision", "Recall", "F1", "ROC-AUC", "Params", "Epochs", "Thời gian"], [
                  ["PyTorch RNN", fmt4(pm['accuracy']), fmt4(pm['precision']), fmt4(pm['recall']), fmt4(pm['f1']), fmt4(pm['roc_auc']), f"{ptc['architecture']['parameter_count']:,}", str(ptc['training']['epochs_actually_run']), f"{ptc['training']['training_duration_seconds']:.2f}s"],
                  ["Keras SimpleRNN", fmt4(km['accuracy']), fmt4(km['precision']), fmt4(km['recall']), fmt4(km['f1']), fmt4(km['roc_auc']), f"{kc['architecture']['parameter_count']:,}", str(kc['training']['epochs_actually_run']), f"{kc['training']['training_duration_seconds']:.2f}s"],
              ], widths=[2.7, 1.5, 1.5, 1.4, 1.3, 1.6, 1.3, 1.1, 1.9], font_size=7.8)
    add_body(doc, f"Hai mô hình đồng ý nhãn ở {comp['customer']['prediction_agreement_rate']*100:.2f}% mẫu và probability correlation = {comp['customer']['probability_correlation']:.4f}; chỉ bất đồng {comp['customer']['prediction_disagreement_count']:,}/53.052 dự báo. Keras cao hơn rất nhẹ về Accuracy, Precision và F1; PyTorch cao hơn rất nhẹ về Recall và ROC-AUC. Các delta nhỏ không xác lập một framework chiến thắng.")
    add_figure(doc, "7.1", "results/figures/comparison/customer_metrics_comparison.png", "So sánh metric customer giữa PyTorch và Keras", 15.5)
    add_figure(doc, "7.2", "results/figures/comparison/customer_confusion_matrices.png", "So sánh confusion matrix customer giữa hai framework", 15.5)
    add_figure(doc, "7.3", "results/figures/comparison/customer_roc_comparison.png", "So sánh ROC customer giữa hai framework", 14.5)

    heading(doc, 2, "7.2. Stock: naive baseline thắng rõ rệt")
    add_table(doc, "7.2", "So sánh kết quả stock trên cùng TEST",
              ["Model", "MAE (USD) ↓", "RMSE (USD) ↓", "R² ↑", "Kết luận so với naive"], [
                  ["PyTorch RNN", fmt4(pts['MAE']), fmt4(pts['RMSE']), fmt4(pts['R2']), "Kém hơn"],
                  ["Keras SimpleRNN", fmt4(ks['MAE']), fmt4(ks['RMSE']), fmt4(ks['R2']), "Kém hơn"],
                  ["Naive last Close", fmt4(pts['naive_MAE']), fmt4(pts['naive_RMSE']), fmt4(pts['naive_R2']), "Tốt nhất"],
              ], widths=[3.8, 2.7, 2.7, 2.1, 4.3], font_size=8.8)
    add_body(doc, "Naive values giống tuyệt đối giữa hai notebook vì được tính trên cùng target dates và cùng last Close. Cả hai RNN đều có R² âm và sai số lớn hơn nhiều; câu trả lời thực nghiệm là RNN không đánh bại naive baseline. Điều này nhấn mạnh rằng low training/validation loss không thể thay thế một baseline phù hợp với cấu trúc chuỗi.")
    add_figure(doc, "7.4", "results/figures/comparison/stock_metrics_comparison.png", "So sánh metric stock giữa hai RNN và naive baseline", 15.5)
    add_figure(doc, "7.5", "results/figures/comparison/stock_actual_vs_predicted.png", "So sánh dự báo stock theo thời gian trên cùng TEST", 15.5)

    heading(doc, 2, "7.3. Khác biệt triển khai, không đồng nhất với chất lượng dự báo", new_page=True)
    add_table(doc, "7.3", "Khác biệt thực hành giữa PyTorch và Keras",
              ["Khía cạnh", "PyTorch", "Keras"], [
                  ["Data input", "Dataset + DataLoader; shuffle chỉ TRAIN", "NumPy arrays đưa vào fit(); validation_data riêng"],
                  ["Training loop", "Forward, zero_grad, backward, step viết tường minh", "model.fit() quản lý vòng lặp"],
                  ["Early stopping", "Tự lưu best state theo val loss", "Callback + restore_best_weights"],
                  ["Device", "Chuyển tensor/model explicit", "Backend/runtime quản lý ở mức cao hơn"],
                  ["Output customer", "Logit + BCEWithLogitsLoss", "Sigmoid probability + binary crossentropy"],
                  ["Mức kiểm soát", "Chi tiết, thuận tiện debug từng bước", "Ngắn gọn, thuận tiện baseline nhanh"],
              ], widths=[3.2, 6.2, 6.2], font_size=8.5)
    add_body(doc, "Duration 151,06s so với 86,01s ở customer và 5,61s so với 21,47s ở stock chỉ là số đo mô tả của lần chạy CPU hiện tại. Khác biệt framework implementation, batch loop, callback và số epoch đều ảnh hưởng thời gian. Ít dòng code hơn cũng không đồng nghĩa predictive quality tốt hơn.")

    # VIII
    heading(doc, 1, "VIII. HẠN CHẾ VÀ HƯỚNG PHÁT TRIỂN", new_page=True)
    heading(doc, 2, "8.1. Hạn chế của nghiên cứu")
    add_numbered(doc, [
        "Vanilla RNN dễ gặp vanishing gradient và có năng lực lưu phụ thuộc dài hạn hạn chế; 8 tuần/30 ngày vẫn là một cửa sổ hữu hạn.",
        "Customer target đơn giản hóa hành vi thành có/không mua. Nó chưa phân biệt giá trị đơn hàng, loại sản phẩm, vòng đời khách hàng hay churn.",
        "Mất cân bằng customer mạnh làm precision thấp khi tăng độ nhạy cho class 1; một threshold 0,5 chưa tối ưu cho mọi chi phí kinh doanh.",
        "Stock price level là chuỗi non-stationary, chịu sự kiện ngoại sinh, tin tức và regime change không có trong năm OHLCV features.",
        "Mỗi baseline dùng một chronological split và một seed; kết quả không đại diện cho mọi giai đoạn thị trường hoặc mọi khởi tạo.",
        "CPU-only duration không thể dùng làm kết luận tổng quát về tốc độ framework trên GPU/phần cứng khác.",
    ])
    heading(doc, 2, "8.2. Hướng phát triển hợp lý")
    add_bullets(doc, [
        "Thử LSTM và GRU trên đúng split/artifact hiện tại, giữ nguyên validation protocol để đánh giá lợi ích gating.",
        "Dùng rolling-origin evaluation hoặc nhiều chronological folds thay vì chỉ một cut point.",
        "Customer: chọn threshold trên validation theo chi phí false positive/false negative; xem thêm PR-AUC và calibration.",
        "Customer: bổ sung đặc trưng hợp lệ theo thời gian như recency, tenure, category mix nhưng phải tính chỉ từ lịch sử trước target.",
        "Stock: cân nhắc target return/difference thay vì price level, thêm exogenous features có timestamp rõ và so với các baseline mạnh hơn.",
        "Đánh giá bất định và prediction interval; tránh diễn giải point forecast như khuyến nghị đầu tư.",
    ])
    add_callout(doc, "Nguyên tắc mở rộng", "Mọi cải tiến phải giữ test set ngoài vòng tuning, fit preprocessing trên TRAIN và tái so sánh cùng naive baseline. Kiến trúc phức tạp hơn chỉ có ý nghĩa khi protocol vẫn công bằng.", PALE_GREEN)

    # IX
    heading(doc, 1, "IX. KẾT LUẬN", new_page=True)
    add_body(doc, "A06 đã hoàn thành một pipeline RNN có thể kiểm chứng từ lý thuyết đến thực nghiệm. Phần customer cho thấy hai Vanilla RNN ở hai framework tạo behavior gần như tương đương trên dữ liệu mất cân bằng: ROC-AUC khoảng 0,70, recall khoảng 0,57 nhưng precision chỉ khoảng 0,16. Bài học chính là phải đánh giá nhiều metric và hiểu confusion matrix, không dựa vào accuracy.")
    add_body(doc, "Phần stock cho thấy một kết quả âm nhưng quan trọng: cả PyTorch và Keras RNN đều thua rõ naive last-Close baseline trên đúng 410 TEST dates. Kết luận này được giữ nguyên thay vì tìm cách hợp thức hóa RNN; nó minh họa vai trò bắt buộc của baseline và real-scale evaluation.")
    add_body(doc, "PyTorch làm rõ cơ chế training ở cấp thấp qua DataLoader, forward/backward, optimizer step và device handling. Keras diễn đạt cùng ý tưởng ở cấp API cao hơn qua SimpleRNN, model.fit() và callbacks. Sự khác nhau về cách viết và runtime không đồng nghĩa một framework có năng lực dự báo bẩm sinh tốt hơn.")
    add_body(doc, "Về phương pháp, thành quả lớn nhất là contract chống leakage: window chỉ dùng quá khứ, split theo target time, scaler/class weight chỉ học từ TRAIN, model selection dựa validation và hai framework dùng cùng artifact. Đây là điều kiện để mọi so sánh mô hình sau này có ý nghĩa.")
    add_callout(doc, "Knowledge checklist", "✓ Hiểu hidden state và weight sharing; ✓ phân biệt classification/regression; ✓ tạo many-to-one sequence đúng target; ✓ chia thời gian và scale TRAIN-only; ✓ dùng early stopping; ✓ đọc Precision/Recall/F1/ROC-AUC và MAE/RMSE/R²; ✓ luôn so với baseline; ✓ không biến dự báo stock thành khuyến nghị đầu tư.", PALE_BLUE)

    # References
    heading(doc, 1, "TÀI LIỆU THAM KHẢO", new_page=True)
    references = [
        ("[1]", "D. Chen, “Online Retail II,” UCI Machine Learning Repository, 2012. DOI: 10.24432/C5CG6D.", "https://archive.ics.uci.edu/dataset/502/online+retail+ii"),
        ("[2]", "Yahoo Finance, “Apple Inc. (AAPL) Historical Data.” Snapshot A06 được tải một lần và khóa cục bộ.", "https://finance.yahoo.com/quote/AAPL/history/"),
        ("[3]", "yfinance documentation, “yfinance.download.”", "https://ranaroussi.github.io/yfinance/reference/api/yfinance.download.html"),
        ("[4]", "PyTorch documentation, “torch.nn.RNN.”", "https://docs.pytorch.org/docs/stable/generated/torch.nn.RNN.html"),
        ("[5]", "Keras 3 documentation, “SimpleRNN layer.”", "https://keras.io/api/layers/recurrent_layers/simple_rnn/"),
        ("[6]", "D. E. Rumelhart, G. E. Hinton, and R. J. Williams, “Learning representations by back-propagating errors,” Nature, vol. 323, pp. 533–536, 1986.", "https://doi.org/10.1038/323533a0"),
        ("[7]", "Y. Bengio, P. Simard, and P. Frasconi, “Learning long-term dependencies with gradient descent is difficult,” IEEE Trans. Neural Networks, vol. 5, no. 2, pp. 157–166, 1994.", "https://doi.org/10.1109/72.279181"),
        ("[8]", "S. Hochreiter and J. Schmidhuber, “Long Short-Term Memory,” Neural Computation, vol. 9, no. 8, pp. 1735–1780, 1997.", "https://doi.org/10.1162/neco.1997.9.8.1735"),
        ("[9]", "K. Cho et al., “Learning Phrase Representations using RNN Encoder–Decoder for Statistical Machine Translation,” EMNLP, pp. 1724–1734, 2014.", "https://aclanthology.org/D14-1179/"),
        ("[10]", "A06 PROJECT_SPEC.md, CURRENT_STATE.md, preprocessing_metadata.json, bốn metrics JSON và A06_AUDIT_REPORT.md (artifact nội bộ đã audit).", ""),
    ]
    for label, citation, url in references:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.left_indent = Cm(0.75)
        p.paragraph_format.first_line_indent = Cm(-0.75)
        p.paragraph_format.space_after = Pt(5)
        r = p.add_run(f"{label} {citation} ")
        set_run_font(r, size=10.5)
        if url:
            add_hyperlink(p, "Nguồn trực tuyến", url)

    doc.core_properties.title = "A06 Assignment Report — Recurrent Neural Network"
    doc.core_properties.subject = "RNN customer classification and AAPL stock regression"
    doc.core_properties.author = "A06 Project"
    doc.core_properties.keywords = "RNN, PyTorch, Keras, Online Retail II, AAPL"
    doc.core_properties.comments = "Generated only from audited A06 artifacts; no retraining performed."
    doc.save(OUTPUT_PATH)
    validate_docx(OUTPUT_PATH, ptc, kc, pts, ks)
    print(f"Created: {OUTPUT_PATH}")


def validate_docx(path: Path, ptc: dict, kc: dict, pts: dict, ks: dict) -> None:
    if not path.is_file() or path.stat().st_size < 500_000:
        raise RuntimeError("DOCX missing or unexpectedly small")
    with zipfile.ZipFile(path) as archive:
        names = set(archive.namelist())
        if "word/document.xml" not in names:
            raise RuntimeError("Invalid DOCX package")
        xml = archive.read("word/document.xml").decode("utf-8")
        media = [name for name in names if name.startswith("word/media/")]
    plain = re.sub(r"<[^>]+>", " ", xml)
    for required in [
        "I. TỔNG QUAN", "II. CƠ SỞ LÝ THUYẾT VỀ RNN", "III. PHÂN TÍCH DỮ LIỆU",
        "IV. TIỀN XỬ LÝ", "V. THỰC NGHIỆM PYTORCH", "VI. THỰC NGHIỆM KERAS",
        "VII. SO SÁNH HAI FRAMEWORK", "VIII. HẠN CHẾ", "IX. KẾT LUẬN",
        "TÀI LIỆU THAM KHẢO",
    ]:
        if required not in plain:
            raise RuntimeError(f"Missing section: {required}")
    for forbidden in ["TODO", "FIXME", "Lorem ipsum", "[[TOC]]", "TBD", "PLACEHOLDER"]:
        if forbidden in plain:
            raise RuntimeError(f"Placeholder leaked into report: {forbidden}")
    if len(media) != 46:
        # 24 saved analytical PNGs + 20 rendered tables + 2 report diagrams.
        raise RuntimeError(f"Unexpected embedded image count: {len(media)}")
    expected_metric_strings = [
        fmt4(ptc['evaluation']['metrics']['roc_auc']), fmt4(kc['evaluation']['metrics']['roc_auc']),
        fmt4(pts['MAE']), fmt4(pts['RMSE']), fmt4(pts['R2']),
        fmt4(ks['MAE']), fmt4(ks['RMSE']), fmt4(ks['R2']),
        fmt4(pts['naive_MAE']), fmt4(pts['naive_RMSE']), fmt4(pts['naive_R2']),
    ]
    normalized_plain = plain.replace("−", "-")
    for value in expected_metric_strings:
        if value not in normalized_plain and value.replace(".", ",") not in normalized_plain:
            raise RuntimeError(f"Metric missing from report: {value}")
    reopened = Document(path)
    if len(reopened.paragraphs) < 250 or len(reopened.inline_shapes) != 46:
        raise RuntimeError("DOCX content count is unexpectedly low")


if __name__ == "__main__":
    build_report()
