"""Verify the locked A06 environment, project structure, and raw datasets."""

from __future__ import annotations

import hashlib
import importlib
import json
import platform
import sys
from datetime import datetime
from importlib import metadata
from pathlib import Path
from typing import Any

import openpyxl
import pandas as pd
import tensorflow as tf
import torch


PROJECT_ROOT = Path(__file__).resolve().parents[1]
CUSTOMER_PATH = PROJECT_ROOT / "datasets" / "customer" / "online_retail_II.xlsx"
STOCK_PATH = PROJECT_ROOT / "datasets" / "stock" / "AAPL_2015_2025.csv"
JSON_REPORT_PATH = PROJECT_ROOT / "results" / "metrics" / "environment_verification.json"
MARKDOWN_REPORT_PATH = PROJECT_ROOT / "results" / "metrics" / "environment_verification.md"

REQUIRED_DIRECTORIES = [
    "datasets/customer",
    "datasets/stock",
    "datasets/processed",
    "notebooks",
    "src",
    "tests",
    "models/pytorch",
    "models/keras",
    "results/figures",
    "results/metrics",
    "results/predictions",
    "report",
]

PACKAGES = {
    "numpy": "numpy",
    "pandas": "pandas",
    "matplotlib": "matplotlib",
    "scikit-learn": "scikit-learn",
    "openpyxl": "openpyxl",
    "jupyter": "jupyter",
    "ipykernel": "ipykernel",
    "pytest": "pytest",
    "joblib": "joblib",
    "torch": "torch",
    "tensorflow": "tensorflow",
    "keras": "keras",
    "yfinance": "yfinance",
}

IMPORT_MODULES = {
    "numpy": "numpy",
    "pandas": "pandas",
    "matplotlib": "matplotlib",
    "scikit-learn": "sklearn",
    "openpyxl": "openpyxl",
    "jupyter": "jupyter",
    "ipykernel": "ipykernel",
    "pytest": "pytest",
    "joblib": "joblib",
    "torch": "torch",
    "tensorflow": "tensorflow",
    "keras": "keras",
    "yfinance": "yfinance",
}

EXPECTED_STOCK_COLUMNS = [
    "Date",
    "Open",
    "High",
    "Low",
    "Close",
    "Adj Close",
    "Volume",
]


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def inspect_customer_workbook() -> dict[str, Any]:
    workbook = openpyxl.load_workbook(CUSTOMER_PATH, read_only=True, data_only=True)
    try:
        sheets: list[dict[str, Any]] = []
        for worksheet in workbook.worksheets:
            header = [cell.value for cell in next(worksheet.iter_rows(min_row=1, max_row=1))]
            sample = pd.read_excel(CUSTOMER_PATH, sheet_name=worksheet.title, nrows=200)
            sheets.append(
                {
                    "name": worksheet.title,
                    "data_rows": max(worksheet.max_row - 1, 0),
                    "columns": header,
                    "sample_dtypes": {column: str(dtype) for column, dtype in sample.dtypes.items()},
                }
            )
    finally:
        workbook.close()

    return {
        "path": str(CUSTOMER_PATH),
        "size_bytes": CUSTOMER_PATH.stat().st_size,
        "sha256": sha256(CUSTOMER_PATH),
        "sheet_count": len(sheets),
        "total_data_rows": sum(sheet["data_rows"] for sheet in sheets),
        "sheets": sheets,
    }


def inspect_stock_csv() -> dict[str, Any]:
    stock = pd.read_csv(STOCK_PATH)
    dates = pd.to_datetime(stock["Date"], errors="raise")
    return {
        "path": str(STOCK_PATH),
        "size_bytes": STOCK_PATH.stat().st_size,
        "sha256": sha256(STOCK_PATH),
        "rows": len(stock),
        "columns": stock.columns.tolist(),
        "dtypes": {column: str(dtype) for column, dtype in stock.dtypes.items()},
        "first_trading_date": dates.min().date().isoformat(),
        "last_trading_date": dates.max().date().isoformat(),
        "dates_monotonic_increasing": bool(dates.is_monotonic_increasing),
        "duplicate_dates": int(dates.duplicated().sum()),
        "missing_values": {column: int(value) for column, value in stock.isna().sum().items()},
    }


def build_report() -> dict[str, Any]:
    errors: list[str] = []
    actual_python = str(Path(sys.executable).resolve())
    expected_python = str(
        Path(r"C:\Users\anhca\anaconda3\envs\rnn312\python.exe").resolve()
    )
    if actual_python.casefold() != expected_python.casefold():
        errors.append(f"Unexpected interpreter: {actual_python}")
    if sys.version_info[:2] != (3, 12):
        errors.append(f"Expected Python 3.12, found {platform.python_version()}")

    directory_status = {
        relative: (PROJECT_ROOT / relative).is_dir() for relative in REQUIRED_DIRECTORIES
    }
    errors.extend(
        f"Missing directory: {relative}"
        for relative, exists in directory_status.items()
        if not exists
    )

    package_versions = {
        label: metadata.version(distribution) for label, distribution in PACKAGES.items()
    }
    import_status: dict[str, str] = {}
    for label, module_name in IMPORT_MODULES.items():
        try:
            importlib.import_module(module_name)
            import_status[label] = "PASS"
        except Exception as exc:  # pragma: no cover - used to report environment failures
            import_status[label] = f"FAIL: {type(exc).__name__}: {exc}"
            errors.append(f"Import failed for {label}: {exc}")

    customer: dict[str, Any] | None = None
    if CUSTOMER_PATH.is_file():
        customer = inspect_customer_workbook()
        if customer["sheet_count"] == 0:
            errors.append("Customer workbook has no sheets")
    else:
        errors.append(f"Missing customer workbook: {CUSTOMER_PATH}")

    stock: dict[str, Any] | None = None
    if STOCK_PATH.is_file():
        stock = inspect_stock_csv()
        if stock["columns"] != EXPECTED_STOCK_COLUMNS:
            errors.append(f"Unexpected stock columns: {stock['columns']}")
        if stock["first_trading_date"] < "2015-01-01":
            errors.append("Stock data starts before the locked interval")
        if stock["last_trading_date"] > "2025-12-31":
            errors.append("Stock data ends after the locked interval")
        if not stock["dates_monotonic_increasing"]:
            errors.append("Stock dates are not chronological")
        if stock["duplicate_dates"]:
            errors.append("Stock dates contain duplicates")
    else:
        errors.append(f"Missing stock CSV: {STOCK_PATH}")

    tf_devices = [
        {"name": device.name, "type": device.device_type}
        for device in tf.config.list_physical_devices()
    ]
    torch_cuda_devices = [
        torch.cuda.get_device_name(index) for index in range(torch.cuda.device_count())
    ]

    return {
        "generated_at_local": datetime.now().astimezone().isoformat(timespec="seconds"),
        "status": "PASS" if not errors else "FAIL",
        "errors": errors,
        "project_root": str(PROJECT_ROOT),
        "environment": {
            "name": "rnn312",
            "interpreter": actual_python,
            "python_version": platform.python_version(),
            "platform": platform.platform(),
            "package_versions": package_versions,
            "import_status": import_status,
        },
        "devices": {
            "pytorch_cuda_available": torch.cuda.is_available(),
            "pytorch_cuda_devices": torch_cuda_devices,
            "tensorflow_physical_devices": tf_devices,
        },
        "directories": directory_status,
        "datasets": {"customer": customer, "stock": stock},
    }


def render_markdown(report: dict[str, Any]) -> str:
    environment = report["environment"]
    devices = report["devices"]
    customer = report["datasets"]["customer"]
    stock = report["datasets"]["stock"]
    lines = [
        "# A06 Environment Verification",
        "",
        f"- Status: **{report['status']}**",
        f"- Generated: `{report['generated_at_local']}`",
        f"- Interpreter: `{environment['interpreter']}`",
        f"- Python: `{environment['python_version']}`",
        f"- Direct package imports: `{'PASS' if all(value == 'PASS' for value in environment['import_status'].values()) else 'FAIL'}`",
        f"- PyTorch CUDA available: `{devices['pytorch_cuda_available']}`",
        f"- TensorFlow physical devices: `{devices['tensorflow_physical_devices']}`",
        "",
        "## Required package versions",
        "",
    ]
    lines.extend(
        f"- `{name}=={version}`"
        for name, version in environment["package_versions"].items()
    )
    lines.extend(["", "## Dataset checks", ""])
    if customer:
        lines.append(
            f"- Customer workbook: `{customer['sheet_count']}` sheets, "
            f"`{customer['total_data_rows']}` data rows, SHA-256 "
            f"`{customer['sha256']}`."
        )
        for sheet in customer["sheets"]:
            lines.append(
                f"  - `{sheet['name']}`: `{sheet['data_rows']}` rows; columns: "
                + ", ".join(f"`{column}`" for column in sheet["columns"])
            )
    if stock:
        lines.append(
            f"- Stock CSV: `{stock['rows']}` rows, "
            f"`{stock['first_trading_date']}` through `{stock['last_trading_date']}`, "
            f"SHA-256 `{stock['sha256']}`."
        )
        lines.append(
            "  - Columns: " + ", ".join(f"`{column}`" for column in stock["columns"])
        )
        lines.append(
            f"  - Chronological: `{stock['dates_monotonic_increasing']}`; "
            f"duplicate dates: `{stock['duplicate_dates']}`."
        )
    if report["errors"]:
        lines.extend(["", "## Errors", ""])
        lines.extend(f"- {error}" for error in report["errors"])
    return "\n".join(lines) + "\n"


def main() -> None:
    report = build_report()
    JSON_REPORT_PATH.parent.mkdir(parents=True, exist_ok=True)
    JSON_REPORT_PATH.write_text(
        json.dumps(report, indent=2, ensure_ascii=False), encoding="utf-8"
    )
    MARKDOWN_REPORT_PATH.write_text(render_markdown(report), encoding="utf-8")
    print(json.dumps(report, indent=2, ensure_ascii=False))
    if report["status"] != "PASS":
        raise SystemExit(1)


if __name__ == "__main__":
    main()
