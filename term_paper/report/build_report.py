"""Build the Phase 8 DOCX from the reviewed Phase 7 manuscript.

The document is deliberately generated with native Word paragraphs, headings,
tables, captions, hyperlinks, bookmarks, page fields, and editable code blocks.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import shutil
import subprocess
import sys
import zipfile
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from PIL import Image
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Cm, Inches, Pt, RGBColor


WORKSPACE = Path(__file__).resolve().parents[2]
TERM_ROOT = WORKSPACE / "term_paper"
REPORT_ROOT = TERM_ROOT / "report"
MANUSCRIPT_PATH = REPORT_ROOT / "manuscript.md"
REPORT_MANIFEST_PATH = REPORT_ROOT / "report_manifest.json"
DEFAULT_OUTPUT = REPORT_ROOT / "Tieu_luan_AI_ML_CNN_RNN.docx"
DEFAULT_LOG = REPORT_ROOT / "build_log.json"
LOGO_PATH = REPORT_ROOT / "assets" / "branding" / "ptit_logo.png"

RUNTIME_ROOT = Path.home() / ".cache" / "codex-runtimes" / "codex-primary-runtime" / "dependencies"
RUNTIME_NODE = RUNTIME_ROOT / "node" / "bin" / "node.exe"
RUNTIME_NODE_MODULES = RUNTIME_ROOT / "node" / "node_modules"

NAVY = "17365D"
BLUE = "2F75B5"
LIGHT_BLUE = "D9EAF7"
PALE_BLUE = "F4F8FC"
BODY = "1F2933"
MUTED = "5B6570"
GRID = "9FB7C9"
CODE_BG = "F3F6F8"
CODE_BORDER = "D9E2E8"

METADATA = {
    "school": "HỌC VIỆN CÔNG NGHỆ BƯU CHÍNH VIỄN THÔNG",
    "faculty": "KHOA CÔNG NGHỆ THÔNG TIN",
    "course": "PHÁT TRIỂN CÁC HỆ THỐNG THÔNG MINH",
    "document_type": "TIỂU LUẬN MÔN HỌC",
    "title": "LỊCH SỬ PHÁT TRIỂN AI VÀ THỰC NGHIỆM ML, CNN, RNN",
    "class": "D23CTPM01",
    "student_name": "Triệu Tuấn Anh",
    "student_id": "B23DCCN053",
    "instructor": "PGS.TS Trần Đình Quế",
    "semester": "Học kỳ 1 năm học 2026 – 2027",
    "location_year": "HÀ NỘI — 2026",
}


@dataclass
class Block:
    kind: str
    data: Any
    line: int


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def strip_markdown(text: str) -> str:
    text = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1", text)
    text = text.replace("**", "").replace("*", "").replace("`", "")
    return re.sub(r"\s+", " ", text).strip()


def caption_parts(text: str) -> tuple[str, str, str] | None:
    plain = strip_markdown(text)
    match = re.match(r"^(Hình|Bảng)\s+((?:\d+|K)\.\d+)\.?(?:\s+)(.+)$", plain)
    if not match:
        return None
    label, number, title = match.groups()
    return label, number, title.strip().rstrip(".")


def is_caption_paragraph(text: str) -> bool:
    """Return true only for deliberately emphasized Markdown caption lines."""
    return bool(re.match(r"^\*{1,2}(?:Hình|Bảng)\s+", text.strip()))


def parse_markdown(text: str) -> list[Block]:
    lines = text.splitlines()
    blocks: list[Block] = []
    index = 0

    def is_boundary(line: str, next_line: str = "") -> bool:
        stripped = line.strip()
        return bool(
            not stripped
            or stripped.startswith("#")
            or stripped.startswith("<!--")
            or stripped.startswith("```")
            or stripped == r"\["
            or re.match(r"^!\[.*\]\(.*\)$", stripped)
            or re.match(r"^(?:[-*]|\d+\.)\s+", stripped)
            or (stripped.startswith("|") and re.match(r"^\|?\s*:?-+", next_line.strip()))
        )

    while index < len(lines):
        raw = lines[index]
        stripped = raw.strip()
        if not stripped or stripped.startswith("<!--"):
            index += 1
            continue
        heading = re.match(r"^(#{1,3})\s+(.+)$", stripped)
        if heading:
            blocks.append(Block("heading", (len(heading.group(1)), heading.group(2).strip()), index + 1))
            index += 1
            continue
        if stripped.startswith("```"):
            language = stripped[3:].strip()
            start = index + 1
            code_lines: list[str] = []
            index += 1
            while index < len(lines) and not lines[index].strip().startswith("```"):
                code_lines.append(lines[index])
                index += 1
            if index >= len(lines):
                raise ValueError(f"Unclosed code fence at line {start}")
            index += 1
            blocks.append(Block("code", (language, "\n".join(code_lines)), start))
            continue
        if stripped == r"\[":
            start = index + 1
            equation: list[str] = []
            index += 1
            while index < len(lines) and lines[index].strip() != r"\]":
                equation.append(lines[index].strip())
                index += 1
            if index >= len(lines):
                raise ValueError(f"Unclosed equation at line {start}")
            index += 1
            blocks.append(Block("equation", " ".join(equation), start))
            continue
        image_match = re.match(r"^!\[([^\]]*)\]\(([^)]+)\)$", stripped)
        if image_match:
            blocks.append(Block("image", image_match.groups(), index + 1))
            index += 1
            continue
        if stripped.startswith("|") and index + 1 < len(lines) and re.match(
            r"^\|?\s*:?-+", lines[index + 1].strip()
        ):
            start = index + 1
            table_lines = [stripped]
            index += 2  # skip alignment separator
            while index < len(lines) and lines[index].strip().startswith("|"):
                table_lines.append(lines[index].strip())
                index += 1
            rows = [
                [cell.strip() for cell in row.strip().strip("|").split("|")]
                for row in table_lines
            ]
            blocks.append(Block("table", rows, start))
            continue
        list_match = re.match(r"^([-*]|\d+\.)\s+(.+)$", stripped)
        if list_match:
            marker, content = list_match.groups()
            blocks.append(Block("list_item", ("bullet" if marker in {"-", "*"} else "number", marker, content), index + 1))
            index += 1
            continue

        start = index + 1
        parts = [stripped]
        index += 1
        while index < len(lines):
            next_line = lines[index]
            after = lines[index + 1] if index + 1 < len(lines) else ""
            if is_boundary(next_line, after):
                break
            parts.append(next_line.strip())
            index += 1
        blocks.append(Block("paragraph", " ".join(parts), start))
    return blocks


def set_run_font(run, name: str, size: float | None = None, color: str | None = None) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def set_style_font(style, name: str, size: float, color: str = BODY, bold: bool | None = None) -> None:
    style.font.name = name
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        style.font.bold = bold
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    for key in ("w:ascii", "w:hAnsi", "w:eastAsia"):
        rfonts.set(qn(key), name)
    lang = OxmlElement("w:lang")
    lang.set(qn("w:val"), "vi-VN")
    lang.set(qn("w:eastAsia"), "vi-VN")
    rpr.append(lang)


def first_child(parent, *tags: str):
    """Return the first matching child without relying on lxml element truthiness."""
    for tag in tags:
        child = parent.find(qn(tag))
        if child is not None:
            return child
    return None


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        # w:shd precedes w:tcMar and w:vAlign in CT_TcPr.
        before = first_child(tc_pr, "w:tcMar", "w:vAlign")
        if before is None:
            tc_pr.append(shd)
        else:
            before.addprevious(shd)
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top: int = 80, start: int = 100, bottom: int = 80, end: int = 100) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        # w:tcMar must precede w:textDirection/w:tcFitText/w:vAlign.
        before = first_child(tc_pr, "w:textDirection", "w:tcFitText", "w:vAlign")
        if before is None:
            tc_pr.append(tc_mar)
        else:
            before.addprevious(tc_mar)
    # left/right are supported by all WordprocessingML versions used by Word.
    for margin, value in (("top", top), ("left", start), ("bottom", bottom), ("right", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tr_pr.append(tbl_header)


def set_row_cant_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_table_fixed_widths(table, widths_cm: list[float]) -> None:
    """Set tblW, tblGrid and every tcW so renderers cannot collapse columns."""
    if len(table.columns) != len(widths_cm):
        raise ValueError("Column width count does not match table")
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        # tblW follows tblStyle/tblpPr/tblOverlap and precedes jc/tblLayout.
        before = first_child(tbl_pr, "w:jc", "w:tblLayout", "w:tblLook")
        if before is None:
            tbl_pr.append(tbl_w)
        else:
            before.addprevious(tbl_w)
    total_twips = sum(Cm(width).twips for width in widths_cm)
    tbl_w.set(qn("w:w"), str(total_twips))
    tbl_w.set(qn("w:type"), "dxa")
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        before = first_child(tbl_pr, "w:tblCellMar", "w:tblLook")
        if before is None:
            tbl_pr.append(layout)
        else:
            before.addprevious(layout)
    layout.set(qn("w:type"), "fixed")

    grid_columns = list(table._tbl.tblGrid.iterchildren(qn("w:gridCol")))
    for column_index, width_cm in enumerate(widths_cm):
        width = Cm(width_cm)
        table.columns[column_index].width = width
        if column_index < len(grid_columns):
            grid_columns[column_index].set(qn("w:w"), str(width.twips))
        for cell in table.columns[column_index].cells:
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width.twips))
            tc_w.set(qn("w:type"), "dxa")


def add_bookmark(paragraph, name: str, bookmark_id: int) -> None:
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(bookmark_id))
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(bookmark_id))
    # In WordprocessingML, w:pPr must remain the first child of w:p.  Placing a
    # bookmark before it is accepted by some importers but Microsoft Word
    # rejects the package as malformed.
    insert_at = 1 if paragraph._p.pPr is not None else 0
    paragraph._p.insert(insert_at, start)
    paragraph._p.append(end)


def add_hyperlink(paragraph, text: str, url: str | None = None, anchor: str | None = None, *, color: str = BLUE, underline: bool = True, bold: bool = False) -> None:
    hyperlink = OxmlElement("w:hyperlink")
    if url:
        relationship_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
        hyperlink.set(qn("r:id"), relationship_id)
    elif anchor:
        hyperlink.set(qn("w:anchor"), anchor)
        hyperlink.set(qn("w:history"), "1")
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    rfonts = OxmlElement("w:rFonts")
    for key in ("w:ascii", "w:hAnsi", "w:eastAsia"):
        rfonts.set(qn(key), "Times New Roman")
    rpr.append(rfonts)
    if bold:
        rpr.append(OxmlElement("w:b"))
    if color:
        node = OxmlElement("w:color")
        node.set(qn("w:val"), color)
        rpr.append(node)
    if underline:
        node = OxmlElement("w:u")
        node.set(qn("w:val"), "single")
        rpr.append(node)
    run.append(rpr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_page_field(paragraph) -> None:
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    display_text = OxmlElement("w:t")
    display_text.text = "1"
    for node in (begin, instr, separate, display_text, end):
        field_run = OxmlElement("w:r")
        field_run.append(node)
        paragraph._p.append(field_run)


def add_paragraph_border(paragraph, side: str, color: str, size: str = "8", space: str = "3") -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        # pBdr precedes spacing/ind/jc in CT_PPr.
        before = first_child(p_pr, "w:shd", "w:tabs", "w:spacing", "w:ind", "w:jc")
        if before is None:
            p_pr.append(borders)
        else:
            before.addprevious(borders)
    border = OxmlElement(f"w:{side}")
    border.set(qn("w:val"), "single")
    border.set(qn("w:sz"), size)
    border.set(qn("w:space"), space)
    border.set(qn("w:color"), color)
    borders.append(border)


def latex_to_text(text: str) -> str:
    replacements = {
        r"\mathbb{1}": "𝟙",
        r"\hat{y}": "ŷ",
        r"\sigma": "σ",
        r"\mu": "μ",
        r"\epsilon": "ε",
        r"\theta": "θ",
        r"\delta": "δ",
        r"\eta": "η",
        r"\odot": "⊙",
        r"\geq": "≥",
        r"\leq": "≤",
        r"\approx": "≈",
        r"\leftarrow": "←",
        r"\sum": "Σ",
        r"\min": "min",
        r"\max": "max",
        r"\log": "log",
        r"\exp": "exp",
        r"\lVert": "‖",
        r"\rVert": "‖",
        r"\quad": "  ",
        r"\qquad": "    ",
        r"\,": " ",
        r"\;": " ",
        r"\left": "",
        r"\right": "",
        r"\tanh": "tanh",
        r"\operatorname": "",
    }
    text = text.replace(r"\begin{aligned}", "").replace(r"\end{aligned}", "")
    text = text.replace(r"\\", " ; ")
    for old, new in replacements.items():
        text = text.replace(old, new)
    fraction = re.compile(r"\\frac\{([^{}]+)\}\{([^{}]+)\}")
    while fraction.search(text):
        text = fraction.sub(r"(\1)/(\2)", text)
    text = re.sub(r"\^\{([^{}]+)\}", r"^(\1)", text)
    text = re.sub(r"_\{([^{}]+)\}", r"_\1", text)
    text = text.replace("{", "").replace("}", "")
    return re.sub(r"\s+", " ", text).strip()


INLINE_PATTERN = re.compile(
    r"(\*\*.+?\*\*|\*[^*]+?\*|`[^`]+?`|\[[^\]]+\]\([^)]+\)|\\\(.+?\\\)|https?://[^\s]+)"
)


def add_inline_runs(paragraph, text: str, *, base_size: float | None = None) -> None:
    position = 0
    for match in INLINE_PATTERN.finditer(text):
        if match.start() > position:
            run = paragraph.add_run(text[position : match.start()])
            set_run_font(run, "Times New Roman", base_size)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
            set_run_font(run, "Times New Roman", base_size)
        elif token.startswith("*"):
            run = paragraph.add_run(token[1:-1])
            run.italic = True
            set_run_font(run, "Times New Roman", base_size)
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, "Consolas", (base_size - 0.5) if base_size else 10.5, NAVY)
        elif token.startswith("["):
            link_match = re.match(r"\[([^\]]+)\]\(([^)]+)\)", token)
            assert link_match
            add_hyperlink(paragraph, link_match.group(1), url=link_match.group(2))
        elif token.startswith(r"\("):
            run = paragraph.add_run(latex_to_text(token[2:-2]))
            run.italic = True
            set_run_font(run, "Cambria Math", base_size)
        elif token.startswith("http"):
            trailing = ""
            url = token
            while url and url[-1] in ".,;":
                trailing = url[-1] + trailing
                url = url[:-1]
            add_hyperlink(paragraph, url, url=url)
            if trailing:
                run = paragraph.add_run(trailing)
                set_run_font(run, "Times New Roman", base_size)
        position = match.end()
    if position < len(text):
        run = paragraph.add_run(text[position:])
        set_run_font(run, "Times New Roman", base_size)


class ReportBuilder:
    def __init__(self, page_map_path: Path | None = None) -> None:
        self.manuscript = MANUSCRIPT_PATH.read_text(encoding="utf-8")
        self.report_manifest = json.loads(REPORT_MANIFEST_PATH.read_text(encoding="utf-8"))
        self.blocks = parse_markdown(self.manuscript)
        self.page_map: dict[str, int] = {}
        if page_map_path and page_map_path.is_file():
            self.page_map = json.loads(page_map_path.read_text(encoding="utf-8")).get("pages", {})
        self.doc = Document()
        self.bookmark_id = 1
        self.navigation_targets: list[dict[str, Any]] = []
        self.nav_by_block: dict[int, dict[str, Any]] = {}
        self.counts = {"headings": 0, "figures": 0, "tables": 0, "code_blocks": 0, "equations": 0}
        self.after_heading = False
        self.pending_table_caption = False
        self._plan_navigation()
        self._configure_document()

    def _plan_navigation(self) -> None:
        heading_count = figure_count = table_count = 0
        for block_index, block in enumerate(self.blocks):
            target: dict[str, Any] | None = None
            if block.kind == "heading":
                level, title = block.data
                heading_count += 1
                target = {
                    "bookmark": f"heading_{heading_count:03d}",
                    "display": strip_markdown(title),
                    "kind": "heading",
                    "level": level,
                }
            elif block.kind == "image":
                alt, _ = block.data
                parts = caption_parts(alt)
                if parts and parts[0] == "Hình":
                    figure_count += 1
                    target = {
                        "bookmark": f"figure_{parts[1].replace('.', '_')}",
                        "display": f"Hình {parts[1]}. {parts[2]}",
                        "kind": "figure",
                        "level": 0,
                    }
            elif block.kind == "paragraph" and is_caption_paragraph(block.data):
                parts = caption_parts(block.data)
                if parts:
                    label, number, title = parts
                    if label == "Hình":
                        figure_count += 1
                        bookmark = f"figure_{number.replace('.', '_')}"
                        kind = "figure"
                    else:
                        table_count += 1
                        bookmark = f"table_{number.replace('.', '_')}"
                        kind = "table"
                    target = {
                        "bookmark": bookmark,
                        "display": f"{label} {number}. {title}",
                        "kind": kind,
                        "level": 0,
                    }
            if target:
                self.navigation_targets.append(target)
                self.nav_by_block[block_index] = target

        figure_ids = [item["bookmark"] for item in self.navigation_targets if item["kind"] == "figure"]
        table_ids = [item["bookmark"] for item in self.navigation_targets if item["kind"] == "table"]
        if len(figure_ids) != 18 or len(set(figure_ids)) != 18:
            raise ValueError(f"Expected 18 unique figures, got {len(figure_ids)}/{len(set(figure_ids))}")
        if len(table_ids) != 16 or len(set(table_ids)) != 16:
            raise ValueError(f"Expected 16 unique tables, got {len(table_ids)}/{len(set(table_ids))}")

    def _configure_document(self) -> None:
        doc = self.doc
        core = doc.core_properties
        core.title = METADATA["title"]
        core.subject = "Tiểu luận môn Phát triển các hệ thống thông minh"
        core.author = METADATA["student_name"]
        core.keywords = "AI, machine learning, CNN, RNN, NumPy, Keras, PyTorch"
        core.comments = "Generated deterministically from term_paper/report/manuscript.md"

        normal = doc.styles["Normal"]
        set_style_font(normal, "Times New Roman", 11.5, BODY)
        normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
        normal.paragraph_format.space_after = Pt(5)
        normal.paragraph_format.first_line_indent = Cm(0.75)
        normal.paragraph_format.widow_control = True

        title = doc.styles["Title"]
        set_style_font(title, "Times New Roman", 21, NAVY, True)
        title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        title.paragraph_format.space_after = Pt(12)

        heading_specs = {
            "Heading 1": (15, NAVY, 10, 7),
            "Heading 2": (13, BLUE, 8, 5),
            "Heading 3": (11.5, NAVY, 6, 3),
        }
        for name, (size, color, before, after) in heading_specs.items():
            style = doc.styles[name]
            set_style_font(style, "Times New Roman", size, color, True)
            style.paragraph_format.space_before = Pt(before)
            style.paragraph_format.space_after = Pt(after)
            style.paragraph_format.keep_with_next = True
            style.paragraph_format.keep_together = True
            style.paragraph_format.first_line_indent = Cm(0)
        doc.styles["Heading 1"].paragraph_format.page_break_before = True

        caption = doc.styles["Caption"]
        set_style_font(caption, "Times New Roman", 10, BODY)
        caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption.paragraph_format.space_before = Pt(4)
        caption.paragraph_format.space_after = Pt(6)
        caption.paragraph_format.keep_with_next = True
        caption.paragraph_format.first_line_indent = Cm(0)

        self._ensure_custom_style("Source Note", 9, MUTED, italic=True)
        self._ensure_custom_style("Reference Entry", 10, BODY)
        self.doc.styles["Reference Entry"].paragraph_format.left_indent = Cm(0.75)
        self.doc.styles["Reference Entry"].paragraph_format.first_line_indent = Cm(-0.75)
        self.doc.styles["Reference Entry"].paragraph_format.line_spacing = 1.1
        self.doc.styles["Reference Entry"].paragraph_format.space_after = Pt(4)

        for level in (1, 2, 3):
            style = self._ensure_custom_style(f"Static TOC {level}", 10.5 if level == 1 else 10, NAVY if level == 1 else BODY, bold=(level == 1))
            style.paragraph_format.left_indent = Cm((level - 1) * 0.55)
            style.paragraph_format.first_line_indent = Cm(0)
            style.paragraph_format.space_after = Pt(2)
            style.paragraph_format.line_spacing = 1.0

        section = doc.sections[0]
        self._configure_section_geometry(section)
        section.header.is_linked_to_previous = False
        section.footer.is_linked_to_previous = False

    def _ensure_custom_style(self, name: str, size: float, color: str, *, italic: bool = False, bold: bool = False):
        if name in [style.name for style in self.doc.styles]:
            style = self.doc.styles[name]
        else:
            style = self.doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        set_style_font(style, "Times New Roman", size, color, bold)
        style.font.italic = italic
        style.paragraph_format.first_line_indent = Cm(0)
        style.paragraph_format.space_after = Pt(3)
        return style

    @staticmethod
    def _configure_section_geometry(section) -> None:
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(2.25)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(2.7)
        section.right_margin = Cm(2.2)
        section.header_distance = Cm(0.9)
        section.footer_distance = Cm(0.8)

    def _set_header_footer(self, section, *, show: bool) -> None:
        section.header.is_linked_to_previous = False
        section.footer.is_linked_to_previous = False
        header = section.header
        footer = section.footer
        header_p = header.paragraphs[0]
        footer_p = footer.paragraphs[0]
        header_p.clear()
        footer_p.clear()
        if not show:
            return
        header_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        header_p.paragraph_format.space_after = Pt(0)
        run = header_p.add_run("TIỂU LUẬN MÔN HỌC  •  AI / ML / CNN / RNN")
        set_run_font(run, "Times New Roman", 9, MUTED)
        run.bold = True
        add_paragraph_border(header_p, "bottom", "B8C7D1", "5", "3")

        footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        footer_p.paragraph_format.tab_stops.add_tab_stop(Cm(15.5), WD_TAB_ALIGNMENT.RIGHT)
        run = footer_p.add_run("Triệu Tuấn Anh • B23DCCN053")
        set_run_font(run, "Times New Roman", 9, MUTED)
        footer_p.add_run("\tTrang ")
        add_page_field(footer_p)

    def _ensure_assets(self) -> None:
        LOGO_PATH.parent.mkdir(parents=True, exist_ok=True)
        if not LOGO_PATH.is_file():
            source_logo = WORKSPACE / "A06" / "report" / "docx_assets" / "ptit_logo.png"
            if source_logo.is_file():
                shutil.copy2(source_logo, LOGO_PATH)
            else:
                sample = Path.home() / "Documents" / "Intel_sys_dev" / "A06_Assignment_Report.docx"
                if not sample.is_file():
                    raise FileNotFoundError("PTIT logo source is unavailable")
                with zipfile.ZipFile(sample) as archive:
                    LOGO_PATH.write_bytes(archive.read("word/media/image12.png"))

        svg = REPORT_ROOT / "assets" / "ch1" / "ai_history_timeline.svg"
        png = REPORT_ROOT / "assets" / "ch1" / "ai_history_timeline.png"
        if svg.is_file() and (not png.is_file() or svg.stat().st_mtime_ns > png.stat().st_mtime_ns):
            if not RUNTIME_NODE.is_file():
                raise FileNotFoundError(RUNTIME_NODE)
            subprocess.run(
                [
                    str(RUNTIME_NODE),
                    str(REPORT_ROOT / "convert_svg_to_png.mjs"),
                    str(svg),
                    str(png),
                    str(RUNTIME_NODE_MODULES),
                ],
                check=True,
            )

    def _centered_line(self, text: str, size: float, *, bold: bool = False, color: str = BODY, before: float = 0, after: float = 0) -> None:
        paragraph = self.doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.first_line_indent = Cm(0)
        paragraph.paragraph_format.space_before = Pt(before)
        paragraph.paragraph_format.space_after = Pt(after)
        run = paragraph.add_run(text)
        set_run_font(run, "Times New Roman", size, color)
        run.bold = bold

    def _add_label_value(self, label: str, value: str) -> None:
        paragraph = self.doc.add_paragraph()
        paragraph.paragraph_format.left_indent = Cm(3.25)
        paragraph.paragraph_format.first_line_indent = Cm(0)
        paragraph.paragraph_format.space_after = Pt(5)
        label_run = paragraph.add_run(f"{label}: ")
        set_run_font(label_run, "Times New Roman", 12, BODY)
        label_run.bold = True
        value_run = paragraph.add_run(value)
        set_run_font(value_run, "Times New Roman", 12, BODY)

    def add_cover_and_title_page(self) -> None:
        self._ensure_assets()
        self._set_header_footer(self.doc.sections[0], show=False)
        self._centered_line(METADATA["school"], 15, bold=True, after=2)
        self._centered_line(METADATA["faculty"], 14, bold=True, after=10)
        logo_p = self.doc.add_paragraph()
        logo_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        logo_p.paragraph_format.first_line_indent = Cm(0)
        logo_shape = logo_p.add_run().add_picture(str(LOGO_PATH), height=Cm(3.2))
        logo_shape._inline.docPr.set("descr", "Logo Học viện Công nghệ Bưu chính Viễn thông (PTIT)")
        logo_shape._inline.docPr.set("title", "Logo PTIT")
        self._centered_line(METADATA["document_type"], 17, bold=True, color=NAVY, before=8, after=4)
        self._centered_line(METADATA["course"], 14, bold=True, after=8)
        self._centered_line(METADATA["title"], 19, bold=True, color=NAVY, after=18)
        for label, key in (
            ("Lớp", "class"),
            ("Họ và tên", "student_name"),
            ("Mã sinh viên", "student_id"),
            ("Giảng viên hướng dẫn", "instructor"),
            ("Học kỳ", "semester"),
        ):
            self._add_label_value(label, METADATA[key])
        self._centered_line(METADATA["location_year"], 11, bold=True, color=MUTED, before=22)

        self.doc.add_page_break()
        self._centered_line(METADATA["school"], 14, bold=True, after=2)
        self._centered_line(METADATA["faculty"], 13, bold=True, after=25)
        self._centered_line("BÁO CÁO TIỂU LUẬN", 18, bold=True, color=NAVY, after=8)
        self._centered_line(METADATA["title"], 20, bold=True, color=NAVY, after=24)
        self._add_label_value("Môn học", METADATA["course"].title())
        self._add_label_value("Sinh viên thực hiện", f"{METADATA['student_name']} — {METADATA['student_id']}")
        self._add_label_value("Lớp", METADATA["class"])
        self._add_label_value("Giảng viên", METADATA["instructor"])
        self._add_label_value("Học kỳ", METADATA["semester"])
        statement = self.doc.add_paragraph()
        statement.paragraph_format.left_indent = Cm(2.2)
        statement.paragraph_format.right_indent = Cm(2.2)
        statement.paragraph_format.first_line_indent = Cm(0)
        statement.paragraph_format.space_before = Pt(24)
        statement.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        run = statement.add_run(
            "Báo cáo trình bày lịch sử phát triển trí tuệ nhân tạo, các kỹ thuật học máy cơ bản, "
            "CNN, RNN, thực nghiệm so sánh NumPy scratch–Keras–PyTorch và hệ thống triển khai inference."
        )
        set_run_font(run, "Times New Roman", 11.5, MUTED)
        run.italic = True
        self._centered_line(METADATA["location_year"], 11, bold=True, color=MUTED, before=42)

        front_section = self.doc.add_section(WD_SECTION.NEW_PAGE)
        self._configure_section_geometry(front_section)
        self._set_header_footer(front_section, show=True)

    def _add_front_title(self, text: str) -> None:
        paragraph = self.doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.first_line_indent = Cm(0)
        paragraph.paragraph_format.space_before = Pt(8)
        paragraph.paragraph_format.space_after = Pt(12)
        run = paragraph.add_run(text)
        set_run_font(run, "Times New Roman", 17, NAVY)
        run.bold = True
        add_paragraph_border(paragraph, "bottom", BLUE, "8", "5")

    def _add_static_nav_entry(self, target: dict[str, Any], style_name: str) -> None:
        paragraph = self.doc.add_paragraph(style=style_name)
        paragraph.paragraph_format.tab_stops.add_tab_stop(Cm(15.25), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        add_hyperlink(
            paragraph,
            target["display"],
            anchor=target["bookmark"],
            color=NAVY if target.get("level") == 1 else BODY,
            underline=False,
            bold=target.get("level") == 1,
        )
        paragraph.add_run("\t")
        page = str(self.page_map.get(target["bookmark"], "—"))
        add_hyperlink(paragraph, page, anchor=target["bookmark"], color=BODY, underline=False)

    def add_front_lists(self) -> None:
        self._add_front_title("MỤC LỤC")
        for target in self.navigation_targets:
            if target["kind"] == "heading":
                self._add_static_nav_entry(target, f"Static TOC {target['level']}")

        self.doc.add_page_break()
        self._add_front_title("DANH MỤC HÌNH")
        for target in self.navigation_targets:
            if target["kind"] == "figure":
                self._add_static_nav_entry(target, "Static TOC 2")

        self.doc.add_page_break()
        self._add_front_title("DANH MỤC BẢNG")
        for target in self.navigation_targets:
            if target["kind"] == "table":
                self._add_static_nav_entry(target, "Static TOC 2")
        self.doc.add_page_break()

    def _add_heading(self, block_index: int, level: int, text: str) -> None:
        if strip_markdown(text) == "MỞ ĐẦU":
            section = self.doc.add_section(WD_SECTION.NEW_PAGE)
            self._configure_section_geometry(section)
            self._set_header_footer(section, show=True)
        paragraph = self.doc.add_paragraph(style=f"Heading {level}")
        if strip_markdown(text) in {"TÓM TẮT", "MỞ ĐẦU"}:
            paragraph.paragraph_format.page_break_before = False
        add_inline_runs(paragraph, text)
        target = self.nav_by_block[block_index]
        add_bookmark(paragraph, target["bookmark"], self.bookmark_id)
        self.bookmark_id += 1
        self.counts["headings"] += 1
        self.after_heading = True

    def _add_caption(self, block_index: int, parts: tuple[str, str, str], target: dict[str, Any] | None = None) -> None:
        label, number, title = parts
        paragraph = self.doc.add_paragraph(style="Caption")
        label_run = paragraph.add_run(f"{label} {number}. ")
        set_run_font(label_run, "Times New Roman", 10, BODY)
        label_run.bold = True
        label_run.italic = True
        title_run = paragraph.add_run(title)
        set_run_font(title_run, "Times New Roman", 10, BODY)
        title_run.italic = True
        target = target or self.nav_by_block.get(block_index)
        if target:
            add_bookmark(paragraph, target["bookmark"], self.bookmark_id)
            self.bookmark_id += 1
        self.after_heading = False

    def _add_source_note(self, text: str) -> None:
        paragraph = self.doc.add_paragraph(style="Source Note")
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_inline_runs(paragraph, strip_markdown(text), base_size=9)
        for run in paragraph.runs:
            run.italic = True
            run.font.color.rgb = RGBColor.from_string(MUTED)

    def _add_body_paragraph(self, text: str) -> None:
        plain = strip_markdown(text)
        if plain.startswith("Nguồn:") or plain.startswith("Nguồn bảng:"):
            self._add_source_note(text)
            return
        style = "Reference Entry" if re.match(r"^\[\d+\]\s", plain) else "Normal"
        paragraph = self.doc.add_paragraph(style=style)
        if self.after_heading or style == "Reference Entry":
            paragraph.paragraph_format.first_line_indent = Cm(0)
        add_inline_runs(paragraph, text)
        self.after_heading = False

    def _add_list_item(self, list_type: str, marker: str, text: str) -> None:
        paragraph = self.doc.add_paragraph()
        paragraph.paragraph_format.left_indent = Cm(0.75)
        paragraph.paragraph_format.first_line_indent = Cm(-0.55)
        paragraph.paragraph_format.space_after = Pt(3)
        paragraph.paragraph_format.line_spacing = 1.25
        prefix = "•" if list_type == "bullet" else marker
        prefix_run = paragraph.add_run(f"{prefix} ")
        set_run_font(prefix_run, "Times New Roman", 11.5, NAVY)
        prefix_run.bold = True
        add_inline_runs(paragraph, text)
        self.after_heading = False

    def _resolve_image_path(self, relative_path: str) -> Path:
        path = REPORT_ROOT / relative_path
        if path.suffix.lower() == ".svg":
            path = path.with_suffix(".png")
        if not path.is_file():
            raise FileNotFoundError(path)
        return path

    def _add_image(self, block_index: int, alt: str, relative_path: str) -> None:
        path = self._resolve_image_path(relative_path)
        with Image.open(path) as image:
            width_px, height_px = image.size
        max_width_cm = 15.6
        max_height_cm = 12.6
        aspect = width_px / height_px
        width_cm = min(max_width_cm, max_height_cm * aspect)
        height_cm = width_cm / aspect
        if height_cm > max_height_cm:
            height_cm = max_height_cm
            width_cm = height_cm * aspect
        paragraph = self.doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.first_line_indent = Cm(0)
        paragraph.paragraph_format.space_before = Pt(5)
        paragraph.paragraph_format.space_after = Pt(2)
        paragraph.paragraph_format.keep_with_next = True
        inline_shape = paragraph.add_run().add_picture(str(path), width=Cm(width_cm), height=Cm(height_cm))
        inline_shape._inline.docPr.set("descr", strip_markdown(alt))
        inline_shape._inline.docPr.set("title", strip_markdown(alt)[:120])
        self.counts["figures"] += 1
        parts = caption_parts(alt)
        if parts:
            self._add_caption(block_index, parts, self.nav_by_block[block_index])
        self.after_heading = False

    @staticmethod
    def _looks_numeric(text: str) -> bool:
        plain = strip_markdown(text).replace("±", "").replace("–", "-")
        return bool(re.fullmatch(r"[\d\s.,/%:+\-×→]+", plain))

    def _column_widths(self, rows: list[list[str]]) -> list[float]:
        columns = len(rows[0])
        weights: list[float] = []
        for column in range(columns):
            lengths = [min(max(len(strip_markdown(row[column])), 4), 45) for row in rows if column < len(row)]
            average = sum(lengths) / len(lengths)
            weights.append(max(0.7, average ** 0.65))
        total = sum(weights)
        return [15.6 * weight / total for weight in weights]

    def _add_table(self, rows: list[list[str]]) -> None:
        if not rows or any(len(row) != len(rows[0]) for row in rows):
            raise ValueError("Malformed Markdown table")
        columns = len(rows[0])
        table = self.doc.add_table(rows=len(rows), cols=columns)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        table.style = "Table Grid"
        widths = self._column_widths(rows)
        set_table_fixed_widths(table, widths)
        font_size = 8.5 if columns >= 5 else 9.2
        for row_index, (word_row, values) in enumerate(zip(table.rows, rows)):
            set_row_cant_split(word_row)
            if row_index == 0:
                set_repeat_table_header(word_row)
            for column_index, (cell, value) in enumerate(zip(word_row.cells, values)):
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                set_cell_margins(cell)
                if row_index == 0:
                    set_cell_shading(cell, LIGHT_BLUE)
                elif row_index % 2 == 0:
                    set_cell_shading(cell, PALE_BLUE)
                paragraph = cell.paragraphs[0]
                paragraph.paragraph_format.first_line_indent = Cm(0)
                paragraph.paragraph_format.space_before = Pt(1)
                paragraph.paragraph_format.space_after = Pt(1)
                paragraph.paragraph_format.line_spacing = 1.05
                numeric_column = all(self._looks_numeric(row[column_index]) for row in rows[1:] if row[column_index].strip())
                if row_index == 0:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                elif numeric_column:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                elif len(strip_markdown(value)) <= 18:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                else:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
                add_inline_runs(paragraph, value, base_size=font_size)
                for run in paragraph.runs:
                    if row_index == 0:
                        run.bold = True
                        run.font.color.rgb = RGBColor.from_string(NAVY)
        self.counts["tables"] += 1
        self.after_heading = False

    def _add_code(self, language: str, code: str) -> None:
        label = self.doc.add_paragraph()
        label.paragraph_format.first_line_indent = Cm(0)
        label.paragraph_format.space_before = Pt(4)
        label.paragraph_format.space_after = Pt(2)
        run = label.add_run(f"Mã nguồn rút gọn{f' ({language})' if language else ''}")
        set_run_font(run, "Times New Roman", 9, MUTED)
        run.bold = True
        table = self.doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        table.style = "Table Grid"
        set_table_fixed_widths(table, [15.6])
        cell = table.cell(0, 0)
        set_cell_shading(cell, CODE_BG)
        set_cell_margins(cell, 100, 140, 100, 140)
        paragraph = cell.paragraphs[0]
        paragraph.paragraph_format.first_line_indent = Cm(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.0
        run = paragraph.add_run(code)
        set_run_font(run, "Consolas", 8.5, BODY)
        set_row_cant_split(table.rows[0])
        self.counts["code_blocks"] += 1
        self.after_heading = False

    def _add_equation(self, equation: str) -> None:
        paragraph = self.doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.first_line_indent = Cm(0)
        paragraph.paragraph_format.space_before = Pt(5)
        paragraph.paragraph_format.space_after = Pt(5)
        paragraph.paragraph_format.keep_together = True
        run = paragraph.add_run(latex_to_text(equation))
        set_run_font(run, "Cambria Math", 10.5, BODY)
        run.italic = True
        self.counts["equations"] += 1
        self.after_heading = False

    def add_manuscript(self) -> None:
        for block_index, block in enumerate(self.blocks):
            if block.kind == "heading":
                self._add_heading(block_index, *block.data)
            elif block.kind == "paragraph":
                parts = caption_parts(block.data) if is_caption_paragraph(block.data) else None
                if parts:
                    self._add_caption(block_index, parts)
                else:
                    self._add_body_paragraph(block.data)
            elif block.kind == "image":
                self._add_image(block_index, *block.data)
            elif block.kind == "table":
                self._add_table(block.data)
            elif block.kind == "code":
                self._add_code(*block.data)
            elif block.kind == "equation":
                self._add_equation(block.data)
            elif block.kind == "list_item":
                self._add_list_item(*block.data)
            else:
                raise ValueError(f"Unknown block type: {block.kind}")

    def save(self, output: Path, build_log_path: Path) -> None:
        output.parent.mkdir(parents=True, exist_ok=True)
        self.doc.save(output)
        log = {
            "schema_version": 1,
            "phase": 8,
            "output": output.relative_to(WORKSPACE).as_posix(),
            "output_sha256": sha256(output),
            "output_bytes": output.stat().st_size,
            "source_manuscript": MANUSCRIPT_PATH.relative_to(WORKSPACE).as_posix(),
            "source_manuscript_sha256": sha256(MANUSCRIPT_PATH),
            "source_manifest_sha256": sha256(REPORT_MANIFEST_PATH),
            "metadata": METADATA,
            "counts": self.counts,
            "navigation_targets": self.navigation_targets,
            "page_map_entries_used": len(self.page_map),
            "style_contract": {
                "paper": "A4 portrait",
                "font": "Times New Roman 11.5 pt",
                "body_spacing": "1.5 lines",
                "palette": {"navy": f"#{NAVY}", "blue": f"#{BLUE}"},
                "native_tables": True,
                "editable_code_blocks": True,
                "static_linked_navigation": True,
            },
        }
        build_log_path.parent.mkdir(parents=True, exist_ok=True)
        build_log_path.write_text(json.dumps(log, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        print(
            f"Built {output}: {output.stat().st_size:,} bytes; "
            f"{self.counts['headings']} headings, {self.counts['figures']} figures, "
            f"{self.counts['tables']} content tables."
        )


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--page-map", type=Path)
    parser.add_argument("--build-log", type=Path, default=DEFAULT_LOG)
    args = parser.parse_args()
    builder = ReportBuilder(args.page_map)
    builder.add_cover_and_title_page()
    builder.add_front_lists()
    builder.add_manuscript()
    builder.save(args.output.resolve(), args.build_log.resolve())


if __name__ == "__main__":
    main()
