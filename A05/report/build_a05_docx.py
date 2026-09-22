"""Generate the final academic DOCX report for Assignment 05.

This script is intentionally read-only with respect to datasets, metrics,
splits, checkpoints, and notebooks. It builds the report from the current
artifacts already saved under A05.
"""

from __future__ import annotations

import json
import math
import re
import sys
import textwrap
import zipfile
from pathlib import Path
from typing import Iterable


EXPECTED_PYTHON = "C:/Users/anhca/anaconda3/envs/tf312/python.exe"


def _normalize_path(path: str | Path) -> str:
    return Path(path).resolve().as_posix().lower()


def verify_environment() -> None:
    actual = _normalize_path(sys.executable)
    expected = _normalize_path(EXPECTED_PYTHON)
    if actual != expected:
        raise RuntimeError(
            "Sai Python executable. "
            f"Expected {EXPECTED_PYTHON}, got {sys.executable}"
        )

    missing: list[str] = []
    for package_name, import_name in [
        ("python-docx", "docx"),
        ("pandas", "pandas"),
        ("numpy", "numpy"),
        ("matplotlib", "matplotlib"),
        ("Pillow", "PIL"),
    ]:
        try:
            __import__(import_name)
        except ModuleNotFoundError:
            missing.append(package_name)
    if missing:
        raise RuntimeError(
            "Missing required package(s): " + ", ".join(missing) + ". "
            "Do not install automatically."
        )


verify_environment()

import matplotlib.pyplot as plt
import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from matplotlib.patches import FancyBboxPatch
from PIL import Image


SCRIPT_PATH = Path(__file__).resolve()
REPORT_DIR = SCRIPT_PATH.parent
ROOT = REPORT_DIR.parent
RESULTS_DIR = ROOT / "results"
FIGURES_DIR = RESULTS_DIR / "figures"
ASSETS_DIR = REPORT_DIR / "docx_assets"
OUTPUT_PATH = REPORT_DIR / "A05_Assignment_Report.docx"


def read_csv(relative_path: str) -> pd.DataFrame:
    path = ROOT / relative_path
    if not path.is_file():
        raise FileNotFoundError(f"Required CSV not found: {path}")
    return pd.read_csv(path)


def read_json(relative_path: str) -> dict:
    path = ROOT / relative_path
    if not path.is_file():
        raise FileNotFoundError(f"Required JSON not found: {path}")
    return json.loads(path.read_text(encoding="utf-8"))


def read_text(relative_path: str) -> str:
    path = ROOT / relative_path
    if not path.is_file():
        raise FileNotFoundError(f"Required text file not found: {path}")
    return path.read_text(encoding="utf-8")


def fmt_float(value: object, digits: int = 4) -> str:
    if value is None:
        return ""
    try:
        if pd.isna(value):
            return ""
    except TypeError:
        pass
    return f"{float(value):.{digits}f}"


def fmt_seconds(value: object) -> str:
    if value is None:
        return ""
    try:
        if pd.isna(value):
            return ""
    except TypeError:
        pass
    return f"{float(value):.2f}"


def fmt_int(value: object) -> str:
    if value is None:
        return ""
    try:
        if pd.isna(value):
            return ""
    except TypeError:
        pass
    return f"{int(value):,}"


def text_or_blank(value: object) -> str:
    if value is None:
        return ""
    try:
        if pd.isna(value):
            return ""
    except TypeError:
        pass
    return str(value)


def top_rows(df: pd.DataFrame, n: int) -> pd.DataFrame:
    return df.head(n).reset_index(drop=True)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_width(cell, width_cm: float) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = OxmlElement("w:tcW")
    tc_w.set(qn("w:w"), str(int(width_cm * 567)))
    tc_w.set(qn("w:type"), "dxa")
    tc_pr.append(tc_w)


def add_field(paragraph, instruction: str) -> None:
    run = paragraph.add_run()
    fld_char_begin = OxmlElement("w:fldChar")
    fld_char_begin.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = instruction
    fld_char_sep = OxmlElement("w:fldChar")
    fld_char_sep.set(qn("w:fldCharType"), "separate")
    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char_begin)
    run._r.append(instr_text)
    run._r.append(fld_char_sep)
    run._r.append(fld_char_end)


def set_run_font(run, size: float | None = None, bold: bool | None = None, italic: bool | None = None) -> None:
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def add_code_paragraph(document: Document, code: str) -> None:
    for line in code.strip("\n").splitlines():
        paragraph = document.add_paragraph(style="CodeBlock")
        run = paragraph.add_run(line)
        run.font.name = "Consolas"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
        run.font.size = Pt(9)


def _wrap_table_value(value: str, width_hint: float) -> str:
    text = str(value)
    if len(text) <= 1:
        return text
    max_chars = max(10, int(width_hint * 10.5))
    return "\n".join(
        textwrap.wrap(
            text,
            width=max_chars,
            break_long_words=False,
            break_on_hyphens=False,
        )
        or [text]
    )


def render_table_image(
    path: Path,
    headers: list[str],
    rows: list[list[str]],
    widths: list[float] | None = None,
) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    if widths is None:
        widths = [1.0 for _ in headers]
    total_width = sum(widths)
    col_widths = [w / total_width for w in widths]

    wrapped_headers = [_wrap_table_value(h, w) for h, w in zip(headers, widths)]
    wrapped_rows = [[_wrap_table_value(v, w) for v, w in zip(row, widths)] for row in rows]
    line_count = sum(max(cell.count("\n") + 1 for cell in row) for row in wrapped_rows)
    header_lines = max(cell.count("\n") + 1 for cell in wrapped_headers)
    fig_height = min(14.5, max(2.2, 0.42 * line_count + 0.55 * header_lines + 0.75))
    fig_width = 12.2

    fig, ax = plt.subplots(figsize=(fig_width, fig_height))
    ax.axis("off")
    table = ax.table(
        cellText=wrapped_rows,
        colLabels=wrapped_headers,
        cellLoc="center",
        colLoc="center",
        colWidths=col_widths,
        loc="center",
    )
    table.auto_set_font_size(False)
    font_size = 8.3 if len(headers) >= 8 else 9.4
    table.set_fontsize(font_size)
    table.scale(1, 1.52)

    for (row_idx, _), cell in table.get_celld().items():
        cell.set_edgecolor("#5B6770")
        cell.set_linewidth(0.45)
        if row_idx == 0:
            cell.set_facecolor("#CFE3F4")
            cell.set_text_props(weight="bold", color="#17365D")
        else:
            cell.set_facecolor("#FFFFFF" if row_idx % 2 else "#F8FBFD")
    fig.tight_layout(pad=0.25)
    fig.savefig(path, dpi=220, bbox_inches="tight")
    plt.close(fig)


class ReportBuilder:
    def __init__(self) -> None:
        self.document = Document()
        self.figure_count = 0
        self.table_count = 0
        self.code_count = 0
        self.figure_entries: list[str] = []
        self.table_entries: list[str] = []
        self._setup_document()

    def _setup_document(self) -> None:
        document = self.document
        section = document.sections[0]
        section.page_width = Cm(21)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(2.2)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(2.0)
        section.right_margin = Cm(2.0)

        styles = document.styles
        normal = styles["Normal"]
        normal.font.name = "Times New Roman"
        normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        normal.font.size = Pt(13)
        normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        normal.paragraph_format.first_line_indent = Cm(0.8)
        normal.paragraph_format.line_spacing = 1.3
        normal.paragraph_format.space_after = Pt(6)

        for style_name, size, color in [
            ("Title", 18, "1F4E79"),
            ("Heading 1", 15, "1F4E79"),
            ("Heading 2", 13.5, "1F4E79"),
            ("Heading 3", 13, "365F91"),
        ]:
            style = styles[style_name]
            style.font.name = "Times New Roman"
            style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
            style.font.size = Pt(size)
            style.font.bold = True
            style.font.color.rgb = RGBColor.from_string(color)
            style.paragraph_format.space_before = Pt(10)
            style.paragraph_format.space_after = Pt(6)
            style.paragraph_format.keep_with_next = True

        caption = styles["Caption"]
        caption.font.name = "Times New Roman"
        caption._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        caption.font.size = Pt(11)
        caption.font.italic = True
        caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption.paragraph_format.first_line_indent = Cm(0)
        caption.paragraph_format.space_after = Pt(8)

        if "CodeBlock" not in [style.name for style in styles]:
            code_style = styles.add_style("CodeBlock", WD_STYLE_TYPE.PARAGRAPH)
        else:
            code_style = styles["CodeBlock"]
        code_style.font.name = "Consolas"
        code_style._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
        code_style.font.size = Pt(9)
        code_style.paragraph_format.first_line_indent = Cm(0)
        code_style.paragraph_format.left_indent = Cm(0.5)
        code_style.paragraph_format.space_before = Pt(0)
        code_style.paragraph_format.space_after = Pt(0)
        code_style.paragraph_format.line_spacing = 1.05

        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        footer.add_run("Trang ")
        add_field(footer, "PAGE")

    def add_cover(self) -> None:
        document = self.document
        for _ in range(3):
            document.add_paragraph()
        title = document.add_paragraph()
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = title.add_run("ASSIGNMENT 05")
        set_run_font(run, 20, True)

        subtitle = document.add_paragraph()
        subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = subtitle.add_run(
            "TÌM HIỂU, XÂY DỰNG VÀ ĐÁNH GIÁ\n"
            "CÁC KIẾN TRÚC CONVOLUTIONAL NEURAL NETWORK"
        )
        set_run_font(run, 17, True)
        run.font.color.rgb = RGBColor.from_string("1F4E79")

        note = document.add_paragraph()
        note.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = note.add_run("Báo cáo học thuật tổng hợp từ các thí nghiệm đã hoàn thành trong A05")
        set_run_font(run, 13, False, True)

        for _ in range(6):
            document.add_paragraph()
        meta = [
            "Phạm vi: Basic CNN, AlexNet-inspired, VGG-inspired, ResNet-inspired",
            "Dữ liệu: EuroSAT RGB, Oxford-IIIT Pet, CDC Diabetes Health Indicators",
            "Môi trường: Python 3.12 - TensorFlow, CPU-only",
            "Nguồn số liệu: CSV kết quả, notebooks đã thực thi và audit artifacts trong A05",
        ]
        for line in meta:
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.first_line_indent = Cm(0)
            run = paragraph.add_run(line)
            set_run_font(run, 12, False)
        document.add_page_break()

    def add_front_matter(self) -> None:
        self.document.add_heading("MỤC LỤC", level=1)
        paragraph = self.document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        add_field(paragraph, r'TOC \o "1-3" \h \z \u')
        paragraph.add_run("Nhấn chuột phải và chọn Update Field trong Microsoft Word để cập nhật số trang.")
        self.document.add_page_break()

        self.document.add_heading("DANH MỤC HÌNH", level=1)
        paragraph = self.document.add_paragraph(
            "Danh mục hình được tạo trong báo cáo với caption chuẩn. Sau khi mở bằng Microsoft Word, "
            "có thể cập nhật lại danh mục tự động nếu cần."
        )
        paragraph.paragraph_format.first_line_indent = Cm(0)
        self.document.add_heading("DANH MỤC BẢNG", level=1)
        paragraph = self.document.add_paragraph(
            "Các bảng được đánh số theo thứ tự xuất hiện và sử dụng số liệu đọc trực tiếp từ các CSV kết quả."
        )
        paragraph.paragraph_format.first_line_indent = Cm(0)
        self.document.add_page_break()

    def heading(self, text: str, level: int = 1) -> None:
        self.document.add_heading(text, level=level)

    def paragraph(self, text: str, *, no_indent: bool = False) -> None:
        paragraph = self.document.add_paragraph(text)
        if no_indent:
            paragraph.paragraph_format.first_line_indent = Cm(0)

    def bullets(self, items: Iterable[str]) -> None:
        for item in items:
            paragraph = self.document.add_paragraph(style="List Bullet")
            paragraph.add_run(item)

    def table_caption(self, caption: str) -> str:
        self.table_count += 1
        label = f"Bảng {self.table_count}. {caption}"
        self.table_entries.append(label)
        paragraph = self.document.add_paragraph(label, style="Caption")
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        return label

    def add_table(self, caption: str, headers: list[str], rows: list[list[str]], widths: list[float] | None = None) -> None:
        self.table_caption(caption)
        table_number = self.table_count
        image_path = ASSETS_DIR / "tables" / f"table_{table_number:02d}.png"
        render_table_image(image_path, headers, rows, widths)
        self.document.add_picture(str(image_path), width=Inches(6.65))
        self.document.add_paragraph()

    def figure(self, relative_or_absolute_path: str | Path, caption: str, width_inches: float = 6.35) -> None:
        path = Path(relative_or_absolute_path)
        if not path.is_absolute():
            path = ROOT / path
        if not path.is_file():
            raise FileNotFoundError(f"Figure not found: {path}")
        self.document.add_picture(str(path), width=Inches(width_inches))
        self.figure_count += 1
        label = f"Hình {self.figure_count}. {caption}"
        self.figure_entries.append(label)
        paragraph = self.document.add_paragraph(label, style="Caption")
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    def code(self, caption: str, code: str) -> None:
        self.code_count += 1
        paragraph = self.document.add_paragraph(f"Mã {self.code_count}. {caption}", style="Caption")
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_code_paragraph(self.document, code)
        self.document.add_paragraph()

    def page_break(self) -> None:
        self.document.add_page_break()


def create_diagram(path: Path, title: str, labels: list[str], arrows: bool = True) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    fig, ax = plt.subplots(figsize=(11.6, 3.0))
    ax.axis("off")
    x_positions = [0.075 + i * (0.85 / max(1, len(labels) - 1)) for i in range(len(labels))]
    for x, label in zip(x_positions, labels):
        ax.text(
            x,
            0.52,
            label,
            ha="center",
            va="center",
            fontsize=12.5,
            linespacing=1.12,
            bbox=dict(boxstyle="round,pad=0.52", fc="#EEF5FB", ec="#1F4E79", lw=1.7),
            transform=ax.transAxes,
        )
    if arrows:
        for left, right in zip(x_positions[:-1], x_positions[1:]):
            ax.annotate(
                "",
                xy=(right - 0.065, 0.52),
                xytext=(left + 0.065, 0.52),
                arrowprops=dict(arrowstyle="->", color="#1F4E79", lw=2.0),
                xycoords=ax.transAxes,
                textcoords=ax.transAxes,
            )
    ax.set_title(title, fontsize=15, color="#1F4E79", weight="bold", pad=16)
    fig.tight_layout()
    fig.savefig(path, dpi=220, bbox_inches="tight")
    plt.close(fig)


def create_residual_diagram(path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    fig, ax = plt.subplots(figsize=(10.5, 4.0))
    ax.axis("off")
    nodes = {
        "x": (0.08, 0.48, "x"),
        "f": (0.42, 0.48, "F(x)\nConv - ReLU - Conv"),
        "add": (0.71, 0.48, "+"),
        "y": (0.9, 0.48, "y = F(x) + shortcut(x)"),
    }
    for _, (x, y, label) in nodes.items():
        ax.text(
            x,
            y,
            label,
            ha="center",
            va="center",
            fontsize=12.5,
            linespacing=1.15,
            bbox=dict(boxstyle="round,pad=0.55", fc="#EEF5FB", ec="#1F4E79", lw=1.8),
            transform=ax.transAxes,
        )
    arrow = dict(arrowstyle="->", color="#1F4E79", lw=2.1)
    ax.annotate("", xy=(0.30, 0.48), xytext=(0.15, 0.48), arrowprops=arrow, xycoords=ax.transAxes)
    ax.annotate("", xy=(0.64, 0.48), xytext=(0.55, 0.48), arrowprops=arrow, xycoords=ax.transAxes)
    ax.annotate("", xy=(0.83, 0.48), xytext=(0.76, 0.48), arrowprops=arrow, xycoords=ax.transAxes)
    ax.annotate("", xy=(0.71, 0.62), xytext=(0.08, 0.76), arrowprops=arrow, xycoords=ax.transAxes)
    ax.text(0.4, 0.82, "shortcut hoặc projection 1x1 nếu channel khác", ha="center", fontsize=11.5, transform=ax.transAxes)
    ax.set_title("Residual block trong ResNet-inspired CNN", fontsize=15, color="#1F4E79", weight="bold", pad=14)
    fig.tight_layout()
    fig.savefig(path, dpi=220, bbox_inches="tight")
    plt.close(fig)


def create_convolution_diagram(path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    fig, ax = plt.subplots(figsize=(9.8, 4.4))
    ax.axis("off")
    grid = [
        [1, 2, 3, 4],
        [5, 6, 7, 8],
        [9, 10, 11, 12],
        [13, 14, 15, 16],
    ]
    kernel = [[1, 0], [0, -1]]
    for r in range(4):
        for c in range(4):
            fc = "#FFF2CC" if r < 2 and c < 2 else "#F5F5F5"
            ax.add_patch(plt.Rectangle((c, 4 - r), 0.82, 0.82, fc=fc, ec="#777777", lw=1.1))
            ax.text(c + 0.41, 4.41 - r, str(grid[r][c]), ha="center", va="center", fontsize=12)
    for r in range(2):
        for c in range(2):
            ax.add_patch(plt.Rectangle((5.6 + c, 3.5 - r), 0.82, 0.82, fc="#E2F0D9", ec="#777777", lw=1.1))
            ax.text(6.01 + c, 3.91 - r, str(kernel[r][c]), ha="center", va="center", fontsize=12)
    ax.text(1.65, 5.05, "Input patch", ha="center", fontsize=13, weight="bold", color="#17365D")
    ax.text(6.42, 4.55, "Kernel", ha="center", fontsize=13, weight="bold", color="#17365D")
    ax.annotate(
        "nhân từng phần tử\nrồi cộng",
        xy=(5.15, 3.8),
        xytext=(3.75, 3.8),
        fontsize=11,
        ha="center",
        va="center",
        arrowprops=dict(arrowstyle="->", color="#1F4E79", lw=1.8),
    )
    ax.text(
        6.05,
        1.65,
        "1*1 + 2*0 + 5*0 + 6*(-1) = -5",
        ha="center",
        fontsize=12,
        bbox=dict(boxstyle="round,pad=0.35", fc="#EEF5FB", ec="#1F4E79", lw=1.1),
    )
    ax.set_xlim(-0.3, 8.2)
    ax.set_ylim(1.0, 5.45)
    ax.set_title("Ví dụ convolution/cross-correlation 2D", fontsize=15, color="#1F4E79", weight="bold")
    fig.tight_layout()
    fig.savefig(path, dpi=220, bbox_inches="tight")
    plt.close(fig)


def create_conceptual_diagrams() -> dict[str, Path]:
    assets = {
        "composition": ASSETS_DIR / "cnn_function_composition.png",
        "convolution": ASSETS_DIR / "convolution_example.png",
        "basic_arch": ASSETS_DIR / "basic_cnn_architecture.png",
        "vgg_block": ASSETS_DIR / "vgg_block.png",
        "residual": ASSETS_DIR / "residual_block.png",
        "split_roles": ASSETS_DIR / "train_val_test_roles.png",
        "conv1d": ASSETS_DIR / "diabetes_conv1d_representation.png",
    }
    create_diagram(
        assets["composition"],
        "CNN dưới góc nhìn hợp của các hàm",
        ["Input", "Convolution", "ReLU", "Pooling", "Feature\nrepresentation", "Classifier", "Prediction"],
    )
    create_convolution_diagram(assets["convolution"])
    create_diagram(
        assets["basic_arch"],
        "BasicCNN2D đã triển khai",
        ["RGB image", "Conv2D 32\nReLU", "MaxPool", "Conv2D 64\nReLU", "MaxPool", "GlobalAvgPool", "Dense", "Softmax"],
    )
    create_diagram(
        assets["vgg_block"],
        "VGG-inspired block",
        ["Input", "Conv 3x3\nReLU", "Conv 3x3\nReLU", "MaxPool", "Next block"],
    )
    create_residual_diagram(assets["residual"])
    create_diagram(
        assets["split_roles"],
        "Vai trò của train, validation và test",
        ["Train\nfit weights", "Validation\nselect protocol", "Test\nfinal evaluation"],
    )
    create_diagram(
        assets["conv1d"],
        "Biểu diễn Conv1D cho CDC Diabetes",
        ["21 predictor\nfeatures", "preserve CSV\nfeature order", "reshape", "(21, 1)", "Conv1D"],
    )
    return assets


def notebook_overview() -> list[list[str]]:
    rows: list[list[str]] = []
    for path in sorted((ROOT / "notebooks").glob("*.ipynb")):
        nb = json.loads(path.read_text(encoding="utf-8"))
        markdown = sum(1 for cell in nb["cells"] if cell.get("cell_type") == "markdown")
        code = sum(1 for cell in nb["cells"] if cell.get("cell_type") == "code")
        errors = sum(
            1
            for cell in nb["cells"]
            if cell.get("cell_type") == "code"
            for output in cell.get("outputs", [])
            if output.get("output_type") == "error"
        )
        rows.append([path.name, str(markdown), str(code), str(errors)])
    return rows


def metrics_table_rows(df: pd.DataFrame, dataset: str) -> list[list[str]]:
    if dataset == "diabetes":
        rows = []
        for _, row in df.iterrows():
            rows.append(
                [
                    row["model_family"],
                    fmt_float(row["test_accuracy"]),
                    fmt_float(row["test_macro_precision"]),
                    fmt_float(row["test_macro_recall"]),
                    fmt_float(row["test_macro_f1"]),
                    fmt_float(row["test_weighted_f1"]),
                    fmt_float(row["test_class1_prediabetes_recall"]),
                    fmt_float(row["test_class2_diabetes_recall"]),
                    fmt_int(row["trainable_parameter_count"]),
                    fmt_seconds(row["training_time_seconds"]),
                    fmt_seconds(row["test_inference_time_seconds"]),
                ]
            )
        return rows
    rows = []
    for _, row in df.iterrows():
        rows.append(
            [
                row["model_family"],
                fmt_float(row["test_accuracy"]),
                fmt_float(row.get("macro_precision", row.get("test_macro_precision"))),
                fmt_float(row.get("macro_recall", row.get("test_macro_recall"))),
                fmt_float(row.get("macro_f1", row.get("test_macro_f1"))),
                fmt_int(row["trainable_parameter_count"]),
                fmt_seconds(row["training_time_seconds"]),
                fmt_seconds(row.get("inference_time_seconds", row.get("test_inference_time_seconds"))),
            ]
        )
    return rows


def add_intro(builder: ReportBuilder) -> None:
    builder.heading("I. GIỚI THIỆU", 1)
    builder.heading("1.1. Bối cảnh", 2)
    builder.paragraph(
        "Convolutional Neural Network (CNN) là một họ mô hình quan trọng trong deep learning, "
        "đặc biệt phù hợp với dữ liệu có cấu trúc lưới như ảnh. Thay vì học trực tiếp một ánh xạ "
        "từ toàn bộ pixel đến nhãn bằng một tầng Dense rất lớn, CNN dùng convolution để khai thác "
        "các mẫu cục bộ, chia sẻ trọng số theo không gian và tạo feature map ở nhiều mức trừu tượng. "
        "Cách tổ chức này giúp mô hình học từ cạnh, texture, vùng màu và cấu trúc đối tượng trước khi "
        "đưa biểu diễn cuối vào classifier."
    )
    builder.heading("1.2. Mục tiêu Assignment 05", 2)
    builder.paragraph(
        "Assignment 05 tập trung vào việc hiểu CNN như một hợp của các hàm, khảo sát sự phát triển "
        "của kiến trúc CNN, xây dựng bốn family kiến trúc từ scratch và đánh giá chúng trên hai image "
        "datasets cùng một dataset tabular diabetes. Bốn family gồm Basic CNN, AlexNet-inspired, "
        "VGG-inspired và ResNet-inspired."
    )
    builder.bullets(
        [
            "Giải thích các thành phần convolution, ReLU, pooling, Flatten/Global Average Pooling, Dense và Softmax.",
            "Xây dựng các biến thể Conv2D cho EuroSAT và Oxford-IIIT Pet.",
            "Xây dựng các biến thể Conv1D cho CDC Diabetes như một thích nghi bắt buộc theo yêu cầu bài tập.",
            "Chọn siêu tham số bằng validation evidence, giữ test set độc lập cho đánh giá cuối.",
            "So sánh performance với parameter count, training time và inference time thay vì chỉ nhìn một metric.",
        ]
    )
    builder.heading("1.3. Phạm vi thực hiện", 2)
    builder.paragraph(
        "Toàn bộ mô hình được triển khai bằng TensorFlow/Keras, không dùng pretrained weights và không dùng transfer learning. "
        "Các kiến trúc là phiên bản educational scaled variants để phù hợp CPU-only, không phải bản lịch sử đầy đủ "
        "của AlexNet, VGG hay ResNet. Các quyết định tuning dùng validation set; test set chỉ xuất hiện ở bước "
        "đánh giá cuối cùng sau khi protocol đã đóng băng."
    )

def add_theory(builder: ReportBuilder, diagrams: dict[str, Path]) -> None:
    builder.heading("II. CƠ SỞ LÝ THUYẾT VỀ CNN", 1)
    builder.heading("2.1. CNN dưới góc nhìn hợp của các hàm", 2)
    builder.paragraph(
        "Một CNN có thể được mô tả như một chuỗi hàm biến đổi dữ liệu: "
        "f(x) = f_n(...f_2(f_1(x))). Trong bài toán ảnh, đầu vào đi qua convolution để tạo feature map, "
        "qua ReLU để đưa phi tuyến, qua pooling để giảm kích thước không gian, rồi qua phần classifier "
        "để tạo phân phối xác suất trên các lớp. Ý tưởng composition có thể viết ngắn gọn là "
        "Softmax ∘ Dense ∘ Pool ∘ ReLU ∘ Conv. Kết quả của mỗi phép biến đổi trở thành đầu vào của phép biến đổi kế tiếp."
    )
    builder.figure(diagrams["composition"], "CNN dưới góc nhìn hợp của các hàm.", 6.3)
    builder.heading("2.2. Phép convolution", 2)
    builder.paragraph(
        "Trong CNN, convolution thường được cài đặt dưới dạng cross-correlation: kernel trượt qua input, "
        "nhân từng phần tử trong receptive field với trọng số kernel rồi cộng lại để tạo một điểm trên feature map. "
        "Với input X và kernel K, một biểu diễn đơn giản là Y(i,j) = Σ_m Σ_n X(i+m,j+n)K(m,n). "
        "Stride quy định bước trượt, padding quy định cách xử lý biên, còn feature map là bản đồ phản hồi "
        "của một filter với mẫu cục bộ mà nó học được."
    )
    builder.figure(diagrams["convolution"], "Ví dụ convolution/cross-correlation 2D trên một patch nhỏ.", 5.6)
    builder.code(
        "Hàm convolution 2D từ notebook nền tảng, với chú thích tiếng Việt.",
        """
def conv2d_valid(input_matrix, kernel):
    # Chuyển input và kernel sang số thực để tính toán ổn định.
    input_matrix = np.asarray(input_matrix, dtype=float)
    kernel = np.asarray(kernel, dtype=float)
    out_height = input_matrix.shape[0] - kernel.shape[0] + 1
    out_width = input_matrix.shape[1] - kernel.shape[1] + 1
    output = np.zeros((out_height, out_width), dtype=float)

    # Trượt kernel qua từng vùng hợp lệ của ảnh.
    for i in range(out_height):
        for j in range(out_width):
            patch = input_matrix[i:i + kernel.shape[0], j:j + kernel.shape[1]]
            output[i, j] = np.sum(patch * kernel)
    return output
""",
    )
    builder.paragraph(
        "Đoạn mã trên giữ đúng thuật toán trong notebook: mỗi vị trí output là tổng của tích từng phần tử "
        "giữa patch và kernel. Đây là phép toán cốt lõi giúp CNN học các mẫu cục bộ thay vì nối phẳng toàn bộ ảnh ngay từ đầu."
    )
    builder.heading("2.3. Hàm kích hoạt ReLU", 2)
    builder.paragraph(
        "ReLU được định nghĩa bởi ReLU(x) = max(0, x). Hàm này giữ các giá trị dương và đưa giá trị âm về 0, "
        "tạo tính phi tuyến cho mạng. Nếu chỉ ghép nhiều phép biến đổi tuyến tính, toàn bộ mạng vẫn tương đương "
        "một phép tuyến tính lớn; ReLU làm cho chuỗi convolution có khả năng biểu diễn quan hệ phức tạp hơn."
    )
    builder.code(
        "Hàm ReLU trong notebook nền tảng.",
        """
def relu(values):
    # Giữ giá trị dương và đưa giá trị âm về 0.
    return np.maximum(values, 0)
""",
    )
    builder.heading("2.4. Pooling", 2)
    builder.paragraph(
        "MaxPooling lấy giá trị lớn nhất trong từng cửa sổ nhỏ, ví dụ cửa sổ 2x2 với stride 2. "
        "Pooling làm giảm kích thước feature map, giảm chi phí tính toán và giúp biểu diễn ít nhạy hơn với "
        "dịch chuyển nhỏ. Đổi lại, pooling cũng làm mất một phần thông tin vị trí chi tiết; vì vậy placement "
        "và số lần pooling phải được cân nhắc theo kích thước input."
    )
    builder.code(
        "MaxPooling 2D minh họa từ notebook nền tảng.",
        """
def max_pool2d(feature_map, pool_size=(2, 2), stride=2):
    # Tính kích thước output sau khi cửa sổ pooling trượt qua feature map.
    feature_map = np.asarray(feature_map, dtype=float)
    pool_height, pool_width = pool_size
    out_height = (feature_map.shape[0] - pool_height) // stride + 1
    out_width = (feature_map.shape[1] - pool_width) // stride + 1
    output = np.zeros((out_height, out_width), dtype=float)

    for i in range(out_height):
        for j in range(out_width):
            row, col = i * stride, j * stride
            window = feature_map[row:row + pool_height, col:col + pool_width]
            output[i, j] = np.max(window)
    return output
""",
    )
    builder.heading("2.5. Từ feature map đến dự đoán", 2)
    builder.paragraph(
        "Sau khi CNN đã trích xuất feature map, mô hình cần chuyển biểu diễn này thành logits hoặc xác suất. "
        "Flatten nối toàn bộ feature map thành vector, nhưng có thể làm parameter count tăng mạnh nếu feature map còn lớn. "
        "Global Average Pooling lấy trung bình theo chiều không gian cho mỗi channel, giúp giảm parameter count và phù hợp "
        "với ràng buộc CPU trong assignment. Dense layer kết hợp các đặc trưng đã học; Softmax biến logits thành "
        "phân phối xác suất. Với nhãn nguyên nhiều lớp, Sparse Categorical Cross-Entropy đo mức sai lệch giữa phân phối dự đoán "
        "và lớp đúng."
    )


def add_architecture_chapter(builder: ReportBuilder, diagrams: dict[str, Path]) -> None:
    builder.heading("III. SỰ PHÁT TRIỂN CỦA CÁC KIẾN TRÚC CNN", 1)
    builder.paragraph(
        "Các kiến trúc trong `models/architectures.py` được xây dựng từ scratch bằng các layer Keras cơ bản. "
        "Biến thể Conv2D dùng cho ảnh RGB, còn biến thể Conv1D dùng cho 21 predictor của CDC Diabetes. "
        "Những mô hình này giữ ý tưởng kiến trúc nhưng được scale nhỏ để chạy được trong môi trường CPU-only."
    )
    builder.heading("3.1. Basic CNN", 2)
    builder.paragraph(
        "Basic CNN là baseline đơn giản nhất. Với ảnh, mô hình dùng hai stage Conv2D-ReLU-MaxPooling, sau đó "
        "Global Average Pooling, Dense và Softmax. Baseline này giúp kiểm tra mức hiệu quả tối thiểu của convolution "
        "trước khi thêm depth hoặc residual connection."
    )
    builder.figure(diagrams["basic_arch"], "Sơ đồ kiến trúc BasicCNN2D đã triển khai.", 6.3)
    builder.code(
        "Khối đại diện trong BasicCNN2D.",
        """
inputs = keras.Input(shape=input_shape, name="image")
x = _conv_relu_2d(inputs, 32, 3, "basic_block1")
x = layers.MaxPooling2D(pool_size=2, name="basic_pool1")(x)
x = _conv_relu_2d(x, 64, 3, "basic_block2")
x = layers.MaxPooling2D(pool_size=2, name="basic_pool2")(x)
outputs = _classification_head_2d(x, num_classes, dense_units=64, name="basic_head")
""",
    )
    builder.heading("3.2. AlexNet-inspired CNN", 2)
    builder.paragraph(
        "AlexNet-inspired tăng capacity so với Basic CNN bằng bốn convolution layer và ba pooling stage trong biến thể 2D. "
        "Khối đầu dùng kernel 5x5 để mở receptive field, các tầng sau dùng kernel 3x3. Báo cáo không xem đây là bản sao "
        "AlexNet lịch sử; mô hình bỏ các Dense layer rất lớn và dùng Global Average Pooling để tránh parameter explosion."
    )
    builder.heading("3.3. VGG-inspired CNN", 2)
    builder.paragraph(
        "VGG-inspired nhấn mạnh cấu trúc block lặp có hệ thống: mỗi block gồm hai convolution 3x3 liên tiếp rồi pooling. "
        "Khác với AlexNet-inspired, điểm chính ở đây là sự đều đặn của repeated small-kernel blocks, cho phép tăng độ sâu "
        "theo cách có cấu trúc."
    )
    builder.figure(diagrams["vgg_block"], "Một VGG-inspired block với hai convolution nhỏ liên tiếp.", 5.8)
    builder.code(
        "Vòng lặp tạo các VGG-inspired blocks.",
        """
for block_index, filters in enumerate((32, 64, 128), start=1):
    # Mỗi block dùng hai convolution 3x3 để tăng độ sâu có cấu trúc.
    x = _conv_relu_2d(x, filters, 3, f"vgg_block{block_index}_conv1")
    x = _conv_relu_2d(x, filters, 3, f"vgg_block{block_index}_conv2")
    # Pooling được đặt sau block để giảm kích thước feature map.
    x = layers.MaxPooling2D(pool_size=2, name=f"vgg_block{block_index}_pool")(x)
""",
    )
    builder.heading("3.4. ResNet-inspired CNN", 2)
    builder.paragraph(
        "ResNet-inspired đưa residual learning vào kiến trúc: y = F(x) + shortcut(x). Main branch học phần biến đổi F(x), "
        "shortcut truyền thông tin đầu vào sang phép cộng. Khi số channel hoặc stride khác nhau, shortcut dùng projection "
        "1x1 convolution để tensor có shape tương thích; mô hình không dựa vào broadcasting ngẫu nhiên."
    )
    builder.figure(diagrams["residual"], "Residual block với main branch, shortcut và phép cộng.", 5.9)
    builder.code(
        "Residual block 2D đã triển khai trong architectures.py.",
        """
def residual_block_2d(x, filters, *, name, stride=1):
    shortcut = x
    # Main branch học residual mapping F(x).
    x = layers.Conv2D(filters, 3, strides=stride, padding="same", name=f"{name}_conv1")(x)
    x = layers.ReLU(name=f"{name}_relu1")(x)
    x = layers.Conv2D(filters, 3, padding="same", name=f"{name}_conv2")(x)

    # Projection 1x1 khi số channel hoặc stride không khớp.
    if int(shortcut.shape[-1]) != filters or stride != 1:
        shortcut = layers.Conv2D(filters, 1, strides=stride, padding="same", name=f"{name}_projection")(shortcut)

    x = layers.Add(name=f"{name}_add")([x, shortcut])
    return layers.ReLU(name=f"{name}_relu_out")(x)
""",
    )
    builder.heading("3.5. So sánh về mặt kiến trúc", 2)
    builder.add_table(
        "So sánh các family kiến trúc theo implementation thực tế",
        ["Family", "Ý tưởng chính", "Depth/blocks", "Residual", "Representative parameters", "Khác biệt chính"],
        [
            ["Basic CNN", "Baseline convolution đơn giản", "2 Conv2D hoặc 2 Conv1D", "Không", "24,202 / 2,787", "Mốc so sánh nhỏ nhất"],
            ["AlexNet-inspired", "Tăng depth và capacity", "4 convolution layers", "Không", "260,170 / 29,635", "Nhiều stage hơn Basic CNN"],
            ["VGG-inspired", "Repeated small-kernel blocks", "6 convolution layers", "Không", "304,810 / 38,299", "Block 3x3 lặp có hệ thống"],
            ["ResNet-inspired", "Residual learning y = F(x) + x", "Stem + 3 residual blocks", "Có", "324,490 / 44,387", "Shortcut và projection khi cần"],
        ],
        [2.8, 3.8, 3.0, 1.8, 2.8, 3.8],
    )


def add_data_and_method(builder: ReportBuilder, diagrams: dict[str, Path], audit: dict) -> None:
    datasets = audit["datasets"]
    euro = datasets["eurosat"]["observed"]
    oxford = datasets["oxford_pets"]["observed"]
    diabetes = datasets["diabetes"]["observed"]
    builder.heading("IV. DỮ LIỆU VÀ THIẾT KẾ THỰC NGHIỆM", 1)
    builder.heading("4.1. EuroSAT RGB", 2)
    builder.paragraph(
        f"EuroSAT RGB gồm {fmt_int(euro['jpeg_image_count'])} ảnh JPEG, {euro['class_count']} lớp, "
        "mỗi ảnh có kích thước 64x64 và 3 kênh RGB. Vì mọi ảnh đã có cùng resolution, input 64x64 là "
        "data-determined. Dataset này đại diện cho phân loại land-use/land-cover từ ảnh vệ tinh."
    )
    builder.add_table(
        "Phân bố lớp EuroSAT",
        ["Class", "Số ảnh"],
        [[name, fmt_int(count)] for name, count in euro["class_distribution"].items()],
        [6.0, 3.0],
    )
    builder.figure("results/figures/eurosat/examples/eurosat_examples.png", "Một số ảnh ví dụ từ EuroSAT.", 6.2)
    builder.heading("4.2. Oxford-IIIT Pet", 2)
    builder.paragraph(
        f"Oxford-IIIT Pet trong assignment chỉ dùng official samples từ `annotations/trainval.txt` và "
        f"`annotations/test.txt`: trainval = {oxford['trainval_record_count']}, test = {oxford['test_record_count']}, "
        f"tổng unique official samples = {oxford['official_unique_sample_count']}, số lớp = {oxford['class_count']}, "
        f"trainval/test overlap = {oxford['trainval_test_overlap_count']}. 41 raw images không xuất hiện trong annotation "
        "được gọi là unreferenced raw images và không đi vào thí nghiệm."
    )
    builder.paragraph(
        "Audit cũng phát hiện bốn file có đuôi `.jpg` nhưng nội dung PNG: Abyssinian_5.jpg, Egyptian_Mau_14.jpg, "
        "Egyptian_Mau_156.jpg và Egyptian_Mau_186.jpg. Vì pipeline decode theo nội dung ảnh bằng Pillow và convert RGB "
        "trong bộ nhớ, bốn sample hợp lệ này vẫn được sử dụng mà không sửa raw file."
    )
    builder.figure("results/figures/oxford_pets/examples/oxford_original_size_distribution.png", "Phân bố kích thước ảnh gốc của Oxford-IIIT Pet.", 6.2)
    builder.figure("results/figures/oxford_pets/examples/oxford_class_examples.png", "Ví dụ ảnh Oxford-IIIT Pet theo lớp.", 5.3)
    builder.heading("4.3. CDC Diabetes Health Indicators", 2)
    target_counts = diabetes["target_distribution"]
    builder.paragraph(
        f"CDC Diabetes Health Indicators có {fmt_int(diabetes['row_count'])} dòng, {diabetes['predictor_feature_count']} predictor "
        "và target `Diabetes_012` gồm 3 lớp: 0 = No diabetes, 1 = Prediabetes, 2 = Diabetes. "
        f"Phân bố target rất lệch: class 0 có {fmt_int(target_counts['0.0'])}, class 1 có {fmt_int(target_counts['1.0'])}, "
        f"class 2 có {fmt_int(target_counts['2.0'])}."
    )
    builder.paragraph(
        f"Audit ghi nhận {fmt_int(diabetes['duplicate_row_count'])} duplicate rows. Các dòng này không bị xóa tự động vì dataset "
        "gồm nhiều biến khảo sát rời rạc/binary và không có respondent ID để khẳng định duplicate là lỗi thu thập dữ liệu."
    )
    builder.figure("results/figures/diabetes/hyperparameters/feature_distributions.png", "Phân bố một số feature của CDC Diabetes.", 6.0)
    builder.heading("4.4. Phân chia train / validation / test", 2)
    builder.paragraph(
        "Train set dùng để học weights và fit preprocessing. Validation set dùng cho quyết định hyperparameter hoặc protocol. "
        "Test set được giữ độc lập và chỉ dùng sau khi lựa chọn đã đóng băng. Cách tách này tránh biến test metric thành "
        "một metric lựa chọn lạc quan."
    )
    builder.figure(diagrams["split_roles"], "Vai trò của train, validation và test trong thí nghiệm.", 5.8)
    split_rows = []
    for dataset, prefix in [("EuroSAT", "eurosat"), ("Oxford-IIIT Pet", "oxford_pets"), ("CDC Diabetes", "diabetes")]:
        for split in ["train", "val", "test"]:
            path = RESULTS_DIR / "splits" / f"{prefix}_{split}.csv"
            if path.is_file():
                split_rows.append([dataset, split, fmt_int(len(pd.read_csv(path)))])
    builder.add_table("Số mẫu trong các split thực nghiệm", ["Dataset", "Split", "Số mẫu"], split_rows, [5.0, 3.0, 3.0])
    builder.heading("4.5. Tiền xử lý dữ liệu", 2)
    builder.paragraph(
        "EuroSAT dùng ảnh RGB 64x64 sẵn có và chuẩn hóa tensor cho Keras. Oxford resize official samples về resolution đã chọn, "
        "decode theo nội dung ảnh bằng Pillow để không loại nhầm PNG-content `.jpg`. Diabetes giữ thứ tự feature gốc, fit "
        "standard scaling trên train split và reshape thành tensor (21, 1) cho Conv1D."
    )
    builder.figure(diagrams["conv1d"], "Biểu diễn 21 predictor thành chuỗi Conv1D (21, 1).", 5.8)
    builder.heading("4.6. Phương pháp lựa chọn siêu tham số", 2)
    builder.paragraph(
        "Mỗi quyết định quan trọng được gán một loại: data-determined, architecture-determined, experimentally selected, "
        "operational bound hoặc controlled protocol choice with theoretical rationale. Khi không thể suy ra trực tiếp từ data "
        "hoặc architecture, assignment dùng validation evidence thay vì chọn theo thói quen."
    )
    builder.add_table(
        "Ví dụ phân loại quyết định tham số trong assignment",
        ["Loại quyết định", "Ví dụ", "Ý nghĩa"],
        [
            ["data-determined", "EuroSAT 64x64 RGB; Diabetes 21 predictors", "Suy ra trực tiếp từ dữ liệu đã audit"],
            ["architecture-determined", "Conv1D input (21, 1); residual projection 1x1", "Bắt buộc để shape và phép toán hợp lệ"],
            ["experimentally selected", "learning rate, batch size, resolution Oxford", "Chọn bằng validation evidence"],
            ["operational bound", "max_epochs, EarlyStopping patience", "Giới hạn CPU/thời gian, không tuyên bố tối ưu"],
        ],
        [4.0, 5.0, 6.5],
    )
    builder.heading("4.7. Các chỉ số đánh giá", 2)
    builder.paragraph(
        "Accuracy đo tỷ lệ dự đoán đúng tổng thể. Precision, recall và F1 đo hiệu quả theo lớp; Macro F1 lấy trung bình không "
        "trọng số theo lớp nên nhạy hơn với minority classes. Confusion matrix cho thấy mô hình nhầm lớp nào sang lớp nào. "
        "Với CDC Diabetes, accuracy đơn lẻ đặc biệt dễ gây hiểu lầm vì class 0 chiếm đa số."
    )


def add_hyperparameter_table(builder: ReportBuilder, dataset_name: str, selected_path: str) -> None:
    selected = read_csv(selected_path)
    rows = [
        [row["decision"], str(row["selected_value"]), row["classification"], row["evidence_or_reason"]]
        for _, row in selected.iterrows()
    ]
    builder.add_table(
        f"Giao thức đã chọn cho {dataset_name}",
        ["Parameter", "Selected value", "Decision type", "Evidence"],
        rows,
        [3.2, 4.0, 3.8, 6.0],
    )


def add_experiment_eurosat(builder: ReportBuilder) -> None:
    models = read_csv("results/metrics/eurosat_models.csv")
    hp = read_csv("results/hyperparameters/eurosat_hyperparameters.csv")
    pairs = read_csv("results/metrics/eurosat_confusion_pairs.csv")
    builder.heading("V. THỰC NGHIỆM TRÊN EUROSAT", 1)
    builder.heading("5.1. Cấu hình và tiền xử lý", 2)
    builder.paragraph(
        "EuroSAT dùng ảnh RGB 64x64, split stratified 70/15/15 với seed 42. Augmentation không được đưa vào final protocol "
        "vì validation evidence không cải thiện so với no augmentation. Bốn kiến trúc cuối dùng cùng optimizer family, "
        "learning rate, batch size, EarlyStopping, preprocessing và metric definitions."
    )
    builder.heading("5.2. Khảo sát learning rate", 2)
    lr = hp[hp["experiment_type"] == "learning_rate"][
        ["candidate_value", "tuning_stage", "val_loss", "val_accuracy", "val_macro_f1", "training_time_seconds"]
    ]
    builder.add_table(
        "EuroSAT - kết quả khảo sát learning rate",
        ["Candidate", "Stage", "Val loss", "Val accuracy", "Val Macro F1", "Time (s)"],
        [[str(r.candidate_value), r.tuning_stage, fmt_float(r.val_loss), fmt_float(r.val_accuracy), fmt_float(r.val_macro_f1), fmt_seconds(r.training_time_seconds)] for r in lr.itertuples()],
        [2.4, 3.2, 2.2, 2.4, 2.6, 2.4],
    )
    builder.figure("results/figures/eurosat/hyperparameters/learning_rate_stage_a.png", "EuroSAT - validation Macro F1 theo learning rate ở Stage A.", 6.2)
    builder.paragraph(
        "Stage A cho thấy 0.01 và 0.001 là hai ứng viên mạnh. Stage B xác nhận 0.01 đạt validation Macro F1 0.7061, cao hơn "
        "0.001 với 0.6537, nên learning rate 0.01 được chọn."
    )
    builder.heading("5.3. Khảo sát batch size", 2)
    bs = hp[hp["experiment_type"] == "batch_size"][
        ["candidate_value", "val_loss", "val_accuracy", "val_macro_f1", "training_time_seconds", "time_per_epoch_seconds"]
    ]
    builder.add_table(
        "EuroSAT - kết quả khảo sát batch size",
        ["Batch size", "Val loss", "Val accuracy", "Val Macro F1", "Time (s)", "Time/epoch (s)"],
        [[str(r.candidate_value), fmt_float(r.val_loss), fmt_float(r.val_accuracy), fmt_float(r.val_macro_f1), fmt_seconds(r.training_time_seconds), fmt_seconds(r.time_per_epoch_seconds)] for r in bs.itertuples()],
        [2.5, 2.3, 2.5, 2.5, 2.3, 2.6],
    )
    builder.figure("results/figures/eurosat/hyperparameters/batch_size_comparison.png", "EuroSAT - so sánh batch size bằng validation metrics.", 6.2)
    builder.paragraph("Batch size 16 đạt validation Macro F1 cao nhất trong nhóm so sánh nên được chọn, dù thời gian mỗi epoch lớn hơn một số batch size khác.")
    builder.heading("5.4. Các siêu tham số được chọn", 2)
    add_hyperparameter_table(builder, "EuroSAT", "results/hyperparameters/eurosat_selected_protocol.csv")
    builder.heading("5.5. Kết quả bốn mô hình", 2)
    builder.add_table(
        "EuroSAT - final test metrics của bốn Conv2D models",
        ["Model", "Accuracy", "Macro precision", "Macro recall", "Macro F1", "Parameters", "Train time (s)", "Inference (s)"],
        metrics_table_rows(models, "eurosat"),
        [3.8, 2.1, 2.4, 2.2, 2.0, 2.4, 2.4, 2.3],
    )
    builder.figure("results/figures/eurosat/eurosat_complexity_tradeoffs.png", "EuroSAT - trade-off giữa hiệu năng, số tham số và thời gian.", 6.3)
    builder.heading("5.6. Learning curves", 2)
    for key, name in [("basic", "BasicCNN2D"), ("alexnet_inspired", "AlexNetInspired2D"), ("vgg_inspired", "VGGInspired2D"), ("resnet_inspired", "ResNetInspired2D")]:
        builder.figure(f"results/figures/eurosat/training_curves/{key}_history.png", f"EuroSAT - learning curve của {name}.", 5.8)
    builder.heading("5.7. Confusion matrix và phân tích lỗi", 2)
    builder.figure("results/figures/eurosat/confusion_matrices/basic_confusion_matrix.png", "EuroSAT - confusion matrix của BasicCNN2D.", 5.8)
    builder.figure("results/figures/eurosat/errors/basic_top_confusion_errors.png", "EuroSAT - ví dụ lỗi từ các cặp nhầm lẫn quan trọng của BasicCNN2D.", 5.8)
    basic_pairs = pairs[pairs["key"] == "basic"].head(8)
    builder.add_table(
        "EuroSAT - các cặp nhầm lẫn thường gặp của BasicCNN2D",
        ["Rank", "True class", "Predicted class", "Count"],
        [[str(r.rank), r.true_class, r.predicted_class, fmt_int(r.count)] for r in basic_pairs.itertuples()],
        [1.8, 4.8, 4.8, 2.0],
    )
    builder.paragraph(
        "BasicCNN2D nhầm nhiều giữa Industrial và Residential, cũng như Highway và River. Đây là các lớp có cấu trúc "
        "thị giác dễ gần nhau trong ảnh vệ tinh: khu dân cư và khu công nghiệp đều có mật độ xây dựng, còn đường và sông "
        "đều có dạng tuyến tính. Ba mô hình sâu hơn collapse về dự đoán AnnualCrop trong nhiều lớp, nên lỗi chính của chúng "
        "là suy biến protocol thay vì nhầm lẫn tinh tế giữa hai lớp."
    )


def add_experiment_oxford(builder: ReportBuilder) -> None:
    models = read_csv("results/metrics/oxford_pets_models.csv")
    hp = read_csv("results/hyperparameters/oxford_pets_hyperparameters.csv")
    pairs = read_csv("results/metrics/oxford_pets_confusion_pairs.csv")
    builder.heading("VI. THỰC NGHIỆM TRÊN OXFORD-IIIT PET", 1)
    builder.heading("6.1. Đặc điểm xử lý ảnh Oxford", 2)
    builder.paragraph(
        "Oxford-IIIT Pet có ảnh tự nhiên với kích thước gốc đa dạng và 37 lớp breed. Pipeline chỉ lấy official samples từ "
        "trainval/test annotations, sau đó split trainval thành train và validation theo stratification với seed 42. "
        "Giải mã ảnh dùng nội dung file, không suy luận format từ đuôi `.jpg`."
    )
    builder.heading("6.2. Khảo sát input resolution", 2)
    res = hp[hp["experiment_type"] == "resolution"][
        ["candidate_value", "tuning_stage", "val_loss", "val_accuracy", "val_macro_f1", "training_time_seconds"]
    ]
    builder.add_table(
        "Oxford - kết quả khảo sát resolution",
        ["Resolution", "Stage", "Val loss", "Val accuracy", "Val Macro F1", "Time (s)"],
        [[str(r.candidate_value), r.tuning_stage, fmt_float(r.val_loss), fmt_float(r.val_accuracy), fmt_float(r.val_macro_f1), fmt_seconds(r.training_time_seconds)] for r in res.itertuples()],
        [2.4, 3.6, 2.1, 2.4, 2.5, 2.4],
    )
    builder.figure("results/figures/oxford_pets/hyperparameters/resolution_stage_a.png", "Oxford - Stage A khảo sát input resolution.", 6.2)
    builder.paragraph(
        "Stage B xác nhận 96x96 có validation Macro F1 0.0191, cao hơn 160x160 với 0.0157, trong khi thời gian huấn luyện "
        "123.06s thấp hơn nhiều so với 254.21s. Do đó 96x96 được chọn theo evidence và ràng buộc CPU."
    )
    builder.heading("6.3. Khảo sát learning rate", 2)
    lr = hp[hp["experiment_type"] == "learning_rate"][
        ["candidate_value", "tuning_stage", "val_loss", "val_accuracy", "val_macro_f1", "training_time_seconds"]
    ]
    builder.add_table(
        "Oxford - kết quả khảo sát learning rate",
        ["Candidate", "Stage", "Val loss", "Val accuracy", "Val Macro F1", "Time (s)"],
        [[str(r.candidate_value), r.tuning_stage, fmt_float(r.val_loss), fmt_float(r.val_accuracy), fmt_float(r.val_macro_f1), fmt_seconds(r.training_time_seconds)] for r in lr.itertuples()],
        [2.4, 3.6, 2.1, 2.4, 2.5, 2.4],
    )
    builder.figure("results/figures/oxford_pets/hyperparameters/learning_rate_stage_a.png", "Oxford - Stage A khảo sát learning rate.", 6.2)
    builder.heading("6.4. Khảo sát batch size", 2)
    bs = hp[hp["experiment_type"] == "batch_size"][
        ["candidate_value", "val_loss", "val_accuracy", "val_macro_f1", "training_time_seconds"]
    ]
    builder.add_table(
        "Oxford - kết quả khảo sát batch size",
        ["Batch size", "Val loss", "Val accuracy", "Val Macro F1", "Time (s)"],
        [[str(r.candidate_value), fmt_float(r.val_loss), fmt_float(r.val_accuracy), fmt_float(r.val_macro_f1), fmt_seconds(r.training_time_seconds)] for r in bs.itertuples()],
        [2.5, 2.4, 2.6, 2.6, 2.4],
    )
    builder.figure("results/figures/oxford_pets/hyperparameters/batch_size_comparison.png", "Oxford - so sánh batch size bằng validation metrics.", 6.2)
    builder.heading("6.5. Training protocol cuối cùng", 2)
    add_hyperparameter_table(builder, "Oxford-IIIT Pet", "results/hyperparameters/oxford_pets_selected_protocol.csv")
    builder.heading("6.6. Kết quả bốn mô hình", 2)
    builder.add_table(
        "Oxford - final test metrics của bốn Conv2D models",
        ["Model", "Accuracy", "Macro precision", "Macro recall", "Macro F1", "Parameters", "Train time (s)", "Inference (s)"],
        metrics_table_rows(models, "oxford"),
        [3.8, 2.1, 2.4, 2.2, 2.0, 2.4, 2.4, 2.3],
    )
    builder.figure("results/figures/oxford_pets/oxford_pets_complexity_tradeoffs.png", "Oxford - trade-off giữa hiệu năng, số tham số và thời gian.", 6.3)
    builder.heading("6.7. Confusion matrix", 2)
    builder.figure("results/figures/oxford_pets/confusion_matrices/basic_confusion_matrix.png", "Oxford - confusion matrix của BasicCNN2D.", 6.1)
    builder.figure("results/figures/oxford_pets/errors/basic_top_confusion_errors.png", "Oxford - ví dụ lỗi của BasicCNN2D từ các cặp nhầm lẫn lớn.", 5.8)
    builder.heading("6.8. Phân tích các giống dễ nhầm", 2)
    basic_pairs = pairs[pairs["key"] == "basic"].head(10)
    builder.add_table(
        "Oxford - các cặp nhầm lẫn thường gặp của BasicCNN2D",
        ["Rank", "True class", "Predicted class", "Count"],
        [[str(r.rank), r.true_class, r.predicted_class, fmt_int(r.count)] for r in basic_pairs.itertuples()],
        [1.8, 5.0, 5.0, 2.0],
    )
    builder.paragraph(
        "Những cặp như pomeranian -> chihuahua hoặc Ragdoll -> Russian_Blue có thể có tương đồng thị giác trong một số mẫu, "
        "nhưng kết quả tổng thể rất thấp cho thấy mô hình compact và số epoch ngắn chưa tách được đặc trưng breed một cách ổn định. "
        "Báo cáo không gán toàn bộ lỗi cho một nguyên nhân hình ảnh duy nhất."
    )


def add_experiment_diabetes(builder: ReportBuilder) -> None:
    models = read_csv("results/metrics/diabetes_models.csv")
    hp = read_csv("results/hyperparameters/diabetes_hyperparameters.csv")
    per_class = read_csv("results/metrics/diabetes_per_class_metrics.csv")
    builder.heading("VII. THỰC NGHIỆM TRÊN CDC DIABETES", 1)
    builder.heading("7.1. Đặc điểm của bài toán tabular", 2)
    builder.paragraph(
        "CNN khai thác tốt cấu trúc cục bộ tự nhiên trong ảnh, nhưng 21 predictor tabular của CDC Diabetes không có "
        "quan hệ lân cận không gian như pixel. Vì vậy Conv1D ở đây là adaptation theo yêu cầu assignment, không phải tuyên bố "
        "rằng CNN là mô hình tối ưu tự nhiên cho tabular classification."
    )
    builder.heading("7.2. Biểu diễn đầu vào Conv1D", 2)
    builder.paragraph("Thứ tự feature giữ nguyên theo CSV. Mỗi dòng được reshape thành chuỗi ngắn (21, 1) trước khi đi qua Conv1D.")
    builder.heading("7.3. Phân tích duplicated records", 2)
    builder.paragraph(
        "Duplicate rows được giữ lại trong thí nghiệm vì dữ liệu khảo sát có nhiều biến rời rạc và không có respondent ID. "
        "Việc xóa duplicate có thể loại bỏ các quan sát hợp lệ có cùng tổ hợp câu trả lời."
    )
    builder.heading("7.4. Phân bố lớp và class imbalance", 2)
    builder.figure("results/figures/diabetes/hyperparameters/scaling_comparison.png", "Diabetes - so sánh preprocessing scaling bằng validation metrics.", 6.2)
    builder.paragraph(
        "Lớp Prediabetes rất nhỏ so với No diabetes. Vì vậy Macro F1 và recall theo lớp quan trọng hơn accuracy tổng thể khi "
        "đánh giá khả năng nhận diện minority classes."
    )
    builder.heading("7.5. Majority-class baseline", 2)
    majority = models[models["key"] == "majority_baseline"].iloc[0]
    builder.paragraph(
        f"Majority-class baseline đạt accuracy {fmt_float(majority['test_accuracy'])}, nhưng Prediabetes recall = "
        f"{fmt_float(majority['test_class1_prediabetes_recall'])} và Diabetes recall = "
        f"{fmt_float(majority['test_class2_diabetes_recall'])}. Điều này chứng minh accuracy cao có thể chỉ phản ánh class imbalance."
    )
    builder.figure("results/figures/diabetes/confusion_matrices/majority_baseline_confusion_matrix.png", "Diabetes - confusion matrix của majority-class baseline.", 5.4)
    builder.heading("7.6. Khảo sát learning rate", 2)
    lr = hp[hp["experiment_type"] == "learning_rate"][
        ["candidate_value", "val_loss", "val_accuracy", "val_macro_f1", "val_class1_prediabetes_recall", "val_class2_diabetes_recall"]
    ]
    builder.add_table(
        "Diabetes - kết quả khảo sát learning rate",
        ["LR", "Val loss", "Val accuracy", "Val Macro F1", "Prediabetes recall", "Diabetes recall"],
        [[str(r.candidate_value), fmt_float(r.val_loss), fmt_float(r.val_accuracy), fmt_float(r.val_macro_f1), fmt_float(r.val_class1_prediabetes_recall), fmt_float(r.val_class2_diabetes_recall)] for r in lr.itertuples()],
        [1.8, 2.1, 2.4, 2.4, 3.0, 2.8],
    )
    builder.figure("results/figures/diabetes/hyperparameters/learning_rate_comparison.png", "Diabetes - so sánh learning rate.", 6.2)
    builder.heading("7.7. Khảo sát batch size", 2)
    bs = hp[hp["experiment_type"] == "batch_size"][
        ["candidate_value", "val_loss", "val_accuracy", "val_macro_f1", "training_time_seconds"]
    ]
    builder.add_table(
        "Diabetes - kết quả khảo sát batch size",
        ["Batch size", "Val loss", "Val accuracy", "Val Macro F1", "Time (s)"],
        [[str(r.candidate_value), fmt_float(r.val_loss), fmt_float(r.val_accuracy), fmt_float(r.val_macro_f1), fmt_seconds(r.training_time_seconds)] for r in bs.itertuples()],
        [2.5, 2.4, 2.6, 2.6, 2.4],
    )
    builder.figure("results/figures/diabetes/hyperparameters/batch_size_comparison.png", "Diabetes - so sánh batch size.", 6.2)
    builder.heading("7.8. Class-weight experiment", 2)
    cw = hp[hp["experiment_type"] == "class_weight"][
        ["candidate_value", "val_accuracy", "val_macro_f1", "val_class1_prediabetes_recall", "val_class2_diabetes_recall"]
    ]
    builder.add_table(
        "Diabetes - so sánh class-weight policy",
        ["Weighting strategy", "Val accuracy", "Val Macro F1", "Prediabetes recall", "Diabetes recall"],
        [[str(r.candidate_value), fmt_float(r.val_accuracy), fmt_float(r.val_macro_f1), fmt_float(r.val_class1_prediabetes_recall), fmt_float(r.val_class2_diabetes_recall)] for r in cw.itertuples()],
        [4.2, 2.5, 2.5, 3.0, 2.8],
    )
    builder.figure("results/figures/diabetes/hyperparameters/class_weight_comparison.png", "Diabetes - ảnh hưởng của class weighting đến minority recall.", 6.2)
    builder.paragraph(
        "Balanced inverse-frequency weighting giảm accuracy validation nhưng tăng Macro F1 và cải thiện recall của Prediabetes/Diabetes, "
        "nên được chọn cho final comparison."
    )
    builder.heading("7.9. Kết quả bốn Conv1D models", 2)
    builder.add_table(
        "Diabetes - final test metrics của baseline và bốn Conv1D models",
        ["Model", "Accuracy", "Macro precision", "Macro recall", "Macro F1", "Weighted F1", "Recall class 1", "Recall class 2", "Params", "Train (s)", "Infer (s)"],
        metrics_table_rows(models, "diabetes"),
        [3.6, 1.7, 2.0, 1.9, 1.7, 1.9, 2.1, 2.0, 1.7, 1.7, 1.7],
    )
    builder.figure("results/figures/diabetes/diabetes_final_metrics.png", "Diabetes - final metrics của các mô hình.", 6.3)
    builder.heading("7.10. Confusion matrix và phân tích lỗi", 2)
    builder.figure("results/figures/diabetes/confusion_matrices/basic_confusion_matrix.png", "Diabetes - confusion matrix của BasicCNN1D.", 5.4)
    builder.add_table(
        "Diabetes - metric theo từng lớp trên test set",
        ["Model", "Class", "Precision", "Recall", "F1", "Support"],
        [
            [r.model_family, r.class_name, fmt_float(r.precision), fmt_float(r.recall), fmt_float(r.f1), fmt_int(r.support)]
            for r in per_class.itertuples()
        ],
        [4.0, 3.0, 2.1, 2.1, 2.1, 2.0],
    )
    builder.paragraph(
        "BasicCNN1D có Macro F1 cao nhất trong các CNN, nhưng Prediabetes recall chỉ 0.2518. AlexNetInspired1D có Prediabetes recall "
        "cao nhất 0.5094 nhưng Diabetes recall thấp hơn. VGGInspired1D và ResNetInspired1D nghiêng nhiều hơn về class 2 Diabetes. "
        "Do đó lựa chọn mô hình phụ thuộc vào trade-off theo lớp, không thể dựa riêng vào accuracy."
    )


def add_overall_comparison(builder: ReportBuilder) -> None:
    builder.heading("VIII. SO SÁNH TỔNG HỢP CÁC KIẾN TRÚC", 1)
    builder.heading("8.1. Kết quả trên EuroSAT", 2)
    builder.paragraph(
        "Trên EuroSAT, BasicCNN2D đạt accuracy 0.7398 và Macro F1 0.7256, vượt xa ba kiến trúc sâu hơn dưới cùng protocol. "
        "AlexNet-inspired, VGG-inspired và ResNet-inspired đều dừng ở accuracy 0.1111 và Macro F1 0.0200, cho thấy tăng depth "
        "không tự động tạo lợi ích khi protocol cố định không phù hợp."
    )
    builder.heading("8.2. Kết quả trên Oxford-IIIT Pet", 2)
    builder.paragraph(
        "Trên Oxford-IIIT Pet, tất cả mô hình đều yếu vì bài toán 37 lớp fine-grained, ảnh đa dạng và budget CPU ngắn. "
        "BasicCNN2D vẫn cao nhất với Macro F1 0.0198; mô hình sâu hơn tăng parameter count và runtime nhưng không cải thiện."
    )
    builder.heading("8.3. Kết quả trên CDC Diabetes", 2)
    builder.paragraph(
        "Trên Diabetes, majority baseline có accuracy cao nhất nhưng bỏ sót hoàn toàn class 1 và class 2. Trong các CNN, "
        "BasicCNN1D có Macro F1 cao nhất, AlexNetInspired1D có Prediabetes recall cao nhất, còn VGG/ResNet nghiêng về Diabetes recall."
    )
    builder.heading("8.4. Performance và độ phức tạp", 2)
    for rel, caption in [
        ("results/figures/model_comparison/accuracy_by_dataset.png", "Tổng hợp accuracy theo dataset và kiến trúc."),
        ("results/figures/model_comparison/macro_f1_by_dataset.png", "Tổng hợp Macro F1 theo dataset và kiến trúc."),
        ("results/figures/model_comparison/parameters_vs_macro_f1.png", "Parameter count so với Macro F1."),
        ("results/figures/model_comparison/training_time_vs_macro_f1.png", "Training time so với Macro F1."),
    ]:
        builder.figure(rel, caption, 6.3)
    builder.heading("8.5. Ảnh hưởng của sự phát triển kiến trúc", 2)
    builder.paragraph(
        "Bằng chứng trong assignment không ủng hộ kết luận rằng kiến trúc sâu hơn luôn tốt hơn. Basic CNN cạnh tranh hoặc dẫn đầu "
        "trong cả ba dataset theo Macro F1. Residual connections có ý nghĩa kiến trúc rõ ràng, nhưng trong điều kiện CPU-only, số epoch "
        "ngắn và protocol không retune riêng cho từng family, ResNet-inspired không tạo lợi thế test rõ ràng. Điều này phù hợp với mục tiêu "
        "assignment: hiểu trade-off kiến trúc thay vì tìm một winner phổ quát."
    )


def add_discussion_conclusion_refs(builder: ReportBuilder) -> None:
    builder.heading("IX. THẢO LUẬN", 1)
    builder.heading("9.1. Các kết quả đáng chú ý", 2)
    builder.paragraph(
        "Kết quả quan trọng nhất là độ phức tạp kiến trúc không bảo đảm tăng performance. Trên hai image datasets, mô hình sâu hơn không "
        "cải thiện trong protocol đóng băng; trên Diabetes, các kiến trúc tạo trade-off khác nhau giữa Prediabetes recall và Diabetes recall."
    )
    builder.heading("9.2. Trade-off giữa performance và computational cost", 2)
    builder.paragraph(
        "Các mô hình 2D sâu hơn có parameter count và training time cao hơn rõ rệt. Trong EuroSAT, ResNetInspired2D có 324,490 tham số và "
        "training time 1467.59s nhưng Macro F1 chỉ 0.0200, trong khi BasicCNN2D có 24,202 tham số và Macro F1 0.7256. Với Oxford, xu hướng "
        "tương tự xuất hiện dù mọi mô hình đều thấp. Trên Diabetes, chi phí tăng ít hơn nhưng Macro F1 vẫn không tăng tương ứng."
    )
    builder.heading("9.3. Ảnh hưởng của đặc điểm dataset", 2)
    builder.paragraph(
        "EuroSAT có ảnh 64x64 tương đối đồng nhất nên Basic CNN đủ học một số tín hiệu mạnh. Oxford là fine-grained classification với "
        "37 giống thú cưng, ảnh tự nhiên đa dạng và resolution gốc khác nhau, nên mô hình from-scratch compact khó đạt hiệu quả cao trong "
        "budget ngắn. Diabetes là tabular dataset mất cân bằng mạnh, vì vậy Macro F1 và per-class recall quan trọng hơn raw accuracy."
    )
    builder.heading("9.4. Hạn chế", 2)
    builder.bullets(
        [
            "CPU-only training giới hạn số epoch, kích thước mô hình và phạm vi hyperparameter search.",
            "Các kiến trúc là scaled inspired variants, không phải bản lịch sử đầy đủ.",
            "Hyperparameter search có validation evidence nhưng không exhaustive.",
            "Conv1D trên tabular features không có locality tự nhiên như ảnh.",
            "Oxford-IIIT Pet là bài toán fine-grained khó, nên kết quả thấp cần được đọc trong ràng buộc from-scratch và CPU-aware.",
        ]
    )
    builder.heading("X. KẾT LUẬN", 1)
    builder.paragraph(
        "CNN có thể được hiểu như hợp của các hàm convolution, activation, pooling và classifier. Sự phát triển kiến trúc thay đổi cách "
        "các hàm này được tổ chức: Basic CNN đặt baseline, AlexNet-inspired tăng capacity, VGG-inspired hệ thống hóa repeated small kernels, "
        "và ResNet-inspired thêm residual connections. Tuy nhiên, kết quả thực nghiệm cho thấy độ phức tạp cao hơn không tự động đem lại "
        "performance tốt hơn. Hyperparameter selection cần dựa trên validation evidence, còn test set phải được giữ độc lập. Đặc điểm dataset "
        "và ràng buộc tính toán ảnh hưởng mạnh đến kết quả cuối."
    )
    builder.heading("TÀI LIỆU THAM KHẢO", 1)
    refs = [
        "[1] P. Helber, B. Bischke, A. Dengel, and D. Borth, EuroSAT: A Novel Dataset and Deep Learning Benchmark for Land Use and Land Cover Classification, IEEE JSTARS, 2019. URL: https://github.com/phelber/EuroSAT",
        "[2] O. M. Parkhi, A. Vedaldi, A. Zisserman, and C. V. Jawahar, Cats and Dogs, IEEE CVPR, 2012. URL: https://www.robots.ox.ac.uk/~vgg/data/pets/",
        "[3] UCI Machine Learning Repository, CDC Diabetes Health Indicators. URL: https://archive.ics.uci.edu/dataset/891/cdc+diabetes+health+indicators",
        "[4] Centers for Disease Control and Prevention, Behavioral Risk Factor Surveillance System 2015. URL: https://www.cdc.gov/brfss/annual_data/annual_2015.html",
        "[5] A. Krizhevsky, I. Sutskever, and G. E. Hinton, ImageNet Classification with Deep Convolutional Neural Networks, NeurIPS, 2012.",
        "[6] K. Simonyan and A. Zisserman, Very Deep Convolutional Networks for Large-Scale Image Recognition, arXiv:1409.1556, 2014.",
        "[7] K. He, X. Zhang, S. Ren, and J. Sun, Deep Residual Learning for Image Recognition, IEEE CVPR, 2016.",
        "[8] TensorFlow and Keras documentation. URL: https://www.tensorflow.org/guide/keras",
    ]
    for ref in refs:
        paragraph = builder.document.add_paragraph(ref)
        paragraph.paragraph_format.first_line_indent = Cm(0)


def add_appendices(builder: ReportBuilder) -> None:
    builder.heading("PHỤ LỤC A. MỘT SỐ ĐOẠN MÃ XÂY DỰNG KIẾN TRÚC CNN", 1)
    builder.paragraph("Phụ lục này chỉ trích đoạn ngắn có giá trị giải thích, không sao chép toàn bộ source file.")
    builder.code(
        "Classification head 2D dùng Global Average Pooling để giảm parameter count.",
        """
def _classification_head_2d(x, num_classes, dense_units, name):
    # Global Average Pooling tránh Flatten lớn khi feature map còn nhiều vị trí.
    x = layers.GlobalAveragePooling2D(name=f"{name}_gap")(x)
    x = layers.Dense(dense_units, activation="relu", name=f"{name}_dense")(x)
    return layers.Dense(num_classes, activation="softmax", name=f"{name}_probabilities")(x)
""",
    )
    builder.code(
        "Progression shape Conv1D được gắn vào model để kiểm tra độ dài chuỗi.",
        """
def _attach_shape_progression(model, progression):
    # Lưu shape progression để audit Conv1D không bị collapse chiều dài.
    model.shape_progression = progression
    return model
""",
    )
    builder.heading("PHỤ LỤC B. CÁC HÀM HỖ TRỢ TRAINING/EVALUATION QUAN TRỌNG", 1)
    builder.code(
        "Loader ảnh decode theo nội dung và convert RGB trong bộ nhớ.",
        """
with Image.open(resolved) as image:
    # detected_format đến từ nội dung file, không đến từ phần mở rộng.
    detected_format = image.format
    original_mode = image.mode
    original_size = image.size
    rgb = image.convert("RGB")
    if image_size is not None:
        height, width = image_size
        rgb = rgb.resize((width, height), resample=resample)
    array = np.asarray(rgb, dtype=np.uint8)
""",
    )
    builder.code(
        "Tính metrics phân loại macro từ confusion matrix.",
        """
precision = tp / (tp + fp) if tp + fp else 0.0
recall = tp / (tp + fn) if tp + fn else 0.0
f1 = 2 * precision * recall / (precision + recall) if precision + recall else 0.0
macro_f1 = sum(f1s) / len(f1s) if f1s else 0.0
""",
    )
    builder.heading("PHỤ LỤC C. MÔI TRƯỜNG THỰC NGHIỆM VÀ KHẢ NĂNG TÁI LẬP", 1)
    builder.add_table(
        "Môi trường đã dùng cho report generation",
        ["Thành phần", "Giá trị"],
        [
            ["Python executable", sys.executable],
            ["Python version", sys.version.split()[0]],
            ["TensorFlow environment", "C:/Users/anhca/anaconda3/envs/tf312/python.exe"],
            ["Seed policy", "42 cho Python, NumPy và TensorFlow/Keras trong thí nghiệm"],
            ["GPU", "Không yêu cầu; CPU-only là chủ đích"],
        ],
        [5.0, 10.0],
    )
    builder.paragraph(
        "Bảng dưới đây được giữ ở phụ lục như một thông tin kiểm tra khả năng tái lập, "
        "không phải là kết quả khoa học chính của báo cáo."
    )
    builder.add_table(
        "Tổng quan notebooks nguồn đã kiểm tra",
        ["Notebook", "Markdown cells", "Code cells", "Error outputs"],
        notebook_overview(),
        [7.6, 2.4, 2.2, 2.2],
    )


def validate_docx(path: Path, builder: ReportBuilder) -> dict[str, object]:
    if not path.is_file():
        raise FileNotFoundError(path)
    doc = Document(str(path))
    text = "\n".join(p.text for p in doc.paragraphs)
    required_headings = [
        "I. GIỚI THIỆU",
        "II. CƠ SỞ LÝ THUYẾT VỀ CNN",
        "III. SỰ PHÁT TRIỂN CỦA CÁC KIẾN TRÚC CNN",
        "IV. DỮ LIỆU VÀ THIẾT KẾ THỰC NGHIỆM",
        "V. THỰC NGHIỆM TRÊN EUROSAT",
        "VI. THỰC NGHIỆM TRÊN OXFORD-IIIT PET",
        "VII. THỰC NGHIỆM TRÊN CDC DIABETES",
        "VIII. SO SÁNH TỔNG HỢP CÁC KIẾN TRÚC",
        "IX. THẢO LUẬN",
        "X. KẾT LUẬN",
        "TÀI LIỆU THAM KHẢO",
    ]
    missing_headings = [heading for heading in required_headings if heading not in text]
    suspicious = [token for token in ["�", "ï¿½", "Ãƒ", "Ã„", "Ã¡Âº", "áº"] if token in text]
    with zipfile.ZipFile(path) as archive:
        image_count = len([name for name in archive.namelist() if name.startswith("word/media/")])
    if missing_headings:
        raise RuntimeError("Missing expected headings: " + ", ".join(missing_headings))
    if suspicious:
        raise RuntimeError("Suspicious mojibake tokens found: " + ", ".join(suspicious))
    table_caption_count = len(re.findall(r"Bảng \d+\.", text))
    if table_caption_count < builder.table_count:
        raise RuntimeError("DOCX table caption count is lower than expected.")
    if image_count < builder.figure_count:
        raise RuntimeError("Embedded image count is lower than expected.")
    return {
        "paragraphs": len(doc.paragraphs),
        "tables": len(doc.tables),
        "table_captions": table_caption_count,
        "inline_shapes": len(doc.inline_shapes),
        "embedded_images": image_count,
        "figures": builder.figure_count,
        "tables_expected": builder.table_count,
        "code_excerpts": builder.code_count,
        "references": 8,
        "chapters": 10,
        "language_encoding_ok": True,
    }


def trim_trailing_empty_paragraphs(document: Document) -> None:
    """Remove extra empty paragraphs that can create a blank final page."""

    body = document._body._element
    while len(document.paragraphs) > 1:
        paragraph = document.paragraphs[-1]
        if paragraph.text.strip():
            break
        if paragraph.runs:
            break
        body.remove(paragraph._element)


def build_report() -> dict[str, object]:
    ASSETS_DIR.mkdir(parents=True, exist_ok=True)
    diagrams = create_conceptual_diagrams()
    audit = read_json("results/dataset_audit.json")

    builder = ReportBuilder()
    builder.add_cover()
    builder.add_front_matter()
    add_intro(builder)
    add_theory(builder, diagrams)
    add_architecture_chapter(builder, diagrams)
    add_data_and_method(builder, diagrams, audit)
    add_experiment_eurosat(builder)
    add_experiment_oxford(builder)
    add_experiment_diabetes(builder)
    add_overall_comparison(builder)
    add_discussion_conclusion_refs(builder)
    add_appendices(builder)

    trim_trailing_empty_paragraphs(builder.document)
    builder.document.save(str(OUTPUT_PATH))
    validation = validate_docx(OUTPUT_PATH, builder)
    validation.update(
        {
            "output_path": str(OUTPUT_PATH.relative_to(ROOT)),
            "script_path": str(SCRIPT_PATH.relative_to(ROOT)),
            "python_executable": sys.executable,
            "figure_count": builder.figure_count,
            "table_count": builder.table_count,
            "code_count": builder.code_count,
        }
    )
    (REPORT_DIR / "A05_Assignment_Report_validation.json").write_text(
        json.dumps(validation, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    return validation


if __name__ == "__main__":
    result = build_report()
    print(json.dumps(result, ensure_ascii=False, indent=2))
