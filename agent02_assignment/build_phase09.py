"""Create editable black-text Word candidate from the sealed AG02 manuscript."""

import argparse
import json
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image

from .code.publication_v1 import blocks, digest, display_heading, local_file, plain_inline

BASE = Path(__file__).resolve().parent
STEM = 'agent02_D23CTPM01_TrieuTuanAnh'
BLACK = RGBColor(0, 0, 0)


def font(run, size=None, *, code=False):
    name = 'Consolas' if code else 'Times New Roman'
    run.font.name = name
    run.font.color.rgb = BLACK
    if size:
        run.font.size = Pt(size)
    properties = run._element.get_or_add_rPr()
    fonts = properties.get_or_add_rFonts()
    for attr in list(fonts.attrib):
        if 'theme' in attr.lower():
            del fonts.attrib[attr]
    for key in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
        fonts.set(qn('w:' + key), name)
    return run


def inline(paragraph, text, size=None):
    for part in re.split(r'(\*\*[^*]+\*\*|`[^`]+`|\*[^*]+\*)', text):
        if not part:
            continue
        bold = part.startswith('**') and part.endswith('**')
        code = part.startswith('`') and part.endswith('`')
        italic = part.startswith('*') and part.endswith('*') and not bold
        value = part[2:-2] if bold else part[1:-1] if code or italic else part
        run = font(paragraph.add_run(value), 10 if code else size, code=code)
        run.bold = bold
        run.italic = italic


def configure(doc):
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.left_margin, section.right_margin = Cm(2.3), Cm(2)
    section.top_margin, section.bottom_margin = Cm(2), Cm(2)
    section.footer_distance = Cm(.9)
    for style in doc.styles:
        if style.type in (1, 2):
            style.font.name = 'Times New Roman'
            style.font.color.rgb = BLACK
            fonts = style._element.get_or_add_rPr().get_or_add_rFonts()
            fonts.attrib.clear()
            for attr in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
                fonts.set(qn('w:' + attr), 'Times New Roman')
        for border in list(style._element.iter(qn('w:pBdr'))):
            border.getparent().remove(border)
    normal = doc.styles['Normal']
    normal.font.size = Pt(12.5)
    normal.paragraph_format.line_spacing = 1.12
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.widow_control = True
    for name, size in [('Title', 18), ('Heading 1', 15), ('Heading 2', 13), ('Heading 3', 12.5)]:
        style = doc.styles[name]
        style.font.size = Pt(size)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(9)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_with_next = True
    doc.styles['Caption'].font.size = Pt(11)
    doc.styles['Caption'].font.italic = False
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    r = OxmlElement('w:r')
    t = OxmlElement('w:t')
    t.text = '1'
    r.append(t)
    field.append(r)
    footer._p.append(field)
    doc.core_properties.author = 'Triệu Tuấn Anh'
    doc.core_properties.title = 'So sánh tác tử học tăng cường và tác tử theo luật trong mô phỏng thương mại điện tử'
    doc.core_properties.subject = 'Agent02'


def add_table(doc, rows):
    count = len(rows[0])
    table = doc.add_table(rows=0, cols=count)
    table.autofit = False
    # Numeric tables reserve more room for the configuration than for numeric columns.
    weights = [1.3, *([1] * (count - 1))] if count > 2 else [1, 2]
    widths = [16.7 * value / sum(weights) for value in weights]
    for cell, width in zip(table.columns, widths):
        cell.width = Cm(width)
    for index, values in enumerate(rows):
        row = table.add_row()
        props = row._tr.get_or_add_trPr()
        props.append(OxmlElement('w:cantSplit'))
        if index == 0:
            props.append(OxmlElement('w:tblHeader'))
        for cell, width, value in zip(row.cells, widths, values):
            cell.width = Cm(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_before = Pt(2)
            paragraph.paragraph_format.space_after = Pt(2)
            paragraph.paragraph_format.line_spacing = 1.02
            paragraph.paragraph_format.keep_with_next = index < len(rows) - 1
            inline(paragraph, value, size=10.5 if count >= 5 else 11)
            for run in paragraph.runs:
                run.bold = index == 0
            if index == 0:
                fill = OxmlElement('w:shd')
                fill.set(qn('w:fill'), 'EEEEEE')
                cell._tc.get_or_add_tcPr().append(fill)
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        border = OxmlElement('w:' + edge)
        for name, value in {'val': 'single', 'sz': '4', 'color': 'D9D9D9'}.items():
            border.set(qn('w:' + name), value)
        borders.append(border)
    table._tbl.tblPr.append(borders)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_code(doc, text):
    # Single paragraph retains source line breaks exactly; Word wraps long source lines.
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(5)
    paragraph.paragraph_format.line_spacing = 1
    paragraph.paragraph_format.keep_together = True
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.left_indent = Cm(.2)
    font(paragraph.add_run(text), 9, code=True)


def build(revision):
    if not re.fullmatch(r'r[1-9][0-9]*', revision):
        raise ValueError('Expected a new rN candidate revision')
    active = json.loads((BASE / 'report/active_report.json').read_text())
    manuscript = local_file(BASE, active['manuscript'])
    if digest(manuscript) != active['manuscript_sha256']:
        raise ValueError('Content source changed')
    destination = BASE / 'report/build' / ('candidate-' + revision)
    destination.mkdir(parents=True, exist_ok=False)
    output = destination / (STEM + '.docx')
    items = blocks(manuscript.read_text(encoding='utf-8'))
    doc = Document()
    configure(doc)
    cover = [
        ('HỌC VIỆN CÔNG NGHỆ BƯU CHÍNH VIỄN THÔNG', 14),
        ('PHÁT TRIỂN HỆ THỐNG THÔNG MINH', 13),
        ('BÀI THỰC NGHIỆM AGENT02', 15),
        ('SO SÁNH TÁC TỬ HỌC TĂNG CƯỜNG VÀ TÁC TỬ THEO LUẬT TRONG MÔ PHỎNG THƯƠNG MẠI ĐIỆN TỬ', 18),
        ('Sinh viên  Triệu Tuấn Anh', 13), ('Mã sinh viên  B23DCCN053', 13),
        ('Lớp  D23CTPM01', 13), ('Bài nộp cá nhân', 13), ('Tháng 10 năm 2026', 13)]
    for index, (text, size) in enumerate(cover):
        p = doc.add_paragraph(style='Title' if index == 3 else None)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(25 if index in (2, 4, 8) else 8)
        p.paragraph_format.space_after = Pt(12)
        font(p.add_run(text), size).bold = index < 4
    started = False
    section_number = 0
    image_receipts = []
    for index, item in enumerate(items):
        if not started:
            if item['kind'] == 'heading' and item['text'].startswith('2.'):
                started = True
            else:
                continue
        kind = item['kind']
        if kind == 'heading':
            if item['level'] == 2:
                section_number = int(item['text'].split('.')[0])
            # Manuscript H2 maps to Word Heading1; major sections start a page.
            p = doc.add_heading(display_heading(item['text']), level=item['level'] - 1)
            if item['level'] == 2 and item['text'].startswith(('2.', '6.', '9.')):
                p.paragraph_format.page_break_before = True
        elif kind == 'table':
            add_table(doc, item['rows'])
        elif kind == 'code':
            add_code(doc, item['text'])
        elif kind == 'image':
            relative = (manuscript.parent / item['path']).resolve().relative_to(BASE).as_posix()
            path = local_file(BASE, relative)
            with Image.open(path) as picture:
                width, height = picture.size
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next = True
            # Cap tall two-panel figures to leave room for the explanation.
            cm_width = min(16.5, 12.0 * width / height)
            drawing = p.add_run().add_picture(str(path), width=Cm(cm_width))
            drawing._inline.docPr.set('descr', item['id'])
            image_receipts.append({'id': item['id'], 'path': relative, 'sha256': digest(path), 'width_cm': cm_width})
        else:
            text = item['text']
            caption = plain_inline(text).startswith(('Hình F', 'Bảng B'))
            p = doc.add_paragraph(style='Caption' if caption else None)
            inline(p, text, size=11.5 if section_number == 9 else 11 if caption or text.startswith('Nguồn E') else None)
            if section_number == 9:
                p.paragraph_format.line_spacing = 1.04
                p.paragraph_format.space_after = Pt(4)
            if caption:
                p.paragraph_format.keep_with_next = True
            if text.startswith('Nguồn E'):
                p.paragraph_format.keep_together = True
            if index > 0 and items[index - 1]['kind'] == 'image':
                p.paragraph_format.keep_together = True
            if index + 1 < len(items) and items[index + 1]['kind'] == 'code':
                p.paragraph_format.keep_with_next = True
    # Remove residual theme colors, including hyperlinks and default table styles.
    for root in (doc._element, doc.styles.element, doc.sections[0].footer._element):
        for color in root.iter(qn('w:color')):
            color.attrib.clear()
            color.set(qn('w:val'), '000000')
    doc.save(output)
    receipt = {'status': 'BUILT_PENDING_RENDER_QA', 'revision': revision, 'docx': output.relative_to(BASE).as_posix(),
               'sha256': digest(output), 'manuscript_sha256': digest(manuscript), 'native_tables': len(doc.tables),
               'embedded_figures': len(image_receipts), 'images': image_receipts,
               'source_code_excerpts': len([b for b in items if b['kind'] == 'code']),
               'submission_mode': 'individual', 'metadata_confirmed_by_user': True,
               'scientific_forward_added': 0, 'train_runs_added': 0}
    (destination / 'build_receipt.json').write_text(json.dumps(receipt, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    print(json.dumps({k: v for k, v in receipt.items() if k != 'images'}, indent=2))


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--revision', required=True)
    build(parser.parse_args().revision)
