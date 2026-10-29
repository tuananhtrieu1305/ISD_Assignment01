"""Inspect DOCX/PDF candidate structure and real page geometry before visual review."""

import argparse
import json
import re
import unicodedata
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

import fitz
from docx import Document
from PIL import Image, ImageChops

from .code.publication_v1 import digest, display_heading

BASE = Path(__file__).resolve().parent
STEM = 'agent02_D23CTPM01_TrieuTuanAnh'
W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'


def inspect(revision):
    docx = BASE / f'report/build/candidate-{revision}/{STEM}.docx'
    qa = BASE / f'report/qa/phase09/{revision}'
    pdf = qa / (STEM + '.pdf')
    doc = Document(docx)
    excerpts = json.loads((BASE / 'report/phase08/code_excerpt_index.json').read_text())
    for excerpt in excerpts:
        if sum(p.text == excerpt['text'] for p in doc.paragraphs) != 1:
            raise ValueError('DOCX source excerpt altered or missing: ' + excerpt['id'])
    if len(doc.tables) != 14 or len(doc.inline_shapes) != 20:
        raise ValueError('DOCX table/figure coverage incomplete')
    with zipfile.ZipFile(docx) as archive:
        color_values = []
        for member in ('word/document.xml', 'word/styles.xml', 'word/footer1.xml'):
            root = ET.fromstring(archive.read(member))
            color_values += [node.attrib.get(W + 'val') for node in root.iter(W + 'color')]
        if set(color_values) != {'000000'}:
            raise ValueError('Nonblack DOCX text')
    reader = fitz.open(pdf)
    pages = []
    text_color = set()
    glyph_errors = []
    for index, page in enumerate(reader):
        text = unicodedata.normalize('NFC', page.get_text())
        if len(text.strip()) < 80:
            raise ValueError('Blank or near-empty page:' + str(index + 1))
        spans = [span for block in page.get_text('dict')['blocks'] if 'lines' in block
                 for line in block['lines'] for span in line['spans']]
        text_color.update(span['color'] for span in spans)
        for span in spans:
            x0, y0, x1, y1 = span['bbox']
            if x0 < 15 or y0 < 15 or x1 > page.rect.width - 15 or y1 > page.rect.height - 15:
                glyph_errors.append({'page': index + 1, 'text': span['text'], 'bbox': span['bbox']})
            if '\ufffd' in span['text']:
                raise ValueError('Missing glyph on page:' + str(index + 1))
        image = qa / f'page-{index + 1}.png'
        if not image.is_file():
            raise ValueError('Missing page raster')
        with Image.open(image) as raster:
            r, g, b = raster.convert('RGB').split()
            if ImageChops.difference(r, g).getbbox() or ImageChops.difference(r, b).getbbox():
                raise ValueError('Non-grayscale rendered page:' + str(index + 1))
        pages.append({'page': index + 1, 'text_characters': len(text),
                      'beginning': text[:100], 'ending': text[-110:],
                      'image_occurrences': len(page.get_image_info()), 'png_sha256': digest(image)})
    if text_color != {0} or glyph_errors:
        raise ValueError('PDF text color or page-boundary overflow: ' + str(glyph_errors[:3]))
    fulltext = '\n'.join(page.get_text() for page in reader)
    normalized = re.sub(r'\s+', '', unicodedata.normalize('NFC', fulltext))
    if sum(p['image_occurrences'] for p in pages) != 20:
        raise ValueError('Actual PDF figure placements differ from twenty')
    questions = []
    manuscript = (BASE / 'report/phase08/manuscript.md').read_text(encoding='utf-8')
    for heading in re.findall(r'^### (6\.\d+\. Câu \d+[^\n]*)', manuscript, re.MULTILINE):
        simple = re.sub(r'\s+', '', display_heading(heading))
        if simple not in normalized:
            raise ValueError('Question missing from actual PDF:' + heading)
        questions.append(heading)
    result = {'status': 'STRUCTURAL_PASS_PENDING_VISUAL_REVIEW', 'revision': revision,
              'docx': docx.relative_to(BASE).as_posix(), 'docx_sha256': digest(docx),
              'pdf': pdf.relative_to(BASE).as_posix(), 'pdf_sha256': digest(pdf),
              'page_count': len(reader), 'native_tables': 14, 'figure_occurrences': sum(p['image_occurrences'] for p in pages),
              'exact_code_excerpts': 8, 'questions_in_pdf': len(questions),
              'docx_text_colors': sorted(set(color_values)), 'pdf_text_colors': sorted(text_color),
              'all_page_pixels_grayscale': True, 'page_boundary_violations': glyph_errors,
              'pages': pages}
    (qa / 'structure.json').write_text(json.dumps(result, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    print(json.dumps({k: v for k, v in result.items() if k != 'pages'}, ensure_ascii=False, indent=2))
    print(json.dumps(pages, ensure_ascii=False, indent=2))


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--revision', required=True)
    inspect(parser.parse_args().revision)
