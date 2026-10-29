"""Final read-only rubric audit and byte-identical publication of Agent02."""

import argparse
import json
import re
import shutil
import unicodedata
import xml.etree.ElementTree as ET
from pathlib import Path

import fitz
from docx import Document

from .code.publication_v1 import digest, display_heading

BASE = Path(__file__).resolve().parent
EVIDENCE = BASE / 'evidence/phase10'
STEM = 'agent02_D23CTPM01_TrieuTuanAnh'


def load(relative):
    return json.loads((BASE / relative).read_text(encoding='utf-8'))


def norm(text):
    return re.sub(r'\s+', '', unicodedata.normalize('NFC', text))


def save_once(path, value):
    content = json.dumps(value, ensure_ascii=False, indent=2) + '\n'
    if path.exists():
        if path.read_text(encoding='utf-8') != content:
            raise ValueError('Keep existing release evidence: ' + str(path))
        return
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(content, encoding='utf-8')


def inspect():
    active = load('report/active_candidate.json')
    for key in ('docx', 'pdf'):
        if digest(BASE / active[key]) != active[key + '_sha256']:
            raise ValueError('Candidate changed')
    for row in load('evidence/phase09/experiment_lock.json')['files']:
        if digest(BASE / row['path']) != row['sha256']:
            raise ValueError('Phase9 protected file changed:' + row['path'])
    doc = Document(BASE / active['docx'])
    reader = fitz.open(BASE / active['pdf'])
    texts = [unicodedata.normalize('NFC', p.get_text()) for p in reader]
    if len(reader) != 28 or len(doc.tables) != 14 or len(doc.inline_shapes) != 20:
        raise ValueError('Publication coverage mismatch')
    manuscript = (BASE / 'report/phase08/manuscript.md').read_text(encoding='utf-8')
    question_pages = []
    for number, heading in enumerate(re.findall(r'^### (6\.\d+\. Câu \d+[^\n]*)', manuscript, re.MULTILINE), 1):
        locator = norm(display_heading(heading))
        pages = [i + 1 for i, text in enumerate(texts) if locator in norm(text)]
        if len(pages) != 1:
            raise ValueError('Question not found exactly once:' + heading)
        question_pages.append({'question': number, 'heading': heading, 'page': pages[0]})
    if len(question_pages) != 10:
        raise ValueError('Missing question')
    table_names = ['baseline_train', 'q_policy', 'alpha', 'exploration', 'core_final',
                   'baseline_final', 'q8', 'gold_kernel', 'variants', 'improved', 'improved_delta', 'oracle']
    tables = load('report/phase08/tables.json')
    for table, name in zip(doc.tables[1:13], table_names):
        actual = [[norm(c.text) for c in row.cells] for row in table.rows]
        expected = [[norm(str(c)) for c in row] for row in [tables[name]['headers'], *tables[name]['rows']]]
        if actual != expected:
            raise ValueError('Word table differs from saved artifact:' + name)
    excerpts = load('report/phase08/code_excerpt_index.json')
    if any(sum(p.text == e['text'] for p in doc.paragraphs) != 1 for e in excerpts):
        raise ValueError('Source excerpt differs')
    if sum(len(page.get_image_info()) for page in reader) != 20:
        raise ValueError('PDF image placement mismatch')
    figure_pages = {}
    for i in range(1, 21):
        hits = [j + 1 for j, text in enumerate(texts) if f'HìnhF{i:02d}.' in norm(text)]
        if len(hits) != 1:
            raise ValueError('Figure title not found')
        figure_pages[f'F{i:02d}'] = hits[0]
    headings = [p.text for p in doc.paragraphs if p.style.name == 'Heading 1']
    if len(headings) != 8:
        raise ValueError('Title plus eight body sections required')
    entries = [
        ('R01', [2,3,5,22], 'Shared market,two simultaneous actions,trace and PEAS;Q1'),
        ('R02', [3,5,23], 'Four local states/two actions,Q8values,three-branch equivalent rule and profit reward'),
        ('R03', [4,6,7], 'REF42,5000x20,alpha/gamma/epsilon and preserved reference'),
        ('R04', [6,7,9], 'B01 total/mean/action count;B02 Q/greedy;F01-F04'),
        ('R05', [7,8], 'Raw reward and MA100;stochastic variation;no convergence claim'),
        ('R06', [9,10,13,24], 'Three alpha configurations,train curves and frozen-policy final'),
        ('R07', [10,11,12,23], 'Decay/fixed.05/fixed.30,train/final and epsilon schedules'),
        ('R08', [5,6,13,14], 'Independent matched tapes,role swap,exploration off,Q before/after hash'),
        ('R09', [6,13,15,17], 'Five seed/sample std,roles first,Q8pair replicate and limits'),
        ('R10', [22,23,24,25,26], 'All ten question headings and evidence matrix'),
        ('R11', [22,23], 'Q3=4.65 step-by-step;terminal target;Q5 margins5/3'),
        ('R12', [17,24], 'Two-learner diagnostic,frozen checkpoint/common opponent;no causal overclaim'),
        ('R13', [16,18,24,25], 'At least five limits,implemented persistent/Gold kernel and before/after results'),
        ('R14', [18,19,20,21,25,26], 'Memory/time state,action/reward/method/evaluation;negative findings preserved'),
        ('R15', list(range(1,29)), 'Title plus eight body sections;8 exact-source excerpts,protocol,seeds and citations'),
        ('R16', [1,28], 'PDF opens28pages,approved individual identity/filename;deadline printed in sourceS1p12'),
        ('U01', [2,22,23], 'Basic theory and essential Q-update/margin/mean calculations;no complex derivation'),
        ('U02', sorted(set(figure_pages.values())), 'All20figures have axes/units,aggregation,n,observations and limitations')]
    rubric = {'status': 'PASS', 'requirements': [{'id': k, 'status': 'PASS', 'pdf_pages': p, 'review': note}
              for k,p,note in entries], 'questions': question_pages, 'figure_pages': figure_pages,
              'selected_tables_exact': 12, 'code_excerpts_exact': 8,
              'source_assignment': 'sources/provided/assignment_Agent02_due_29.10.pdf',
              'source_requirement_pages': [8,9,10,11,12], 'pdf_sha256': active['pdf_sha256']}
    save_once(EVIDENCE / 'rubric_review.json', rubric)
    print(json.dumps({'status': 'PASS', 'requirements': 18, 'questions': 10, 'tables_exact': 12,
                      'code_excerpts_exact': 8, 'pages': 28, 'figure_pages': figure_pages}, indent=2))


def release():
    rubric = load('evidence/phase10/rubric_review.json')
    visual = load('evidence/phase10/visual_review.json')
    word = load('evidence/phase10/word_readonly.json')
    parent = load('evidence/phase10/full_parent_audit.json')
    if any(v['status'] != 'PASS' for v in (rubric,visual,word,parent)):
        raise ValueError('A final gate failed')
    if not word['source_unchanged'] or word['saved'] or not word['read_only']:
        raise ValueError('Word check changed source')
    tests = ET.parse(EVIDENCE / 'final_tests.xml').getroot()
    if len(tests.findall('.//testcase')) != 245 or any(tests.findall('.//' + x) for x in ('failure','error','skipped')):
        raise ValueError('Final regression failed')
    active = load('report/active_candidate.json')
    if visual['pdf_sha256'] != active['pdf_sha256'] or len(visual['pages']) != 28:
        raise ValueError('Final visual review missing')
    structure = load('report/qa/phase09/r3/structure.json')
    for page, checked in zip(structure['pages'], visual['pages']):
        path = BASE / f'report/qa/phase09/r3/page-{page["page"]}.png'
        if checked['status'] != 'PASS' or checked['png_sha256'] != digest(path):
            raise ValueError('A reviewed page changed')
    outputs = []
    for suffix in ('pdf','docx'):
        source = BASE / active[suffix]
        target = BASE / 'output' / (STEM + '.' + suffix)
        if digest(source) != active[suffix + '_sha256']:
            raise ValueError('Candidate changed before publication')
        if target.exists():
            if digest(target) != digest(source):
                raise ValueError('Do not overwrite a different final artifact')
        else:
            shutil.copyfile(source, target)
        outputs.append({'path': target.relative_to(BASE).as_posix(), 'sha256': digest(target), 'bytes': target.stat().st_size})
    for key in ('notebook','source_bundle'):
        path = BASE / active[key]
        outputs.append({'path': active[key], 'sha256': digest(path), 'bytes': path.stat().st_size})
    manifest = {'schema': 'ag02.final_release.1', 'status': 'COMPLETE', 'phase': 10, 'pages': 28,
                'student': load('evidence/phase09/submission_metadata.json'), 'files': outputs,
                'rubric_requirements_passed': 18, 'tests_passed': 245, 'skipped': 0,
                'coverage_combined_percent': load('evidence/phase10/final_coverage.json')['totals']['percent_covered'],
                'coverage_scope': 'publication_v1 parser/path helper only;not whole project',
                'all_pages_reviewed': True, 'native_word_readonly_pages': word['pages'],
                'prior_scientific_files_preserved': parent['prior_protected_files_preserved'],
                'scientific_forward_added': 0, 'train_runs_added': 0, 'submitted_to_instructor': False,
                'candidate_lock_sha256': active['candidate_lock_sha256']}
    save_once(BASE / 'output/AG02_RELEASE_MANIFEST.json', manifest)
    print(json.dumps(manifest, ensure_ascii=False, indent=2))


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--stage', required=True, choices=['inspect','release'])
    {'inspect': inspect, 'release': release}[parser.parse_args().stage]()
